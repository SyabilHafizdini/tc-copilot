#!/usr/bin/env python3
"""PRD filters, the references sheet and the PRD rows of a suite
(run: py tools/test_wiki_suite_prds.py). Synthetic test cases; the only file
written is a workbook in a temp directory."""
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import wiki_suite

APP, PAY = "rental-application", "rental-payment"
TWO = {"schema_version": 2, "prds": {
    APP: {"title": "Rental application", "adopted_version": 2, "staged_version": None},
    PAY: {"title": "Rental payment", "adopted_version": 1, "staged_version": 2}}}


def _tc(i, prds):
    fm = {"type": "Test Case", "id": f"1.1-AC01-{i:02d}", "kind": "sit",
          "status": "active", "module": "/modules/m.md",
          "scenario_id": f"SC-X-{i:02d}", "covers": ["/stories/US-X.md#AC1"],
          "section": "S", "confidence": "High", "title": f"Title {i}",
          "priority": "P1", "technique": "UC",
          "generated_from": {"prd_versions": prds}}
    body = "# Steps\n1. Do.\n\n# Expected Results\n1. Done.\n"
    return (f"testcases/sit/m/{fm['id']}", fm, body)


TCS = [_tc(1, {APP: 2}), _tc(2, {PAY: 1}), _tc(3, {APP: 2, PAY: 1}), _tc(4, {})]


def _old_form_tc():
    """A test case still in the schema-1 form: prd_version, no prd_versions."""
    rel, fm, body = _tc(5, {})
    fm["generated_from"] = {"prd_version": None}
    return (rel, fm, body)


def _concepts(tcs=TCS):
    c = {"modules/m": ({"type": "Module", "title": "Alpha Module",
                        "description": "d"}, "", None),
         "stories/US-X": ({"type": "User Story", "title": "Story X",
                           "description": "§1.1"}, "", None)}
    for rel, fm, body in tcs:
        c[rel] = (fm, body, None)
    return c


def _ids(suite, tcs=TCS):
    selected, _retired, _stale = wiki_suite.select(suite, _concepts(tcs))
    return [fm["id"][-2:] for _r, fm, _b in selected]


def test_no_prd_filter_selects_everything():
    assert _ids({"kind": "sit"}) == ["01", "02", "03", "04"]


def test_include_prds_keeps_test_cases_whose_prds_intersect_the_list():
    assert _ids({"kind": "sit", "include_prds": [PAY]}) == ["02", "03"]
    assert _ids({"kind": "sit", "include_prds": [APP, PAY]}) == ["01", "02", "03"]


def test_exclude_prds_drops_test_cases_that_draw_on_a_listed_prd():
    assert _ids({"kind": "sit", "exclude_prds": [PAY]}) == ["01", "04"]


def test_extra_include_still_overrides_a_prd_filter():
    assert _ids({"kind": "sit", "include_prds": [PAY],
                 "extra_include": ["1.1-AC01-04"]}) == ["02", "03", "04"]


def test_an_old_form_test_case_draws_on_no_prd():
    tcs = TCS + [_old_form_tc()]
    assert wiki_suite.tc_prds(tcs[-1][1]) == set()
    assert _ids({"kind": "sit", "include_prds": [APP]}, tcs) == ["01", "03"]
    assert _ids({"kind": "sit", "exclude_prds": [APP]}, tcs) == ["02", "04", "05"]


def test_an_unknown_prd_id_is_refused_and_the_registered_ones_are_listed():
    wiki_suite.check_prd_filters({"include_prds": [PAY], "exclude_prds": []}, TWO)
    for key in ("include_prds", "exclude_prds"):
        try:
            wiki_suite.check_prd_filters({key: [PAY, "rental-fees"]}, TWO)
        except ValueError as e:
            assert key in str(e) and "rental-fees" in str(e), str(e)
            assert "rental-application, rental-payment" in str(e), str(e)
        else:
            raise AssertionError(f"{key}: an unknown PRD id must be refused")


def test_a_non_list_prd_filter_is_refused():
    try:
        wiki_suite.check_prd_filters({"include_prds": PAY}, TWO)
    except ValueError as e:
        assert "include_prds must be a list" in str(e), str(e)
    else:
        raise AssertionError("a bare string must be refused")


def test_a_prd_filter_element_that_is_not_an_id_is_refused_not_a_traceback():
    for bad in ({"id": PAY}, [PAY], 7, None):
        try:
            wiki_suite.check_prd_filters({"include_prds": [PAY, bad]}, TWO)
        except ValueError as e:
            assert "include_prds names unknown PRD id(s)" in str(e), str(e)
        else:
            raise AssertionError(f"{bad!r} must be refused")


def test_a_references_row_never_says_vNone():
    """The tester's sheet: a PRD with nothing adopted says so, and an id the
    registry does not hold is named as unregistered (wiki.prd_version_text)."""
    reg = {"schema_version": 2, "prds": {"later": {
        "title": "Later PRD", "adopted_version": None, "staged_version": 1}}}

    def tc(pid):
        return ("testcases/sit/m/x", {"type": "Test Case", "id": "x",
                "generated_from": {"prd_versions": {pid: None}}}, "")
    assert wiki_suite.prd_names([tc("later")], reg) == \
        ["Later PRD - no adopted version"]
    assert wiki_suite.prd_names([tc("gone")], reg) == ["gone (unregistered)"]
    assert wiki_suite.prd_docs([tc("later"), tc("gone")], reg, "Proj") == [
        ("PRD", "gone (unregistered)"), ("PRD", "Later PRD - no adopted version")]
    for row in wiki_suite.prd_names([tc("later"), tc("gone")], reg):
        assert "None" not in row, row


def test_references_rows_are_one_per_prd_drawn_on():
    assert wiki_suite.prd_docs(TCS, TWO, "Proj") == [
        ("PRD", "Rental application - v2"), ("PRD", "Rental payment - v1")]
    assert wiki_suite.prd_docs([TCS[1]], TWO, "Proj") == [("PRD", "Rental payment - v1")]


def test_references_fall_back_to_one_project_row_with_no_version():
    assert wiki_suite.prd_docs([TCS[3]], TWO, "Proj") == [("PRD", "Proj")]
    assert wiki_suite.prd_docs([], {}, "Proj") == [("PRD", "Proj")]
    assert wiki_suite.prd_docs([_old_form_tc()], {}, "Proj") == [("PRD", "Proj")]


def _refs_sheet(tcs, manifest):
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "w.xlsx"
        wiki_suite.render_xlsx(tcs, "t", out, concepts=_concepts(tcs),
                               manifest=manifest)
        ws = load_workbook(out)["B - Proj Doc References"]
        return [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=3).value)
                for r in range(8, 11)]


def test_workbook_lists_both_prds_for_a_suite_spanning_both():
    rows = _refs_sheet(TCS, TWO)
    assert rows[0] == ("PRD", "Rental application - v2"), rows
    assert rows[1] == ("PRD", "Rental payment - v1"), rows
    assert rows[2] == (None, None), rows


def test_workbook_narrowed_to_one_prd_lists_only_it():
    selected, _r, _s = wiki_suite.select({"kind": "sit", "include_prds": [PAY],
                                          "exclude_prds": [APP]}, _concepts())
    rows = _refs_sheet(selected, TWO)
    assert rows[0] == ("PRD", "Rental payment - v1") and rows[1] == (None, None), rows


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"[PASS] {name}")
    print("test_wiki_suite_prds OK")
