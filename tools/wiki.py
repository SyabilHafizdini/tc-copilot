#!/usr/bin/env python3
"""wiki — deterministic CLI for the TC co-authoring bundle (tc-copilot-spec §11).

Contains NO LLM calls (spec P1/P5). Every mutating command git-commits as the
configured agent identity (spec §15.1) unless --no-commit is given.

Usage:
  py tools/wiki.py <command> [args]

Commands:
  ingest-prd            inputs/prd/vN/*.{pdf,docx,md} -> sources/prd/ (first
                        version adopts; a NEWER version stages + writes a
                        change report)
  ingest-figma          inputs/figma/*.png  -> sources/figma/*.md (+ manifest)
  ingest-decks          inputs/decks/*.pptx -> sources/decks/<deck>/slide-NN.md
  ingest-reference      inputs/reference/** -> sources/reference/*.md
                        (one Reference Document concept per file; material
                        with no acceptance criteria)
  triage [--apply] [--prd-version N]
                        PUT_FILES_HERE/* -> inputs/{prd/vN,figma,decks}/ by file
                        type; dry-run unless --apply. Refuses on an unstable
                        Figma page name or an occupied PRD version
  index                 regenerate all index.md files from frontmatter
  manifest              rebuild manifest.json from frontmatter
  lint                  L1-L14 / W1-W7; exit 1 on any L error
  status                project status table
  next [--story <id>] [--brief] [--json]   what to DO next per story/flow: the
                        literal next command + the skill that owns it (read-only;
                        --brief skips the hand-edit drift scan, --json emits it
                        structured for the operator app)
  assert story|term|figma|flow <id> --by <user> [--card <path>]
                        HUMAN-GATED assertions; story|flow REQUIRE --card
                        (an emitted card for that scope; asserting stamps the
                        human's answer into it)
  gate --story <id> | --flow <id>   hard generation gate (spec §7.1)
  seal                  hash-seal generated TCs into manifest (spec P6)
  cascade               staleness cascade (spec §12.2): flags only, never regenerates
  export --story <id> | --flow <stem> [--name <n>] [--draft]
                        draft: DRAFT r0 workbook + diff snapshot right after seal;
                        final: refuses without a current strict score, fills Change Log
  rtm                   build/rtm/{matrix.md,trace.md,graph.json,gaps.md} (spec §11.3)
  impact <ref> [--json]  downstream impact of a concept/fragment (read-only; no LLM)
  coverage --story <id> [--propose] [--force]   Component x AC matrix + gaps -> build/rtm/
  testmodel --story <id> | --flow <id> [--propose [--force]]
                        29119-4 test model (coverage items) - scaffold a
                        PROPOSED model from ACs/rules/journey, or show the
                        model + per-technique C = N/T (spec tc-rubric 5)
  suite compile <name>  filter ACTIVE TCs per suites/<name>.yaml -> build/ (spec §9)
  inventory <name>      alias of suite compile
  retire <tc-id> --by <u> --reason voided|superseded [...]   (spec §13)
  void-ac <frag> --by <u> --caused-by <src> --cause-version <n>
  unretire <tc-id> --by <u>         resurrection permit (spec §7.4-4)
  release <tc-id> --by <u> | revert <tc-id>   drift resolution (spec §7.4-5)
  approve-cr CR-NNN --by <u> | reject-cr CR-NNN --by <u>    (spec §12.3)
  diff --prd            adopted vs staged section diff
  migrate-ids           rewrite all TC IDs to the (edited) config template (§10)
  migrate-provenance [--apply]
                        backfill `provenance` on stories written before it was
                        required (L13); stamps prd-verbatim where derived_from
                        exists, reports the rest, dry-run unless --apply
  dashboard             build/status/{dashboard.json,dashboard.html} (spec §16)
  app [--port 8765] [--no-open]     operator app on 127.0.0.1 (needs FastAPI:
                        py -m pip install -r tools/app/requirements.txt)
"""
import hashlib
import json
import os
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

import yaml

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
MANIFEST = ROOT / "manifest.json"

# Test-only: tools/test_wiki_triage.py points the real CLI at a throwaway
# root so the card gate can be proven end-to-end without touching the repo's
# own PUT_FILES_HERE/. Unset in normal operation.
if os.environ.get("TC_ROOT_OVERRIDE"):
    ROOT = Path(os.environ["TC_ROOT_OVERRIDE"]).resolve()
    MANIFEST = ROOT / "manifest.json"

# Files that are the DUMP's own furniture, never content: skipped by both
# wiki_triage.collect_plan (PUT_FILES_HERE/) and cmd_ingest_reference
# (inputs/reference/), so a .gitkeep never becomes a permanent concept id.
FURNITURE = {"README.md", ".gitkeep", ".gitignore", ".DS_Store",
             "Thumbs.db", "desktop.ini"}

CONCEPT_DIRS = [
    "sources/prd", "sources/figma", "sources/decks", "sources/reference",
    "glossary", "modules", "stories", "flows", "resolutions",
]
EDGE_KEYS = [
    "derived_from", "defined_in", "illustrated_by", "module", "flow",
    "stories", "covers", "verifies_rules", "uses_terms", "resolves",
    "supersedes", "superseded_by",
]
REQUIRED_FM = ["type", "id", "title", "description"]

_L10_WIKI_PREFIXES = tuple(f"{d}/" for d in CONCEPT_DIRS + ["testcases", "changereports"])
_L10_TOKEN_RE = re.compile(r"^\+\s*(asserted_by:|status: (?:aligned|superseded|retired))")

def _l10_touches_assertion(diff):
    """True iff the commit adds an assertion/disposition-state line to a WIKI
    concept file. Scoped to wiki .md files so application source and docs that
    merely contain 'asserted_by:'/'status: aligned' as literal text (the
    operator app's TypeScript, a plan doc's code example) do not trip L10."""
    in_wiki_file = False
    for line in diff.splitlines():
        if line.startswith("+++ "):
            path = line[4:].strip()
            if path.startswith("b/"):
                path = path[2:]
            in_wiki_file = path.endswith(".md") and path.startswith(_L10_WIKI_PREFIXES)
        elif in_wiki_file and line.startswith("+") and not line.startswith("+++"):
            if _L10_TOKEN_RE.match(line):
                return True
    return False

# ---------------------------------------------------------------- helpers

def now_iso():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")

def sha256(data):
    if isinstance(data, str):
        data = data.encode("utf-8")
    return "sha256:" + hashlib.sha256(data).hexdigest()

def unsealed_tcs(manifest):
    """TC files whose content is not represented in manifest.tc_hashes — either
    never sealed (a fresh render adds NO entry, so lint W4 cannot see it) or
    changed since sealing. Spec P6: the sealed hashes are what every compiled
    view is entitled to assume it is rendering."""
    hashes = manifest.get("tc_hashes", {})
    base = ROOT / "testcases"
    if not base.exists():
        return []
    out = []
    for p in sorted(base.rglob("*.md")):
        if p.name in ("index.md", "log.md"):
            continue
        rel = p.relative_to(ROOT).as_posix()[:-3]
        if hashes.get(rel) != sha256(p.read_bytes()):
            out.append(rel)
    return sorted(out)


def refuse_if_unsealed(command):
    """Guard for every command that compiles sealed content into an output.

    TC_MANIFEST_OVERRIDE (test-only, unset in normal operation) swaps in an
    alternate manifest.json for this check alone, so tools/smoke.py can prove
    the refusal by invoking the real CLI without ever leaving a real TC
    unsealed."""
    override = os.environ.get("TC_MANIFEST_OVERRIDE")
    manifest = (json.loads(Path(override).read_text(encoding="utf-8"))
                if override else load_manifest())
    un = unsealed_tcs(manifest)
    if un:
        shown = "\n".join(f"    {r}" for r in un[:5])
        more = f"\n    ... and {len(un) - 5} more" if len(un) > 5 else ""
        sys.exit(f"{command} refused: {len(un)} test case(s) not sealed.\n"
                 f"  A render writes TC files; `seal` records their hashes. "
                 f"Compiling now would\n"
                 f"  ship content the manifest does not vouch for (spec P6).\n"
                 f"{shown}{more}\n"
                 f"  Run: py tools/wiki.py seal")

def normalize(text):
    """Spec §5.1.4: collapse whitespace runs, strip trailing space, \\n endings."""
    lines = [re.sub(r"[ \t]+", " ", ln).rstrip() for ln in text.replace("\r\n", "\n").split("\n")]
    return "\n".join(lines).strip()

def load_config():
    return yaml.safe_load((ROOT / "config.yaml").read_text(encoding="utf-8"))

def read_concept(path):
    raw = Path(path).read_text(encoding="utf-8")
    m = re.match(r"^---\n(.*?)\n---\n?(.*)$", raw, re.DOTALL)
    if not m:
        return None, raw
    try:
        return yaml.safe_load(m.group(1)), m.group(2)
    except yaml.YAMLError as e:
        # surfaced as lint L1, never a crash
        print(f"[frontmatter parse error] {path}: {e}", file=sys.stderr)
        return None, m.group(2)

def write_concept(path, fm, body):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fm_text = yaml.safe_dump(fm, sort_keys=False, allow_unicode=True, width=100)
    path.write_text(f"---\n{fm_text}---\n{body}", encoding="utf-8", newline="\n")

def all_concepts():
    """Yield (relpath-no-ext, fm, body, abspath) for every concept file."""
    for d in CONCEPT_DIRS + ["testcases", "changereports"]:
        base = ROOT / d
        if not base.exists():
            continue
        for p in sorted(base.rglob("*.md")):
            if p.name in ("index.md", "log.md"):
                continue
            fm, body = read_concept(p)
            rel = p.relative_to(ROOT).as_posix()[:-3]
            yield rel, fm, body, p

def load_all():
    """(concepts {rel: (fm, body, abspath)}, manifest) — the shared read context."""
    concepts = {rel: (fm, body, p) for rel, fm, body, p in all_concepts()}
    return concepts, load_manifest()

def resolve_ref(ref):
    """'/stories/US-x.md#AC-1' -> (relpath-no-ext, fragment|None)."""
    frag = None
    if "#" in ref:
        ref, frag = ref.split("#", 1)
    rel = ref.lstrip("/")
    if rel.endswith(".md"):
        rel = rel[:-3]
    return rel, frag

def arg_after(args, flag):
    """args[index(flag) + 1], or a clean usage error (never a traceback) if
    `flag` is absent or has no value following it."""
    try:
        return args[args.index(flag) + 1]
    except (ValueError, IndexError):
        sys.exit(f"missing required {flag} <value>")

def fragment_ids(fm):
    ids = set()
    for key in ("acceptance_criteria", "business_rules", "components",
                "branches", "journey"):
        for entry in fm.get(key) or []:
            if isinstance(entry, dict) and "id" in entry:
                ids.add(entry["id"])
    return ids

def fragment_hash_map(fm):
    """Per-fragment normalized hashes (spec §12.1 fragment layer)."""
    out = {}
    for key in ("acceptance_criteria", "business_rules", "components",
                "branches", "journey"):
        for entry in fm.get(key) or []:
            if isinstance(entry, dict) and "id" in entry:
                canon = json.dumps(entry, sort_keys=True, ensure_ascii=False)
                out[entry["id"]] = sha256(normalize(canon))
    return out

def covmap_hash(story_fm, frag):
    """Hash of the component list a story's coverage_map assigns to AC `frag`
    (spec 'Staleness semantics': UAT TCs pin these — coverage_map IS their test
    content). load_coverage imported lazily to avoid a wiki<->wiki_coverage cycle."""
    from wiki_coverage import load_coverage
    cmap = load_coverage(story_fm)[0]
    return sha256(normalize(json.dumps(cmap.get(frag, []), sort_keys=True)))

def load_manifest():
    if MANIFEST.exists():
        return json.loads(MANIFEST.read_text(encoding="utf-8"))
    return {"schema_version": 1, "adopted_prd_version": None, "staged_prd_version": None,
            "id_config_frozen": False, "counters": {}, "sources": {}, "concepts": {},
            "bindings": {}, "tc_hashes": {}, "edges": []}

def save_manifest(m):
    MANIFEST.write_text(json.dumps(m, indent=2, ensure_ascii=False) + "\n",
                        encoding="utf-8", newline="\n")

def agent_commit(msg, no_commit=False):
    if no_commit or "--no-commit" in sys.argv:
        return
    # Correctness is not optional: never record a commit that fails L-series
    # lint. Previously lint ran only inside `gate`, so errors introduced by any
    # mutating command sat in history until someone happened to gate a story.
    # Cheap now that lint is ~3s (see the batched L10 check).
    if "--allow-lint-errors" not in sys.argv:
        r = subprocess.run([sys.executable, __file__, "lint", "--no-commit"],
                           cwd=ROOT, capture_output=True, text=True,
                           encoding="utf-8", errors="replace")
        if r.returncode != 0:
            print("[lint] REFUSING TO COMMIT — L-series errors exist:", file=sys.stderr)
            for ln in r.stdout.splitlines():
                if ln.startswith("ERROR"):
                    print("  " + ln, file=sys.stderr)
            print("  Fix them, then re-run. To commit anyway (you are accepting a\n"
                  "  lint-invalid wiki): re-run with --allow-lint-errors.",
                  file=sys.stderr)
            sys.exit(1)
    cfg = load_config()["provenance"]
    subprocess.run(["git", "add", "-A"], cwd=ROOT, capture_output=True)
    r = subprocess.run(["git", "-c", f"user.name={cfg['agent_git_name']}",
                        "-c", f"user.email={cfg['agent_git_email']}",
                        "commit", "-m", msg], cwd=ROOT, capture_output=True, text=True)
    if r.returncode == 0:
        print(f"[git] committed: {msg}")
    elif "nothing to commit" in (r.stdout + r.stderr):
        print("[git] nothing to commit")
    else:
        print(f"[git] commit failed: {r.stdout}{r.stderr}", file=sys.stderr)

def append_log(entry):
    log = ROOT / "log.md"
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    prev = log.read_text(encoding="utf-8") if log.exists() else "# Project Log\n"
    if not prev.endswith("\n"):
        prev += "\n"
    log.write_text(prev + f"- {stamp} {entry}\n", encoding="utf-8", newline="\n")

# ---------------------------------------------------------------- PDF ingestion

HEAD_MULTI = re.compile(r"^(\d+(?:\.\d+)+)\.?\s+(\S.{1,110})$")
HEAD_TOP = re.compile(r"^(\d+)\.\s+([A-Z][A-Z0-9 /&,'()–-]{3,})$")
TOC_LINE = re.compile(r"\.{4,}\s*\d+\s*$")
NOISE = re.compile(r"^(OFFICIAL \(CLOSED\)|\d{1,3})$")

def table_to_md(rows):
    def clean(c):
        return (c or "").replace("\n", " / ").replace("|", "\\|").strip()
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    out = []
    header = [clean(c) for c in rows[0]] + [""] * (width - len(rows[0]))
    out.append("| " + " | ".join(header) + " |")
    out.append("|" + "---|" * width)
    for r in rows[1:]:
        cells = [clean(c) for c in r] + [""] * (width - len(r))
        out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out)

def extract_pdf_stream(pdf_path):
    """Ordered stream of ('text', line) / ('table', markdown) across all pages."""
    import pdfplumber
    stream = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            items = []
            tables = page.find_tables()
            bboxes = [t.bbox for t in tables]
            for t in tables:
                items.append((t.bbox[1], "table", table_to_md(t.extract())))
            try:
                lines = page.extract_text_lines()
            except Exception:
                lines = []
            for ln in lines:
                cy = (ln["top"] + ln["bottom"]) / 2
                if any(bb[1] - 2 <= cy <= bb[3] + 2 for bb in bboxes):
                    continue
                txt = ln["text"].strip()
                if not txt or NOISE.match(txt) or TOC_LINE.search(txt):
                    continue
                items.append((ln["top"], "text", txt))
            items.sort(key=lambda x: x[0])
            stream.extend((kind, val) for _, kind, val in items)
    return stream

_DOCX_HEADING_STYLE = re.compile(r"^Heading (\d+)$", re.IGNORECASE)


def _kebab(text):
    """Heading text -> a slug component. Used for permanent prd# ids."""
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s or "section"


def extract_docx_stream(docx_path):
    """Ordered stream of ('text', line) / ('table', markdown) /
    ('heading', (num, title, slug)) across a Word document.

    Word auto-numbering means a heading's TEXT carries no number -- only its
    style carries a level. So the displayed number is synthesized from
    per-level counters (matching what Word renders), while the slug -- the
    permanent prd# id -- comes from the heading PATH text instead. Ordinal
    numbers churn when a section is inserted in a later version; heading text
    does not. See the design doc, 'docx as a PRD input'.
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(docx_path))
    stream = []
    counters = {}      # level -> current ordinal
    path_titles = {}   # level -> title text

    # doc.paragraphs and doc.tables are separate lists and lose interleaving;
    # the body's child elements keep document order.
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            table = Table(child, doc)
            rows = [[c.text for c in row.cells] for row in table.rows]
            md = table_to_md(rows)
            if md:
                stream.append(("table", md))
            continue
        if tag != "p":
            continue
        para = Paragraph(child, doc)
        txt = para.text.strip()
        if not txt:
            continue
        m = _DOCX_HEADING_STYLE.match(para.style.name or "")
        if m:
            level = int(m.group(1))
            counters[level] = counters.get(level, 0) + 1
            for deeper in [d for d in counters if d > level]:
                del counters[deeper]
            path_titles[level] = txt
            for deeper in [d for d in list(path_titles) if d > level]:
                del path_titles[deeper]
            num = ".".join(str(counters[d]) for d in sorted(counters))
            slug = "--".join(_kebab(path_titles[d]) for d in sorted(path_titles))
            stream.append(("heading", (num, txt, slug)))
            continue
        if NOISE.match(txt) or TOC_LINE.search(txt):
            continue
        stream.append(("text", txt))
    return stream

def extract_md_stream(md_path):
    """Markdown/text PRD input: every non-empty line is a text element."""
    stream = []
    for ln in Path(md_path).read_text(encoding="utf-8").splitlines():
        txt = ln.rstrip()
        if txt and (NOISE.match(txt.strip()) or TOC_LINE.search(txt)):
            continue
        stream.append(("text", txt))   # blank lines kept: hashes must round-trip
    return stream

def chunk_sections(stream):
    """Heading-hierarchy chunking (spec §5.1.2). -> list of dicts.

    Two heading sources: a ('heading', (num, title, slug)) item, which the
    docx extractor emits because Word carries the level in the style and not
    in the text; or a ('text', ...) line matching HEAD_MULTI/HEAD_TOP, which
    is how numbered pdf/md PRDs have always worked.
    """
    sections = []
    level_titles = {}
    current = (None, "Front Matter", ["Front Matter"], [], None)
    for kind, val in stream:
        heading = None
        explicit_slug = None
        if kind == "heading":
            num, title, explicit_slug = val
            heading = (num, title)
        elif kind == "text":
            m = HEAD_MULTI.match(val) or HEAD_TOP.match(val)
            if m:
                heading = (m.group(1).rstrip("."), m.group(2).strip())
        if heading:
            sections.append(current)
            num, title = heading
            depth = num.count(".") + 1
            level_titles[depth] = f"{num} {title}"
            for d in list(level_titles):
                if d > depth:
                    del level_titles[d]
            hpath = [level_titles[d] for d in sorted(level_titles)]
            current = (num, title, hpath, [], explicit_slug)
        else:
            current[3].append(val)
            if kind == "table":
                current[3].append("")
    sections.append(current)
    out = []
    for num, title, hpath, body_lines, slug in sections:
        body = "\n".join(body_lines).strip()
        if not body:
            continue
        if not slug:
            slug = num.replace(".", "-") if num else "front-matter"
        out.append({"num": num, "slug": slug, "title": title,
                    "heading_path": hpath, "body": body,
                    "content_hash": sha256(normalize(body))})
    return out

def parse_prd(path):
    """PDF, DOCX or MD -> section dicts."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        stream = extract_pdf_stream(path)
    elif ext == ".docx":
        stream = extract_docx_stream(path)
        # A Word PRD whose headings are all plain body text would chunk into
        # ONE Front Matter section and commit cleanly -- silent corruption of
        # every prd# id. Refuse instead.
        if not any(k == "heading" for k, _ in stream):
            sys.exit(
                f"ingest refused: {path.name} has no headings -- every "
                f"paragraph would collapse into one 'Front Matter' section.\n"
                f"  Apply Word's Heading 1/2/3 styles to the section titles "
                f"(or supply a numbered pdf/md) and re-run.")
    else:
        stream = extract_md_stream(path)
    return chunk_sections(stream)

def latest_prd_input():
    versions = sorted((ROOT / "inputs/prd").glob("v*"), key=lambda p: int(p.name[1:]))
    if not versions:
        sys.exit("no inputs/prd/vN directory")
    vdir = versions[-1]
    files = (list(vdir.glob("*.pdf")) + list(vdir.glob("*.docx"))
             + list(vdir.glob("*.md")))
    if not files:
        sys.exit(f"no pdf/docx/md in {vdir}")
    return int(vdir.name[1:]), files[0]

def write_prd_sections(sections, version, src_rel, manifest):
    """Section dicts -> sources/prd/<slug>.md concepts + manifest entries.

    The slug is the PERMANENT prd# id, so two sections sharing one is a lost
    section, not a merge: the second write_concept overwrites the first, the
    manifest entry overwrites, and ingest exits 0 having silently dropped
    content every downstream AC may have cited. Detect that across the WHOLE
    list before writing anything -- a refusal partway through would leave
    orphan concepts behind an untouched manifest.

    The check lives here rather than in the docx extractor so it guards every
    input: docx slugs come from the heading PATH, pdf/md slugs from the
    section NUMBER, and either can repeat in a real document.
    """
    outdir = ROOT / "sources/prd"
    by_slug = {}
    for sec in sections:
        prior = by_slug.get(sec["slug"])
        if prior is not None:
            here = " > ".join(sec["heading_path"]) or sec["title"]
            there = " > ".join(prior["heading_path"]) or prior["title"]
            sys.exit(
                f"ingest refused: two sections both produce the id "
                f"prd#{sec['slug']} --\n"
                f"    {there}\n"
                f"    {here}\n"
                f"  One would silently overwrite the other and the section "
                f"would be lost.\n"
                f"  Rename one of the headings in the source document so "
                f"their heading paths differ, then re-run.")
        by_slug[sec["slug"]] = sec

    outdir.mkdir(parents=True, exist_ok=True)
    written = 0
    for sec in sections:
        sec_id = f"prd#{sec['slug']}"
        fm = {
            "type": "PRD Section",
            "id": sec_id,
            "title": (f"{sec['num']} {sec['title']}" if sec["num"] else sec["title"]),
            "description": f"PRD v{version} section: {sec['title']}",
            "prd_version": version,
            "content_hash": sec["content_hash"],
            "source_file": src_rel,
            "heading_path": sec["heading_path"],
        }
        write_concept(outdir / f"{sec['slug']}.md", fm, sec["body"] + "\n")
        manifest["sources"][sec_id] = {"content_hash": sec["content_hash"],
                                       "prd_version": version}
        written += 1
    return written

def cmd_ingest_prd():
    version, src_path = latest_prd_input()
    manifest = load_manifest()
    adopted = manifest["adopted_prd_version"]
    src_rel = "/" + src_path.relative_to(ROOT).as_posix()
    print(f"ingesting {src_path.name} as PRD v{version} ...")
    sections = parse_prd(src_path)
    if adopted in (None, version):
        written = write_prd_sections(sections, version, src_rel, manifest)
        manifest["adopted_prd_version"] = version
        save_manifest(manifest)
        append_log(f"**Ingestion (agent)**: PRD v{version} -> {written} sections")
        print(f"wrote {written} PRD Section concepts")
        agent_commit(f"ingest(prd): v{version} -> {written} sections")
    else:
        # spec §5.1.5: stage + change report; adopted content untouched
        from wiki_change import stage_prd_version
        stage_prd_version(sections, version, src_rel, manifest)

def cmd_ingest_figma():
    manifest = load_manifest()
    outdir = ROOT / "sources/figma"
    outdir.mkdir(parents=True, exist_ok=True)
    count = 0
    for png in sorted((ROOT / "inputs/figma").glob("*.png")):
        stem = png.stem
        ihash = sha256(png.read_bytes())
        fid = f"figma#{stem}"
        dest = outdir / f"{stem}.md"
        if dest.exists():
            fm, body = read_concept(dest)
            if fm["image_hash"] == ihash:
                continue
            fm["image_hash"] = ihash
            fm["version"] = fm.get("version", 1) + 1
            fm["status"] = "needs-review"
            fm["confirmed_unchanged"] = False
            write_concept(dest, fm, body)
        else:
            fm = {
                "type": "Figma Page", "id": fid, "title": stem,
                "description": f"Figma page screenshot: {stem}",
                "image": "/" + png.relative_to(ROOT).as_posix(),
                "image_hash": ihash, "version": 1, "status": "draft",
                "origin": "agent-proposed", "confirmed_unchanged": False,
            }
            write_concept(dest, fm,
                          "# Visual Description\n\n_(written by the alignment skill "
                          "during the next session touching this page — spec §5.2)_\n")
        manifest["sources"][fid] = {"image_hash": ihash, "version": fm["version"]}
        count += 1
    save_manifest(manifest)
    if count:
        append_log(f"**Ingestion (agent)**: figma -> {count} page(s) ingested/updated")
    print(f"figma pages ingested/updated: {count}")
    agent_commit(f"ingest(figma): {count} page(s)")

def cmd_ingest_decks():
    """pptx -> Deck Slide concepts (spec §5.3). Text extraction via the pptx
    zip's slide XML (<a:t> runs) — no python-pptx dependency."""
    import xml.etree.ElementTree as ET
    import zipfile
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    manifest = load_manifest()
    count = 0
    for pptx in sorted((ROOT / "inputs/decks").glob("*.pptx")):
        deck = pptx.stem
        outdir = ROOT / "sources/decks" / deck
        with zipfile.ZipFile(pptx) as z:
            slide_names = sorted(
                (n for n in z.namelist()
                 if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                key=lambda n: int(re.search(r"(\d+)", n).group(1)))
            for name in slide_names:
                no = int(re.search(r"(\d+)", name).group(1))
                root = ET.fromstring(z.read(name))
                paras = []
                for para in root.iter(f"{A}p"):
                    runs = [t.text or "" for t in para.iter(f"{A}t")]
                    if any(r.strip() for r in runs):
                        paras.append("".join(runs))
                body = "\n".join(paras).strip() or "_(no extractable text)_"
                chash = sha256(normalize(body))
                sid = f"deck#{deck}#slide-{no:02d}"
                fm = {
                    "type": "Deck Slide",
                    "id": sid,
                    "title": f"{deck} slide {no}",
                    "description": (paras[0][:120] if paras else "no text"),
                    "deck": "/" + pptx.relative_to(ROOT).as_posix(),
                    "slide_number": no,
                    "content_hash": chash,
                }
                write_concept(outdir / f"slide-{no:02d}.md", fm, body + "\n")
                manifest["sources"][sid] = {"content_hash": chash}
                count += 1
    save_manifest(manifest)
    if count:
        append_log(f"**Ingestion (agent)**: decks -> {count} slide(s)")
    print(f"deck slides ingested: {count}")
    agent_commit(f"ingest(decks): {count} slide(s)")

_REF_TEXT_EXT = {".yml", ".yaml", ".json", ".csv", ".tsv", ".txt", ".md"}

# Extensions this repo HAS a reader for. The `_(no extractable text)_`
# fallback belongs to everything else and nothing else: a reader that exists
# and then raises means the file is corrupt, locked, or password-protected,
# and recording an empty body for it would hash and commit the corruption.
_REF_READER_EXT = _REF_TEXT_EXT | {".pdf", ".docx", ".xlsx", ".xlsm"}


def reference_slug(rel):
    """Relative path (posix, with extension) -> permanent reference# slug.
    Path-derived so repeated filenames across folders do not collide."""
    stem = rel.rsplit(".", 1)[0] if "." in Path(rel).name else rel
    return re.sub(r"[^A-Za-z0-9]+", "-", stem).strip("-")


def extract_reference_text(path):
    """Best-effort text for a reference document.

    A format with NO reader records the honest fallback rather than
    pretending to have read it (same convention as an empty deck slide). A
    format that HAS a reader and fails is a hard refusal: it is the same
    silent-corruption class the headingless-docx guard exists to prevent,
    and the file would otherwise be hashed, registered and committed as
    though it were empty.
    """
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in _REF_READER_EXT:
        return "_(no extractable text)_"
    try:
        if ext in _REF_TEXT_EXT:
            return path.read_text(encoding="utf-8", errors="replace").strip()
        if ext == ".pdf":
            return "\n".join(v for _k, v in extract_pdf_stream(path)).strip()
        if ext == ".docx":
            return "\n".join(
                v[1] if k == "heading" else v
                for k, v in extract_docx_stream(path)).strip()
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = [["" if c is None else str(c) for c in row]
                    for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if rows:
                parts.append(f"## {ws.title}\n\n{table_to_md(rows)}")
        wb.close()
        return "\n\n".join(parts).strip()
    except Exception as e:                     # noqa: BLE001 - see docstring
        sys.exit(f"ingest refused: {path.name} could not be read as "
                 f"'{ext}' ({type(e).__name__}: {e}).\n"
                 f"  A {ext} reader exists, so an empty body here would hash "
                 f"and commit a corrupt document as though it were blank.\n"
                 f"  Repair or re-export the file (or remove it from "
                 f"inputs/reference/), then re-run.")


# Same bar as a Figma page stem: a reference document's name becomes a
# permanent concept id, so a camera/clipboard artefact is refused.
_JUNK_STEM_REF = re.compile(
    r"^(screenshot|screen[ _-]?shot|pasted|paste|image|img|unnamed|untitled|"
    r"document|new file)[\s_\-]*[\d()\[\]\-_. ]*$", re.IGNORECASE)


def cmd_ingest_reference():
    """inputs/reference/** -> sources/reference/*.md Reference Document
    concepts. One concept per file: reference material is not
    heading-structured, so there is no chunking.

    Every NAME is validated before any concept is written. Writing inside the
    walk meant a junk-stem or collision refusal on file N left N-1 orphan
    concepts on disk behind an untouched manifest -- a half-ingested tree that
    reads as real.
    """
    base = ROOT / "inputs/reference"
    if not base.exists():
        print("no inputs/reference/ -- nothing to ingest.")
        return
    manifest = load_manifest()
    outdir = ROOT / "sources/reference"

    # --- pass 1: names only. Nothing is written until every one is legal. ---
    seen = {}
    queue = []
    for p in sorted(base.rglob("*")):
        if not p.is_file():
            continue
        # The dump's own furniture is not reference material -- matches the
        # skip list wiki_triage.collect_plan applies to PUT_FILES_HERE/.
        if p.name in FURNITURE:
            continue
        rel = p.relative_to(base).as_posix()
        slug = reference_slug(rel)
        if _JUNK_STEM_REF.match(p.stem):
            sys.exit(f"ingest-reference refused: '{p.stem}' is not a stable "
                     f"name, and it becomes this document's permanent id "
                     f"(reference#{slug}).\n  Rename "
                     f"inputs/reference/{rel} to what the document actually "
                     f"is, then re-run.")
        if slug in seen:
            sys.exit(f"ingest-reference refused: '{rel}' and '{seen[slug]}' "
                     f"both produce the id reference#{slug}.\n"
                     f"  Rename one -- merging two documents under one id "
                     f"would lose whichever is ingested second.")
        seen[slug] = rel
        queue.append((p, rel, slug))

    # --- pass 2: read every body. extract_reference_text refuses on a corrupt
    # document that HAS a reader, so this pass stays read-only too. ---
    staged = []
    for p, rel, slug in queue:
        body = extract_reference_text(p) or "_(no extractable text)_"
        staged.append((p, rel, slug, body, sha256(normalize(body))))

    # --- pass 3: write. Nothing above can refuse any more. ---
    count = 0
    for p, rel, slug, body, chash in staged:
        rid = f"reference#{slug}"
        fm = {
            "type": "Reference Document",
            "id": rid,
            "title": p.name,
            "description": f"Reference material: {rel}",
            "source_file": "/" + p.relative_to(ROOT).as_posix(),
            "content_hash": chash,
        }
        write_concept(outdir / f"{slug}.md", fm, body + "\n")
        manifest["sources"][rid] = {"content_hash": chash}
        count += 1
    save_manifest(manifest)
    if count:
        append_log(f"**Ingestion (agent)**: reference -> {count} document(s)")
    print(f"reference documents ingested: {count}")
    agent_commit(f"ingest(reference): {count} document(s)")

# ---------------------------------------------------------------- index / manifest

def cmd_index():
    dirs = []
    for d in CONCEPT_DIRS:
        if (ROOT / d).exists():
            dirs.append(d)
    for sub in sorted((ROOT / "testcases").rglob("*")) if (ROOT / "testcases").exists() else []:
        if sub.is_dir():
            dirs.append(sub.relative_to(ROOT).as_posix())
    if (ROOT / "testcases").exists():
        dirs.append("testcases")
    for d in dirs:
        base = ROOT / d
        entries = []
        for p in sorted(base.glob("*.md")):
            if p.name in ("index.md", "log.md"):
                continue
            fm, _ = read_concept(p)
            if fm:
                entries.append(f"- [{fm.get('title', p.stem)}](/{d}/{p.name}) — "
                               f"{fm.get('description', '')}")
        subdirs = [f"- [{c.name}/](/{d}/{c.name}/index.md)"
                   for c in sorted(base.iterdir()) if c.is_dir()]
        body = f"# {d}\n\n"
        if subdirs:
            body += "## Subdirectories\n\n" + "\n".join(subdirs) + "\n\n"
        if entries:
            body += "## Concepts\n\n" + "\n".join(entries) + "\n"
        (base / "index.md").write_text(body, encoding="utf-8", newline="\n")
    # bundle root index
    cfg = load_config()
    root_body = (f"# {cfg['project']['name']}\n\n"
                 "Test-case co-authoring bundle (OKF). Compiled views live under build/.\n\n"
                 "## Sections\n\n")
    for d in ["sources/prd", "sources/figma", "glossary", "modules", "stories",
              "flows", "resolutions", "testcases"]:
        if (ROOT / d / "index.md").exists():
            root_body += f"- [{d}](/{d}/index.md)\n"
    (ROOT / "index.md").write_text(root_body, encoding="utf-8", newline="\n")
    print("index.md files regenerated")
    agent_commit("index: regenerate")

def cmd_manifest():
    m = load_manifest()
    if m.get("id_config_frozen") and not m.get("id_format"):
        # backfill the frozen format record (pre-dates the id_format field);
        # assumes config.yaml has not been tampered with since the freeze
        m["id_format"] = load_config()["ids"]["tc_format"]
    m["concepts"] = {}
    m["edges"] = []
    for rel, fm, body, p in all_concepts():
        if not fm:
            continue
        entry = {"status": fm.get("status"), "version": fm.get("version")}
        fr = fragment_hash_map(fm)
        if fr:
            entry["fragment_hashes"] = fr
        m["concepts"][rel] = entry
        if fm.get("type") == "PRD Section":
            m["sources"][fm["id"]] = {"content_hash": fm["content_hash"],
                                      "prd_version": fm["prd_version"]}
        if fm.get("type") == "Figma Page":
            m["sources"][fm["id"]] = {"image_hash": fm["image_hash"],
                                      "version": fm.get("version", 1)}
        for key in EDGE_KEYS:
            vals = fm.get(key)
            if not vals:
                continue
            if isinstance(vals, str):
                vals = [vals]
            for v in vals:
                tgt, frag = resolve_ref(v)
                m["edges"].append([rel, key, tgt + (f"#{frag}" if frag else "")])
        for e in fm.get("journey") or []:
            if isinstance(e, dict) and e.get("ref"):
                tgt, frag = resolve_ref(e["ref"])
                m["edges"].append([rel, "journey",
                                   tgt + (f"#{frag}" if frag else "")])
    save_manifest(m)
    print(f"manifest: {len(m['concepts'])} concepts, {len(m['edges'])} edges")
    agent_commit("manifest: rebuild")

# ---------------------------------------------------------------- lint

def body_section(body, heading):
    m = re.search(rf"^# {re.escape(heading)}\s*$(.*?)(?=^# |\Z)", body,
                  re.MULTILINE | re.DOTALL)
    return m.group(1).strip() if m else None

PROSE_LINK = re.compile(r"\]\((/[^)#\s]+\.md)(#[^)\s]+)?\)")

def voided_ac_refs(concepts):
    """Set of '<story-rel>#<ac-id>' for every voided AC."""
    out = set()
    for rel, (fm, _b) in concepts.items():
        if fm and fm.get("type") == "User Story":
            for ac in fm.get("acceptance_criteria") or []:
                if isinstance(ac, dict) and ac.get("status") == "voided":
                    out.add(f"{rel}#{ac['id']}")
    return out

def covered_ac_refs(fm):
    return [(resolve_ref(c)[0], resolve_ref(c)[1]) for c in fm.get("covers") or []]

PROVENANCE_VALUES = ("prd-verbatim", "prd-interpreted", "human-stated")
# Only these two make a claim the story itself must evidence: the AC text is
# not the source's own words, so a human has to have confirmed it.
_PROVENANCE_NEEDS_RESOLUTION = ("prd-interpreted", "human-stated")

def l14_errors(concepts):
    """L14: a story's or flow's `test_model` (29119-4 TD1/TD2, spec
    tc-rubric 5) must be structurally valid: known technique and kind, a
    non-empty basis that resolves to this scope's ACs / rules / journey, a
    boolean feasible, a justification on every infeasible item, at most one
    main scenario. An ABSENT model is not an error (unmigrated scope)."""
    from wiki_rubric import validate_test_model
    out = []
    for rel, (fm, _body) in concepts.items():
        if not fm or fm.get("type") not in ("User Story", "Flow"):
            continue
        for e in validate_test_model(fm, fm.get("id") or rel):
            out.append(f"L14 {rel}: {e}")
    return out


def w7_warnings(concepts):
    """W7: an active test case names a `coverage_items` id that its scope's
    test_model does not define. The engine scores it (T1.1 band 1) but the
    operator reads lint, not the score JSON."""
    known = {}
    for rel, (fm, _b) in concepts.items():
        if fm and fm.get("type") in ("User Story", "Flow"):
            items = ((fm.get("test_model") or {}).get("items") or [])
            known[rel] = {i.get("id") for i in items if isinstance(i, dict)}
    out = []
    for rel, (fm, _b) in concepts.items():
        if not fm or fm.get("type") != "Test Case":
            continue
        if fm.get("status") == "retired" or not fm.get("coverage_items"):
            continue
        scopes = {resolve_ref(r)[0] for r in fm.get("covers") or []}
        pool = set()
        modelled = False
        for s in scopes:
            if s in known:
                modelled = modelled or bool(known[s])
                pool |= known[s]
        if not modelled:
            continue
        for cid in fm["coverage_items"]:
            if cid not in pool:
                out.append(f"W7 {rel}: coverage_items '{cid}' is not in the "
                           f"test_model of {', '.join(sorted(scopes))}")
    return out


def l13_errors(concepts):
    """L13: every User Story DECLARES how its acceptance criteria came to
    exist, and the declaration must match what the story carries.

    Without a declaration, silence is the default and the default is
    "copied from the PRD" -- so the one case that needs a deliberate choice
    is the case where saying nothing works. On a project where some modules
    have PRDs and some do not, that is exploitable with no bad intent: point
    derived_from at any section that happens to exist and agent-authored ACs
    read as sourced ones in every downstream artifact. The RTM traces at
    story level (wiki_rtm), so nothing downstream can tell them apart -- this
    rule is the only place the difference is recorded.
    """
    errors = []
    resolved = {}
    for rel, (fm, _body) in concepts.items():
        if not fm or fm.get("type") != "Resolution":
            continue
        for ref in fm.get("resolves") or []:
            tgt, frag = resolve_ref(ref)
            if frag:
                resolved.setdefault(tgt, set()).add(frag)

    for rel, (fm, _body) in concepts.items():
        if not fm or fm.get("type") != "User Story":
            continue
        prov = fm.get("provenance")
        has_source = bool(fm.get("derived_from"))
        allowed = ", ".join(PROVENANCE_VALUES)

        if not prov:
            errors.append(
                f"L13 {rel}: no 'provenance' -- every story must declare how "
                f"its ACs came to exist: {allowed}")
            continue
        if prov not in PROVENANCE_VALUES:
            errors.append(
                f"L13 {rel}: provenance '{prov}' is not one of {allowed} -- "
                f"declare one of these values to indicate the source of the ACs")
            continue
        if prov == "human-stated" and has_source:
            errors.append(
                f"L13 {rel}: provenance 'human-stated' but derived_from is set "
                f"-- a story with a source is not human-stated; declare "
                f"'prd-verbatim' if the AC text is the PRD's own words, or "
                f"'prd-interpreted' if you wrote it from the PRD's prose")
            continue
        if prov != "human-stated" and not has_source:
            errors.append(
                f"L13 {rel}: provenance '{prov}' but no derived_from -- name "
                f"the PRD section(s) the ACs came from, or declare "
                f"'human-stated'")
            continue

        # Reference material is citable but never a substitute for the PRD.
        # Without this, a story could declare prd-interpreted, point
        # derived_from solely at a rules file, and read as PRD-sourced in
        # every downstream artifact -- the exploit this rule exists to stop.
        if prov != "human-stated":
            from_prd = any(resolve_ref(r)[0].startswith("sources/prd/")
                           for r in fm.get("derived_from") or [])
            if not from_prd:
                errors.append(
                    f"L13 {rel}: provenance '{prov}' but no derived_from ref "
                    f"resolves under sources/prd/ -- acceptance criteria "
                    f"originate in the PRD. Reference material may be cited "
                    f"alongside a PRD section, never instead of one; if there "
                    f"is no PRD behind these ACs, declare 'human-stated'")
                continue

        # AC shape is not a provenance question -- ids are needed for test case
        # covers: refs, the coverage map, and fragment hashing regardless of source.
        have = resolved.get(rel, set()) if prov in _PROVENANCE_NEEDS_RESOLUTION else None
        for ac in fm.get("acceptance_criteria") or []:
            if not isinstance(ac, dict):
                errors.append(
                    f"L13 {rel}: acceptance criterion {ac!r} is not a mapping "
                    f"-- every AC must be {{id: ..., text: ...}} so test cases' "
                    f"covers: refs, the coverage map, and fragment hashes can "
                    f"point at its id")
                continue
            ac_id = ac.get("id")
            if not ac_id:
                errors.append(
                    f"L13 {rel}: acceptance criterion {ac!r} has no 'id' -- "
                    f"every AC must carry an id so test cases' covers: refs, "
                    f"the coverage map, and fragment hashes can point at it")
                continue
            if have is not None and ac_id not in have:
                errors.append(
                    f"L13 {rel}: '{prov}' AC {ac_id} has no Resolution behind "
                    f"it -- capture the human's confirmation in resolutions/ "
                    f"and point resolves: at /{rel}.md#{ac_id}")
    return errors


W6_COVERAGE_THRESHOLD = 0.75

def _tokens(s):
    """Lowercased word-tokens, split on any run of non-alphanumerics.

    This subsumes the dash/pipe/whitespace normalization an earlier version
    of W6 did by hand (a dedicated _squash step): a bare [a-z0-9]+ extract
    already treats a hyphen, an em/en dash, a table pipe, and whitespace
    alike as token boundaries, so there is nothing left for a separate
    normalization pass to do. It is dropped rather than kept-for-safety
    because token coverage does not care where a boundary fell, only which
    words are present -- unlike the substring test it replaced, which cared
    about exact contiguous runs and so needed every boundary variant
    special-cased.
    """
    return re.findall(r"[a-z0-9]+", (s or "").lower())


def w6_warnings(concepts):
    """W6: an AC declared 'prd-verbatim' whose word-tokens are mostly
    absent from the PRD sections its story cites (coverage below
    W6_COVERAGE_THRESHOLD).

    A WARNING, never an error, and deliberately so: 'prd-verbatim' is the
    one provenance value L13 records without being able to check, and this
    is the imperfect check that fills that gap. It measures TOKEN COVERAGE,
    not a substring match, because a contiguous-substring test is strictly
    stricter than the definition of 'prd-verbatim' it is meant to police --
    tc-align Phase A authors a verbatim AC by reflowing a PRD table cell
    into one prose sentence and stripping the PDF extractor's line-wrap
    markers, while the stored section body keeps the table's pipe-delimited
    rows and the cell's own internal breaks, so the AC's words are the
    section's words, same order even, but almost never one contiguous run.
    Measured against 37 real prd-verbatim ACs, a substring test produced 37
    warnings (100% false positives; longest contiguous common run on one
    real AC was 36 of 217 characters). Token coverage, measured on the same
    content, does not have that failure mode: min 0.82 / median 1.00 across
    those 37 real ACs, versus 0.14 and 0.00 on invented ACs.

    Coverage = the fraction of the AC's word-tokens (each occurrence, not
    deduplicated) that also appear in the union of tokens from the sections
    the story cites. Warn when coverage is below W6_COVERAGE_THRESHOLD.

    Short ACs are skipped: too few tokens for coverage to mean anything.
    This is a deliberate false-NEGATIVE trade, not a neutral one -- a short
    invented AC gets no signal from L13 (which does not check AC text at
    all) or from W6 (which skips it here) -- accepted because a confident
    false positive on real content is worse than a quiet miss on a short
    one.

    A story is skipped entirely (no warning for any of its ACs) if none of
    its derived_from refs resolved to a concept in this bundle -- that is an
    L2/L13 fault (a bad or missing derived_from), not a W6 one, and warning
    here with 0% coverage would blame the wrong rule for it.
    """
    warns = []
    for rel, (fm, _body) in concepts.items():
        if not fm or fm.get("type") != "User Story":
            continue
        if fm.get("provenance") != "prd-verbatim":
            continue
        hay_tokens = set()
        any_resolved = False
        for ref in fm.get("derived_from") or []:
            tgt, _frag = resolve_ref(ref)
            if tgt in concepts:
                any_resolved = True
                hay_tokens.update(_tokens(concepts[tgt][1]))
        if not any_resolved:
            continue
        for ac in fm.get("acceptance_criteria") or []:
            if not isinstance(ac, dict):
                continue
            text = ac.get("text") or ""
            if len(text.strip()) < 25:
                continue
            ac_tokens = _tokens(text)
            if not ac_tokens:
                continue
            covered = sum(1 for t in ac_tokens if t in hay_tokens)
            coverage = covered / len(ac_tokens)
            if coverage < W6_COVERAGE_THRESHOLD:
                warns.append(
                    f"W6 {rel}: AC {ac.get('id')} is declared prd-verbatim "
                    f"but only {coverage:.0%} of its words appear in the PRD "
                    f"sections this story cites -- if you wrote it rather "
                    f"than copied it, the story is 'prd-interpreted'")
    return warns


def cmd_lint():
    errors, warns = [], []
    concepts = {rel: (fm, body) for rel, fm, body, _ in all_concepts()}
    manifest = load_manifest()
    cfg = load_config()
    voided = voided_ac_refs(concepts)

    for rel, (fm, body) in concepts.items():
        if fm is None:
            errors.append(f"L1 {rel}: missing/unparseable frontmatter")
            continue
        for req in REQUIRED_FM:
            if not fm.get(req):
                errors.append(f"L1 {rel}: missing required field '{req}'")
        for key in EDGE_KEYS:
            vals = fm.get(key)
            if not vals:
                continue
            if isinstance(vals, str):
                vals = [vals]
            for v in vals:
                tgt, frag = resolve_ref(v)
                if tgt not in concepts:
                    errors.append(f"L2 {rel}: {key} -> {v} (concept not found)")
                elif frag and frag not in fragment_ids(concepts[tgt][0]):
                    errors.append(f"L2 {rel}: {key} -> {v} (fragment not found)")

        if fm.get("type") == "Test Case":
            if not fm.get("covers"):
                errors.append(f"L8 {rel}: covers is empty")
            er = body_section(body, "Expected Results")
            if not er:
                errors.append(f"L8 {rel}: Expected Results section empty/missing")
            # L3: TC exists for a story that was never aligned
            if fm.get("status") in ("active", "stale"):
                for srel, _frag in covered_ac_refs(fm):
                    sfm = concepts.get(srel, (None, None))[0]
                    if sfm and sfm.get("status") in ("draft", "in-alignment"):
                        errors.append(f"L3 {rel}: active TC covers never-aligned "
                                      f"story {srel} ({sfm.get('status')})")
                        break
            ret = fm.get("retirement")
            if fm.get("status") == "retired" and ret:
                reason = ret.get("reason")
                if reason == "superseded":
                    succ_ref = ret.get("superseded_by")
                    if not succ_ref:
                        errors.append(f"L5 {rel}: retired:superseded without superseded_by")
                    else:
                        srel_, _f = resolve_ref(succ_ref)
                        succ = concepts.get(srel_, (None, None))[0]
                        if not succ or succ.get("status") != "active":
                            errors.append(f"L5 {rel}: superseded_by {succ_ref} is not "
                                          "an active TC")
                        else:
                            mine = {f"{s}#{f}" for s, f in covered_ac_refs(fm)
                                    if f"{s}#{f}" not in voided}
                            theirs = {f"{s}#{f}" for s, f in covered_ac_refs(succ)}
                            missing = mine - theirs
                            if missing:
                                errors.append(f"L5 {rel}: successor misses active ACs "
                                              f"{sorted(missing)}")
                elif reason == "voided":
                    if ret.get("superseded_by"):
                        errors.append(f"L5 {rel}: retired:voided must not have "
                                      "superseded_by (reason asymmetry)")
                    if not ret.get("caused_by") or not ret.get("cause_version"):
                        errors.append(f"L5 {rel}: retired:voided requires caused_by "
                                      "+ cause_version")
                    for s, f in covered_ac_refs(fm):
                        if f and f"{s}#{f}" not in voided:
                            errors.append(f"L6 {rel}: retired:voided but covered AC "
                                          f"{s}#{f} is still active")
                else:
                    errors.append(f"L5 {rel}: retirement reason must be "
                                  "voided|superseded")
            # L6 (other direction): active TC whose covers are ALL voided —
            # the void cascade should have retired it
            if fm.get("status") == "active" and fm.get("covers"):
                refs = [f"{s}#{f}" for s, f in covered_ac_refs(fm) if f]
                if refs and all(r in voided for r in refs):
                    errors.append(f"L6 {rel}: active TC covers only voided ACs "
                                  "(void cascade missed it)")

        if fm.get("type") == "Flow" and fm.get("status") == "aligned":
            for sref in fm.get("stories") or []:
                tgt, _f = resolve_ref(sref)
                sfm = concepts.get(tgt, (None, None))[0]
                if sfm and sfm.get("status") != "aligned":
                    errors.append(f"L4 {rel}: flow aligned but member story {sref} "
                                  f"is '{sfm.get('status')}'")

        if fm.get("type") == "Flow":
            members = {resolve_ref(s)[0] for s in fm.get("stories") or []}
            for e in fm.get("journey") or []:
                ref = e.get("ref") if isinstance(e, dict) else None
                if not ref or not isinstance(e, dict) or not e.get("id"):
                    errors.append(f"L12 {rel}: journey entry missing id/ref")
                    continue
                if not e.get("end_state"):
                    errors.append(f"L12 {rel}: journey entry {e.get('id')} "
                                  "missing end_state (schema requires it)")
                tgt, frag = resolve_ref(ref)
                if tgt not in concepts:
                    errors.append(f"L2 {rel}: journey -> {ref} (concept not found)")
                    continue
                if frag and frag not in fragment_ids(concepts[tgt][0]):
                    errors.append(f"L2 {rel}: journey -> {ref} (fragment not found)")
                if tgt not in members:
                    errors.append(f"L12 {rel}: journey ref {ref} targets a story "
                                  "outside the flow's stories list")

        for a, b in (("supersedes", "superseded_by"), ("superseded_by", "supersedes")):
            vals = fm.get(a)
            for v in ([vals] if isinstance(vals, str) else vals or []):
                tgt, _ = resolve_ref(v)
                other = concepts.get(tgt)
                if other and other[0]:
                    back = other[0].get(b) or []
                    back = [back] if isinstance(back, str) else back
                    if not any(resolve_ref(x)[0] == rel for x in back):
                        errors.append(f"L9 {rel}: {a} -> {v} lacks reciprocal {b}")
        # retirement.superseded_by reciprocity (L9 flavour of §13.2)
        ret = fm.get("retirement") or {}
        if ret.get("superseded_by"):
            tgt, _ = resolve_ref(ret["superseded_by"])
            other = concepts.get(tgt, (None, None))[0]
            if other:
                back = other.get("supersedes") or []
                back = [back] if isinstance(back, str) else back
                if not any(resolve_ref(x)[0] == rel for x in back):
                    errors.append(f"L9 {rel}: retirement.superseded_by lacks "
                                  "reciprocal supersedes on successor")

        if fm.get("type") == "User Story" and fm.get("status") == "aligned":
            for src, pin in (fm.get("source_pins") or {}).items():
                cur = manifest["sources"].get(src, {})
                cur_hash = cur.get("content_hash") or cur.get("image_hash")
                if cur_hash and cur_hash != pin:
                    warns.append(f"W2 {rel}: aligned but pinned {src} hash differs")
            for t in fm.get("uses_terms") or []:
                tgt, _ = resolve_ref(t)
                if tgt in concepts and concepts[tgt][0] and \
                        concepts[tgt][0].get("status") == "proposed":
                    warns.append(f"W3 {rel}: uses proposed term {t}")

        # W1: broken prose links in bodies
        for m in PROSE_LINK.finditer(body or ""):
            path = m.group(1)
            tgt, _ = resolve_ref(path)
            if tgt not in concepts and not (ROOT / path.lstrip("/")).exists():
                warns.append(f"W1 {rel}: broken prose link {path}")

    errors.extend(l13_errors(concepts))
    errors.extend(l14_errors(concepts))
    warns.extend(w7_warnings(concepts))
    warns.extend(w6_warnings(concepts))

    # L7: duplicate glossary alias/title coverage
    seen_names = {}
    for rel, (fm, _b) in concepts.items():
        if fm and fm.get("type") == "Glossary Term":
            names = [fm.get("title", "")] + list(fm.get("aliases") or [])
            for n in names:
                key = n.strip().lower()
                if not key:
                    continue
                if key in seen_names and seen_names[key] != rel:
                    errors.append(f"L7 {rel}: alias/title '{n}' already on "
                                  f"{seen_names[key]}")
                seen_names.setdefault(key, rel)

    # L10: tc-agent commits touching assertion/disposition state need an
    # assertion event. Covers alignment (+asserted_by / +status: aligned) AND
    # human-gated retirement/supersession (+status: superseded / +status: retired).
    # Short hashes of commits made before L10 existed. Empty on a fresh
    # project; a project branch adds its own if it carries pre-rule history.
    L10_EXEMPT = set()
    agent_name = cfg["provenance"]["agent_git_name"]
    # ONE `git log -p` for all 200 commits, not one `git show` each: on Windows a
    # process spawn costs ~57ms, so the per-commit loop made `lint` (and every
    # `gate`, which shells out to lint) take 30s+. Same parse, ~1 subprocess.
    CSEP, DSEP = "==COMMIT==", "==DIFF=="
    r = subprocess.run(["git", "log", f"--author={agent_name}", "-n", "200",
                        f"--format={CSEP}%H%n%B%n{DSEP}", "--unified=0", "-p"],
                       cwd=ROOT, capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    for chunk in r.stdout.split(CSEP)[1:]:
        h, _, rest = chunk.partition("\n")
        h = h.strip()
        if any(h.startswith(ex) for ex in L10_EXEMPT):
            continue
        msg, _, diff = rest.partition(DSEP)
        if _l10_touches_assertion(diff) and "Assertion-Event:" not in msg:
            errors.append(f"L10 commit {h[:8]}: tc-agent set assertion/disposition "
                          "state without an Assertion-Event reference")

    # ID-template freeze (spec §10): frozen format must not drift from manifest
    if manifest.get("id_config_frozen") and manifest.get("id_format") and \
            cfg["ids"]["tc_format"] != manifest["id_format"]:
        errors.append("L11 config.yaml: ids.tc_format changed after freeze — "
                      "use 'wiki migrate-ids' (spec §10)")

    # W4 drift + W5 suite refs
    for tcrel, h in manifest.get("tc_hashes", {}).items():
        p = ROOT / (tcrel + ".md")
        if p.exists() and sha256(p.read_bytes()) != h:
            warns.append(f"W4 {tcrel}: file hash differs from manifest (hand-edit drift)")
    tc_ids = {fm["id"]: (rel, fm) for rel, (fm, _b) in concepts.items()
              if fm and fm.get("type") == "Test Case"}
    for spath in sorted((ROOT / "suites").glob("*.yaml")) if (ROOT / "suites").exists() else []:
        suite = yaml.safe_load(spath.read_text(encoding="utf-8")) or {}
        for tcid in suite.get("extra_include") or []:
            if tcid not in tc_ids:
                warns.append(f"W5 suites/{spath.name}: extra_include references "
                             f"nonexistent TC {tcid}")
            elif tc_ids[tcid][1].get("status") == "retired":
                warns.append(f"W5 suites/{spath.name}: extra_include references "
                             f"retired TC {tcid}")

    for e in errors:
        print("ERROR", e)
    for w in warns:
        print("warn ", w)
    print(f"lint: {len(errors)} error(s), {len(warns)} warning(s)")
    sys.exit(1 if errors else 0)

# ---------------------------------------------------------------- status / gate

def story_tc_stats(manifest, story_rel):
    tcs = {"active": 0, "stale": 0, "retired": 0}
    for rel, entry in manifest["concepts"].items():
        if rel.startswith("testcases/"):
            covered = [t for s, k, t in manifest["edges"]
                       if s == rel and k == "covers" and t.startswith(story_rel + "#")]
            if covered and entry.get("status") in tcs:
                tcs[entry["status"]] += 1
    return tcs

def cmd_status():
    manifest = load_manifest()
    print(f"{'story':<16} {'status':<13} {'ACs':>4} {'openQ':>5} "
          f"{'TC act/stale/ret':>16}  title")
    for rel, fm, body, _ in all_concepts():
        if not fm or fm.get("type") != "User Story":
            continue
        acs = len(fm.get("acceptance_criteria") or [])
        oq = body_section(body, "Open Questions") or ""
        openq = len([ln for ln in oq.splitlines() if ln.strip().startswith("-")])
        t = story_tc_stats(manifest, rel)
        print(f"{fm['id']:<16} {fm.get('status', '?'):<13} {acs:>4} {openq:>5} "
              f"{t['active']:>5}/{t['stale']}/{t['retired']:<6}  {fm.get('title', '')[:50]}")
    for rel, fm, body, _ in all_concepts():
        if fm and fm.get("type") == "Figma Page" and fm.get("status") == "needs-review":
            print(f"NEEDS-REVIEW figma: {fm['id']}")
    print(f"adopted PRD version: {manifest.get('adopted_prd_version')}")

def cmd_gate(args):
    concepts = {rel: (fm, body) for rel, fm, body, _ in all_concepts()}
    blocked = []
    if "--flow" in args:
        ident = arg_after(args, "--flow")
        rel = f"flows/{ident}"
        if rel not in concepts:
            sys.exit(f"gate: flow {ident} not found")
        fm, _ = concepts[rel]
        if fm.get("status") != "aligned":
            blocked.append(f"flow {ident} is '{fm.get('status')}', not aligned (spec §7.1)")
        scope_stories = [resolve_ref(s)[0] for s in fm.get("stories") or []]
    else:
        ident = arg_after(args, "--story")
        rel = f"stories/{ident}"
        if rel not in concepts:
            sys.exit(f"gate: story {ident} not found")
        scope_stories = [rel]
    for srel in scope_stories:
        sfm, _b = concepts[srel]
        if sfm.get("status") != "aligned":
            blocked.append(f"story {srel} is '{sfm.get('status')}', not aligned (spec §7.1)")
        if not (sfm.get("acceptance_criteria") or []):
            blocked.append(f"story {srel} has no acceptance_criteria -- generation "
                           f"is never where requirements are born; align it first "
                           f"(tc-align)")
        for ref in sfm.get("illustrated_by") or []:
            tgt, _f = resolve_ref(ref)
            if tgt in concepts and concepts[tgt][0].get("status") == "needs-review":
                blocked.append(f"figma page {ref} is needs-review")
    r = subprocess.run([sys.executable, __file__, "lint", "--no-commit"],
                       capture_output=True, text=True)
    if r.returncode != 0:
        blocked.append("L-series lint errors exist:\n" +
                       "\n".join(ln for ln in r.stdout.splitlines() if ln.startswith("ERROR")))
    if blocked:
        print("GATE BLOCKED:")
        for b in blocked:
            print(" -", b)
        sys.exit(1)
    print(f"GATE OPEN: generation permitted for {ident}")

# ---------------------------------------------------------------- assert (human-gated)

CARD_DIR = ROOT / "build/cards"


def _card_open_q_ids(card):
    z = card.get("zones") or {}
    oq = card.get("open_questions") or z.get("open_questions") or []
    return {q.get("id") for q in oq}


def cmd_card(args):
    if not args or args[0] not in ("revise", "discard"):
        sys.exit("card: expected 'revise' or 'discard' <card> --by <who>")
    sub = args[0]
    if len(args) < 2:
        sys.exit(f"card {sub}: a card path is required")
    card_path = Path(args[1])
    if not card_path.is_absolute():
        card_path = ROOT / card_path
    by = arg_after(args, "--by") if "--by" in args else None
    if not by:
        sys.exit(f"card {sub}: --by <who> is required (only humans decide)")
    note = arg_after(args, "--note") if "--note" in args else None
    if not card_path.exists():
        sys.exit(f"card {sub} refused: card not found: {card_path}")
    try:
        card = json.loads(card_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        sys.exit(f"card {sub} refused: {card_path.name} is not valid JSON ({e})")
    if card.get("human_response"):
        sys.exit(f"card {sub} refused: {card_path.name} already answered "
                 f"({card['human_response'].get('answer')})")
    hr = {"answer": sub, "by": by, "at": now_iso()}
    if sub == "revise":
        valid = _card_open_q_ids(card)
        answers = {}
        i = 0
        while i < len(args):
            if args[i] == "--answer":
                pair = args[i + 1] if i + 1 < len(args) else ""
                qid, _, value = pair.partition("=")
                if qid not in valid:
                    sys.exit(f"card revise refused: unknown question id {qid!r} "
                             f"(card has {sorted(valid)})")
                answers[qid] = ({"decision": "accept", "value": None}
                                if value == "accept"
                                else {"decision": "correct", "value": value})
                i += 2
            else:
                i += 1
        hr["answers"] = answers
    if note is not None:
        hr["note"] = note
    card["human_response"] = hr
    card_path.write_text(json.dumps(card, indent=1, ensure_ascii=False) + "\n",
                         encoding="utf-8", newline="\n")
    print(f"card {sub}: {card_path.name} by {by}")
    if sub == "discard":
        session = card.get("session") or ""
        commits = []
        if session:
            r = subprocess.run(["git", "log", "--oneline", "-n", "50",
                                f"--grep={session}"], cwd=ROOT,
                               capture_output=True, text=True,
                               encoding="utf-8", errors="replace")
            commits = [ln for ln in r.stdout.splitlines() if ln.strip()]
        if commits:
            print(f"  session {session} produced these commits (undo manually "
                  f"until embedded chat lands):")
            for ln in commits:
                print("    " + ln)
        else:
            print(f"  no commits matched session {session!r} — nothing to undo, "
                  f"or undo manually.")


def cmd_session(args):
    if not args or args[0] != "revert":
        sys.exit("session: expected 'revert' <session-id> --by <who>")
    if len(args) < 2 or args[1].startswith("--"):
        sys.exit("session revert: a session id is required")
    session = args[1]
    by = arg_after(args, "--by") if "--by" in args else None
    if not by:
        sys.exit("session revert: --by <who> is required (only humans decide)")

    st = subprocess.run(["git", "status", "--porcelain"], cwd=ROOT,
                        capture_output=True, text=True,
                        encoding="utf-8", errors="replace")
    if st.stdout.strip():
        sys.exit("session revert refused: working tree is dirty -- commit or "
                 "clean it first (the undo must apply to a known state)")

    log = subprocess.run(["git", "log", "--format=%H%x00%s", "-n", "50",
                          f"--grep={session}"], cwd=ROOT,
                         capture_output=True, text=True,
                         encoding="utf-8", errors="replace")
    # --grep is an unanchored substring match, and our OWN generated revert
    # commit message embeds the session id ("revert: undo session <id> ...").
    # Without this guard, reverting the same session twice would match that
    # prior revert commit and try to revert the revert. Skip our own reverts.
    shas = []
    for ln in log.stdout.splitlines():
        if not ln.strip():
            continue
        sha, _, subject = ln.partition("\x00")
        if subject.startswith("revert: undo session"):
            continue
        shas.append(sha)
    if not shas:
        sys.exit(f"session revert refused: no commits matched session "
                 f"{session!r} -- nothing to undo")
    # `git log -n 50` caps the scan. Hitting the cap means more than 50 commits
    # may carry this session id and only the newest 50 are being reverted --
    # warn rather than silently doing a partial revert.
    if len(shas) == 50:
        print(f"session revert warning: reached the 50-commit cap for session "
              f"{session!r} -- more commits may match; only the newest 50 are "
              f"being reverted. Re-run to undo any older ones.")

    # Newest-first is the correct order for a clean sequential revert. Apply
    # all inverses to the index WITHOUT committing, then hand the single
    # collapsed commit to agent_commit so it passes through the L-series lint
    # gate exactly like every other mutating verb.
    rv = subprocess.run(["git", "revert", "--no-commit", *shas], cwd=ROOT,
                        capture_output=True, text=True,
                        encoding="utf-8", errors="replace")
    if rv.returncode != 0:
        subprocess.run(["git", "revert", "--abort"], cwd=ROOT,
                       capture_output=True)
        subprocess.run(["git", "reset", "--hard", "HEAD"], cwd=ROOT,
                       capture_output=True)
        sys.exit(f"session revert refused: git revert conflicted -- another "
                 f"session touched the same lines:\n{rv.stdout}{rv.stderr}")

    agent_commit(f"revert: undo session {session} ({len(shas)} commit(s)) "
                 f"by {by}\n\nAssertion-Event: cli-session-revert {session} "
                 f"by {by}")
    print(f"session revert: undid {len(shas)} commit(s) from {session} by {by}")
    for s in shas:
        print("  reverted " + s[:9])


def require_card(kind, ident, args, by):
    """The card fence (spec P3). `assert story|flow` is the moment agent-proposed
    content becomes asserted truth, and the protocol says a card must be emitted
    and presented to a human FIRST. That ordering was previously enforced by
    nothing — the CLI wrote its own `Assertion-Event:` trailer, which satisfied
    lint L10 automatically, so an agent could assert with no card in existence.

    Now: --card must name an existing card for this scope, and asserting stamps
    the human's answer into it. This cannot prove a human spoke; it does make the
    skipped step impossible to skip.
    """
    if "--card" not in args:
        cards = sorted(p.name for p in CARD_DIR.glob("*.json")
                       if ident in p.name) if CARD_DIR.exists() else []
        sys.exit(
            f"assert refused: --card <path> is required for '{kind}' (spec P3).\n"
            f"  A card must be emitted and presented to a human before assertion.\n"
            f"  Cards for {ident}: "
            f"{', '.join(cards) if cards else '(none — emit one first)'}\n"
            f"  Example: py tools/wiki.py assert {kind} {ident} --by {by} "
            f"--card build/cards/<card>.json")
    card_path = Path(arg_after(args, "--card"))
    if not card_path.is_absolute():
        card_path = ROOT / card_path
    if not card_path.exists():
        sys.exit(f"assert refused: card not found: {card_path}")
    try:
        card = json.loads(card_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        sys.exit(f"assert refused: card {card_path.name} is not valid JSON ({e})")
    # Cards in the wild carry either a bare id ("US-XXXX", coverage cards) or a
    # rel path ("stories/US-XXXX", session cards) — compare on the basename.
    scope = card.get("story") or card.get("flow") or ""
    if scope.rsplit("/", 1)[-1] != ident:
        sys.exit(f"assert refused: card {card_path.name} is for '{scope}', "
                 f"not '{ident}' — wrong card.")
    return card_path, card


def confirm_test_model(fm, card, ident, by):
    """The card the human answered carries `test_model`; the assertion is the
    moment it becomes confirmed (spec tc-rubric 5: no second gate). The card's
    item ids must be exactly the story's proposed ids - a card that shows the
    human one model while the file holds another is refused."""
    if not card or "test_model" not in card:
        return
    tm = fm.get("test_model")
    if not isinstance(tm, dict):
        sys.exit(f"assert refused: card carries a test_model but {ident} has "
                 f"none - run `wiki testmodel --propose` and re-emit the card")
    if tm.get("status") == "confirmed":
        return
    card_ids = [i.get("id") for i in
                ((card.get("test_model") or {}).get("items") or [])
                if isinstance(i, dict)]
    file_ids = [i.get("id") for i in (tm.get("items") or [])
                if isinstance(i, dict)]
    if sorted(card_ids) != sorted(file_ids):
        sys.exit(f"assert refused: the card's test_model items "
                 f"{sorted(card_ids)} are not the ones in {ident} "
                 f"{sorted(file_ids)} - the human confirmed a different model")
    from wiki_rubric import validate_test_model
    errs = validate_test_model(fm, ident)
    if errs:
        sys.exit("assert refused: test_model does not validate (lint L14):\n  "
                 + "\n  ".join(errs))
    tm["status"] = "confirmed"
    tm["asserted_by"] = by
    tm["asserted_at"] = now_iso()


def stamp_card(card_path, card, by):
    """Record the human's answer in the card — it is the audit artifact."""
    card["human_response"] = {"answer": "assert", "by": by, "at": now_iso()}
    card_path.write_text(json.dumps(card, indent=1, ensure_ascii=False) + "\n",
                         encoding="utf-8", newline="\n")


def cmd_assert(args):
    kind = args[0]
    ident = args[1]
    by = arg_after(args, "--by") if "--by" in args else None
    if not by:
        sys.exit("assert: --by <user> is required (only humans assert — spec P3)")
    card_path = card = None
    if kind in ("story", "flow"):
        card_path, card = require_card(kind, ident, args, by)
    manifest = load_manifest()
    if kind == "story":
        p = ROOT / "stories" / f"{ident}.md"
        fm, body = read_concept(p)
        oq = body_section(body, "Open Questions") or ""
        if any(ln.strip().startswith("-") for ln in oq.splitlines()):
            sys.exit(f"assert refused: {ident} has open questions (spec §6.1 Phase C)")
        fm["status"] = "aligned"
        fm["asserted_by"] = by
        fm["asserted_at"] = now_iso()
        pins = {}
        for ref in (fm.get("derived_from") or []) + (fm.get("illustrated_by") or []):
            tgt, _ = resolve_ref(ref)
            tfm, _b = read_concept(ROOT / (tgt + ".md"))
            if tfm:
                pins[tfm["id"]] = tfm.get("content_hash") or tfm.get("image_hash")
        fm["source_pins"] = pins
        confirm_test_model(fm, card, ident, by)
        write_concept(p, fm, body)
        append_log(f"**Assertion ({by})**: story {ident} aligned")
    elif kind == "term":
        p = ROOT / "glossary" / f"{ident}.md"
        fm, body = read_concept(p)
        fm["status"] = "aligned"
        fm["asserted_by"] = by
        fm["asserted_at"] = now_iso()
        write_concept(p, fm, body)
        append_log(f"**Assertion ({by})**: term {ident} confirmed")
    elif kind == "figma":
        p = ROOT / "sources/figma" / f"{ident}.md"
        fm, body = read_concept(p)
        if "--unchanged" in args:
            fm["confirmed_unchanged"] = True
        fm["status"] = "asserted"
        fm["asserted_by"] = by
        fm["asserted_at"] = now_iso()
        write_concept(p, fm, body)
        append_log(f"**Assertion ({by})**: figma {ident} confirmed")
    elif kind == "flow":
        p = ROOT / "flows" / f"{ident}.md"
        fm, body = read_concept(p)
        for sref in fm.get("stories") or []:
            tgt, _ = resolve_ref(sref)
            sfm, _b = read_concept(ROOT / (tgt + ".md"))
            if sfm.get("status") != "aligned":
                sys.exit(f"assert refused: flow member story {sref} is "
                         f"'{sfm.get('status')}', not aligned (lint L4)")
        oq = body_section(body, "Open Questions") or ""
        if any(ln.strip().startswith("-") for ln in oq.splitlines()):
            sys.exit(f"assert refused: {ident} has open questions")
        fm["status"] = "aligned"
        fm["asserted_by"] = by
        fm["asserted_at"] = now_iso()
        confirm_test_model(fm, card, ident, by)
        write_concept(p, fm, body)
        append_log(f"**Assertion ({by})**: flow {ident} aligned")
    else:
        sys.exit(f"assert: unknown kind {kind}")
    save_manifest(manifest)
    if card_path:
        stamp_card(card_path, card, by)
    print(f"asserted {kind} {ident} by {by}"
          + (f" (card {card_path.name})" if card_path else ""))
    evidence = f" card {card_path.name}" if card_path else ""
    agent_commit(f"assert({kind}): {ident} by {by}\n\n"
                 f"Assertion-Event: cli-assert {kind} {ident} by {by} "
                 f"at {now_iso()}{evidence}")

# ---------------------------------------------------------------- seal / cascade

TC_SECTIONS = ["Objective", "Preconditions", "Test Data", "Steps",
               "Expected Results", "Postconditions", "Traceability"]

def cmd_seal():
    manifest = load_manifest()
    concepts = {rel: (fm, body) for rel, fm, body, _ in all_concepts()}
    sealed = 0
    for rel, fm, body, p in all_concepts():
        if not fm or fm.get("type") != "Test Case":
            continue
        missing = [s for s in TC_SECTIONS if body_section(body, s) is None]
        if missing:
            sys.exit(f"seal refused: {rel} missing sections {missing} (spec §4.9)")
        sc = fm.get("scenario_id")
        if not sc:
            sys.exit(f"seal refused: {rel} has no scenario_id")
        prior = manifest["bindings"].get(sc)
        if prior and prior["tc"] != rel:
            sys.exit(f"seal refused: scenario {sc} already bound to {prior['tc']} "
                     "(bindings are permanent — spec §7.4)")
        pins = {}
        for ref in (fm.get("covers") or []) + (fm.get("verifies_rules") or []):
            tgt, frag = resolve_ref(ref)
            fr = fragment_hash_map(concepts[tgt][0]) if tgt in concepts else {}
            if frag in fr:
                pins[f"{tgt}#{frag}"] = fr[frag]
        # UAT TCs: the coverage_map IS their test content (element lists in the
        # steps/expected), so pin COVMAP hashes for every covered story-AC — a
        # confirmed-map correction stales exactly the journey TCs of that AC.
        if fm.get("kind") == "uat":
            for ref in fm.get("covers") or []:
                tgt, frag = resolve_ref(ref)
                if tgt in concepts and frag and \
                        concepts[tgt][0].get("type") == "User Story":
                    # always pin, even if the AC has no mapping YET — a
                    # confirmed_map correction adding one later must stale
                    # this TC too (symmetry with the has-a-mapping case).
                    pins[f"{tgt}#COVMAP:{frag}"] = covmap_hash(
                        concepts[tgt][0], frag)
        manifest["bindings"][sc] = {"tc": rel, "status": fm.get("status", "active"),
                                    "fragment_pins": pins}
        manifest["tc_hashes"][rel] = sha256(p.read_bytes())
        m = re.match(r".*-(\d+)$", fm["id"])
        if m:
            key = "sit:" + rel.split("/")[2] if rel.startswith("testcases/sit/") else "uat"
            manifest["counters"][key] = max(manifest["counters"].get(key, 0), int(m.group(1)))
        sealed += 1
    manifest["id_config_frozen"] = True
    manifest.setdefault("id_format", load_config()["ids"]["tc_format"])
    save_manifest(manifest)
    append_log(f"**Generation (agent)**: sealed {sealed} test case(s) into manifest")
    print(f"sealed {sealed} test case(s)")
    agent_commit(f"seal: {sealed} test case(s)")

def cmd_cascade():
    manifest = load_manifest()
    concepts = {rel: (fm, body, p) for rel, fm, body, p in all_concepts()}
    flagged = []
    verdicts = {}   # rel -> "needs-review" | "stale"
    # stories: aligned but source pins differ -> needs-review (spec §12.2)
    for rel, (fm, body, p) in concepts.items():
        if fm and fm.get("type") == "User Story" and fm.get("status") == "aligned":
            for src, pin in (fm.get("source_pins") or {}).items():
                cur = manifest["sources"].get(src, {})
                cur_hash = cur.get("content_hash") or cur.get("image_hash")
                if cur_hash and cur_hash != pin:
                    fm["status"] = "needs-review"
                    fm.setdefault("review_because", []).append(
                        {"cause": src, "at": now_iso()})
                    write_concept(p, fm, body)
                    verdicts[rel] = "needs-review"
                    flagged.append(f"story {fm['id']} -> needs-review ({src} changed)")
                    break
    # TCs: pinned fragment hashes differ -> stale
    for sc, binding in manifest["bindings"].items():
        tcrel = binding["tc"]
        if tcrel not in concepts:
            continue
        fm, body, p = concepts[tcrel]
        causes = []
        for fragref, pin in binding.get("fragment_pins", {}).items():
            tgt, frag = fragref.split("#")
            if frag.startswith("COVMAP:"):
                if tgt in concepts:
                    cur = covmap_hash(concepts[tgt][0], frag[len("COVMAP:"):])
                    if cur != pin:
                        causes.append(fragref)
                continue
            cur = fragment_hash_map(concepts[tgt][0]) if tgt in concepts else {}
            if cur.get(frag) and cur[frag] != pin:
                causes.append(fragref)
        if causes and fm.get("status") == "active":
            fm["status"] = "stale"
            fm["stale_because"] = [{"cause": c, "at": now_iso()} for c in causes]
            write_concept(p, fm, body)
            binding["status"] = "stale"
            verdicts[tcrel] = "stale"
            if tcrel in manifest["concepts"]:
                manifest["concepts"][tcrel]["status"] = "stale"
            flagged.append(f"TC {fm['id']} -> stale ({', '.join(causes)})")
    save_manifest(manifest)
    for f in flagged:
        print("FLAG", f)
    if flagged:
        append_log(f"**Cascade (agent)**: {len(flagged)} item(s) flagged")
    print(f"cascade: {len(flagged)} item(s) flagged (flags only — spec P8: nothing regenerates)")
    if verdicts:
        from wiki_graph import build_model, emit, scope_from_verdicts
        model = build_model(concepts, manifest)
        seed = next(iter(verdicts))
        emit("cascade", scope_from_verdicts(model, seed, verdicts),
             make_html="--graph" in sys.argv)
    agent_commit(f"cascade: {len(flagged)} flag(s)")

# ---------------------------------------------------------------- export

# ---------------------------------------------------------------- main

def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    cmd = sys.argv[1]
    args = [a for a in sys.argv[2:]
            if a not in ("--no-commit", "--allow-lint-errors")]
    if cmd == "ingest-prd":
        cmd_ingest_prd()
    elif cmd == "ingest-figma":
        cmd_ingest_figma()
    elif cmd == "ingest-decks":
        cmd_ingest_decks()
    elif cmd == "ingest-reference":
        cmd_ingest_reference()
    elif cmd == "triage":
        from wiki_triage import cmd_triage
        cmd_triage(args)
    elif cmd == "index":
        cmd_index()
    elif cmd == "manifest":
        cmd_manifest()
    elif cmd == "lint":
        cmd_lint()
    elif cmd == "status":
        cmd_status()
    elif cmd == "card":
        cmd_card(args)
    elif cmd == "assert":
        cmd_assert(args)
    elif cmd == "session":
        cmd_session(args)
    elif cmd == "gate":
        cmd_gate(args)
    elif cmd == "seal":
        cmd_seal()
    elif cmd == "cascade":
        cmd_cascade()
    elif cmd == "export":
        from wiki_export import cmd_export
        cmd_export(args)
    elif cmd == "rtm":
        from wiki_rtm import cmd_rtm
        cmd_rtm(args)
        agent_commit("rtm: build")
    elif cmd == "suite":
        from wiki_suite import cmd_suite
        cmd_suite(args)
    elif cmd == "inventory":
        from wiki_suite import cmd_suite
        cmd_suite(["compile"] + args)
    elif cmd in ("retire", "void-ac", "unretire", "release", "revert"):
        from wiki_lifecycle import cmd_lifecycle
        cmd_lifecycle(cmd, args)
    elif cmd in ("approve-cr", "reject-cr", "diff"):
        from wiki_change import cmd_change
        cmd_change(cmd, args)
    elif cmd == "migrate-ids":
        from wiki_migrate import cmd_migrate_ids
        cmd_migrate_ids(args)
    elif cmd == "migrate-provenance":
        from wiki_provenance import cmd_migrate_provenance
        cmd_migrate_provenance(args)
    elif cmd == "dashboard":
        from wiki_dashboard import cmd_dashboard
        cmd_dashboard()
    elif cmd == "coverage":
        from wiki_coverage import cmd_coverage
        cmd_coverage(args)
    elif cmd == "testmodel":
        from wiki_testmodel import cmd_testmodel
        cmd_testmodel(args)
    elif cmd == "impact":
        from wiki_impact import cmd_impact
        cmd_impact(args)
    elif cmd == "next":
        from wiki_next import cmd_next
        cmd_next(args)
    elif cmd == "app":
        try:
            sys.path.insert(0, str(ROOT / "tools/app"))
            from server import serve
        except ImportError as e:
            sys.exit(f"app: FastAPI is not installed ({e}).\n"
                     f"  Install it with: py -m pip install -r tools/app/requirements.txt")
        port = int(arg_after(args, "--port")) if "--port" in args else 8765
        serve(port=port, open_browser="--no-open" not in args)
    else:
        print(__doc__)
        sys.exit(2)

if __name__ == "__main__":
    main()
