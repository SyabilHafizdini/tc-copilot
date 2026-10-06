#!/usr/bin/env python3
"""One worksheet per run (run: py tools/test_wiki_suite_runs.py)."""
import re
import sys
import tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import wiki_suite


def _tc(i, run=None, module="/modules/m.md"):
    fm = {"id": f"1.1-AC01-{i:02d}", "module": module}
    if run:
        fm["run"] = run
    return (f"testcases/sit/m/1.1-AC01-{i:02d}", fm, "")


def test_run_groups_keep_first_appearance_order():
    tcs = [_tc(1, "Main flow"), _tc(2, "Main flow"), _tc(3, "Variant flow"), _tc(4, "Standalone checks")]
    groups = wiki_suite.run_groups(tcs)
    assert [g[0] for g in groups] == ["Main flow", "Variant flow", "Standalone checks"]
    assert [len(g[1]) for g in groups] == [2, 1, 1]


def test_test_cases_without_a_run_stay_on_one_sheet():
    groups = wiki_suite.run_groups([_tc(1), _tc(2)])
    assert len(groups) == 1 and len(groups[0][1]) == 2


def test_sheet_names_fit_excel_and_are_unique():
    names = [wiki_suite.run_sheet_name(n, i) for i, n in enumerate(
        ["Main flow", "Variant flow (returning tenant, alternative options)", "Standalone checks"])]
    assert all(len(n) <= 31 for n in names), names
    assert len(set(names)) == 3, names


def test_sheet_name_replaces_characters_excel_forbids():
    n = wiki_suite.run_sheet_name("Pass/validity: [edge] *a*? \\ b 'q'", 0)
    assert not re.search(r"[\[\]:*?/\\']", n), n
    assert len(n) <= 31 and n.startswith("C-TC-1 (") and n.endswith(")"), n


def test_long_names_that_truncate_alike_stay_unique_and_fit():
    long = "Variant flow for a very long tenant description"
    names = [wiki_suite.run_sheet_name(long, i) for i in (0, 1, 9, 10, 98)]
    assert all(len(n) <= 31 for n in names), names
    assert len({n.lower() for n in names}) == len(names), names
    # identical titles differ by the index prefix alone
    a, b = wiki_suite.run_sheet_name(long, 0), wiki_suite.run_sheet_name(long, 1)
    assert a.split(" (")[1] == b.split(" (")[1] and a != b


def _concepts(mod_title="Alpha Module", mod_desc="PRD section §1"):
    """Synthetic wiki: one module and one story, nothing from a project."""
    return {
        "modules/m": ({"type": "Module", "title": mod_title,
                       "description": mod_desc}, "", None),
        "stories/US-X": ({"type": "User Story", "title": "Story X",
                          "description": "§1.1"}, "", None),
    }


def _full_tc(i, run):
    fm = {"id": f"1.1-AC01-{i:02d}", "module": "/modules/m.md",
          "covers": ["/stories/US-X.md#AC1"], "section": "S", "confidence": "High", "title": "t"}
    if run:
        fm["run"] = run
    body = "# Steps\n1. Do.\n\n# Expected Results\n1. Done.\n"
    return (f"testcases/sit/m/{fm['id']}", fm, body)


def _render(runs, concepts=None):
    """Render a synthetic selection; `runs` holds one run name (or None) per TC."""
    from openpyxl import load_workbook
    tcs = [_full_tc(k + 1, r) for k, r in enumerate(runs)]
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "w.xlsx"
        wiki_suite.render_xlsx(tcs, "t", out, concepts=concepts or _concepts(),
                               manifest={})
        return load_workbook(out)


def _assert_formulas_resolve(wb):
    names = set(wb.sheetnames)
    refs = 0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for ref in re.findall(r"'([^']+)'!", c.value):
                        assert ref in names, (ws.title, c.coordinate, ref)
                        refs += 1
    return refs


def test_two_runs_give_two_sheets_one_dashboard_each_and_valid_references():
    runs = ["Main flow"] * 3 + ["Variant flow (alt)"] * 2
    wb = _render(runs)
    tc_sheets = [n for n in wb.sheetnames if n.startswith("C-TC")]
    assert tc_sheets == [wiki_suite.run_sheet_name("Main flow", 0),
                         wiki_suite.run_sheet_name("Variant flow (alt)", 1)], tc_sheets
    for sheet, n in zip(tc_sheets, (3, 2)):
        ids = [wb[sheet].cell(row=r, column=1).value for r in range(1, 40)]
        assert sum(1 for v in ids if str(v).startswith("TC-")) == n, (sheet, ids)
    stats = wb["Test Statistics"]
    titles = [stats.cell(row=11, column=2 + 3 * i).value for i in range(3)]
    assert titles == ["Main flow", "Variant flow (alt)", None], titles
    vlookups = [c.value for row in stats.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=VLOOKUP")]
    assert len(vlookups) == 5
    for sheet in tc_sheets:
        assert any(f"'{sheet}'!$A:$N" in v for v in vlookups), sheet
    assert _assert_formulas_resolve(wb) >= 5
    toc = [wb["A - Table of Contents"].cell(row=r, column=1).value for r in range(23, 30)]
    for sheet in tc_sheets:
        assert sheet in toc, toc


def test_mixed_selection_puts_runless_cases_on_the_module_sheet():
    wb = _render([None, None, "Main flow"])
    tc_sheets = [n for n in wb.sheetnames if n.startswith("C-TC")]
    module_sheet = wiki_suite._sheet_name(_concepts()["modules/m"][0], 0)
    assert tc_sheets == [module_sheet, wiki_suite.run_sheet_name("Main flow", 1)], tc_sheets
    assert _assert_formulas_resolve(wb) > 0


def test_no_run_keeps_the_module_sheet():
    wb = _render([None, None])
    module_sheet = wiki_suite._sheet_name(_concepts()["modules/m"][0], 0)
    assert [n for n in wb.sheetnames if n.startswith("C-TC")] == [module_sheet]


def test_a_sheet_name_collision_is_a_refusal_not_a_silent_duplicate():
    # The module sheet takes its number from the module's PRD section, a run
    # sheet from its index: module "§2" titled like run 1 gives the same name.
    concepts = _concepts(mod_title="Main flow", mod_desc="PRD section §2")
    try:
        _render([None, "Main flow"], concepts)
    except SystemExit as e:
        assert "share a name" in str(e), e
    else:
        raise AssertionError("duplicate sheet names were not refused")


def _chain_tc(i, run, chain):
    fm = {"id": f"1.1-AC01-{i:02d}", "module": "/modules/m.md",
          "covers": ["/stories/US-X.md#AC1"], "section": "S", "confidence": "High", "title": "t"}
    if run:
        fm["run"] = run
    pre = "1. Shared setup.\n2. " + chain + "\n"
    body = f"# Preconditions\n\n{pre}\n# Steps\n1. Do.\n\n# Expected Results\n1. Done.\n"
    return (f"testcases/sit/m/{fm['id']}", fm, body)


def _chain_rows(tcs):
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "w.xlsx"
        wiki_suite.render_xlsx(tcs, "t", out, concepts=_concepts(), manifest={})
        wb = load_workbook(out)
    ws = wb[[n for n in wb.sheetnames if n.startswith("C-TC")][0]]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    steps = [r[2] for r in rows if str(r[0]).startswith("TC-")]
    block = [str(r[0]) for r in rows if r[0] and re.match(r"^\d+\. ", str(r[0]))]
    return steps, block


def test_all_start_sheet_keeps_start_of_run_in_every_test_steps_cell():
    line = "Start of run: no test case to continue from."
    steps, block = _chain_rows([_chain_tc(i, "Standalone checks", line) for i in (1, 2, 3)])
    assert len(steps) == 3 and all(str(s).startswith("Start of run:") for s in steps), steps
    assert block == ["1. Shared setup."], block


def test_one_case_sheet_keeps_its_continue_from_line_in_test_steps():
    line = "Continue from TC-1.1-AC01-09: Step 1 completed."
    steps, block = _chain_rows([_chain_tc(1, "Main flow", line)])
    assert len(steps) == 1 and str(steps[0]).startswith("Continue from TC-"), steps
    assert block == ["1. Shared setup."], block


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); print(f"[PASS] {name}")
    print("test_wiki_suite_runs OK")
