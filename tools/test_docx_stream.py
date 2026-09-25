#!/usr/bin/env python3
"""Plain-assert tests for the docx PRD stream extractor
(run: py tools/test_docx_stream.py). No pytest -- matches tools/smoke.py."""
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from wiki import extract_docx_stream

TMP = ROOT / "build/_docx_t"


def _doc(blocks):
    """blocks: list of ("h", level, text) | ("p", text) | ("tbl", [[cells]])
    -> path to a .docx built with python-docx."""
    from docx import Document
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)
    d = Document()
    for b in blocks:
        if b[0] == "h":
            d.add_heading(b[2], level=b[1])
        elif b[0] == "p":
            d.add_paragraph(b[1])
        else:
            rows = b[1]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    t.cell(ri, ci).text = cell
    p = TMP / "sample.docx"
    d.save(str(p))
    return p


def test_heading_levels_produce_synthesized_dotted_numbers():
    p = _doc([("h", 1, "Introduction"), ("p", "body a"),
              ("h", 2, "Purpose"), ("p", "body b"),
              ("h", 2, "Scope"), ("p", "body c"),
              ("h", 1, "Example File Validator"), ("p", "body d")])
    heads = [v for k, v in extract_docx_stream(p) if k == "heading"]
    assert [h[0] for h in heads] == ["1", "1.1", "1.2", "2"], heads
    assert [h[1] for h in heads] == ["Introduction", "Purpose", "Scope",
                                     "Example File Validator"], heads


def test_slug_is_derived_from_the_heading_path_not_the_number():
    p = _doc([("h", 1, "Example File Validator"), ("p", "x"),
              ("h", 2, "File Validator"), ("p", "y"),
              ("h", 3, "Configurations"), ("p", "z")])
    heads = [v for k, v in extract_docx_stream(p) if k == "heading"]
    assert heads[-1][2] == "example-file-validator--file-validator--configurations", heads


def test_repeated_titles_under_different_parents_do_not_collide():
    p = _doc([("h", 1, "File Validator"), ("p", "a"),
              ("h", 2, "Configurations"), ("p", "b"),
              ("h", 1, "File Summary"), ("p", "c"),
              ("h", 2, "Configurations"), ("p", "d")])
    slugs = [v[2] for k, v in extract_docx_stream(p) if k == "heading"]
    assert slugs[1] != slugs[3], slugs
    assert slugs[1] == "file-validator--configurations", slugs
    assert slugs[3] == "file-summary--configurations", slugs


def test_tables_interleave_in_document_order():
    p = _doc([("h", 1, "Intro"), ("p", "before"),
              ("tbl", [["A", "B"], ["1", "2"]]), ("p", "after")])
    kinds = [k for k, _ in extract_docx_stream(p)]
    texts = [v for k, v in extract_docx_stream(p)]
    ti = kinds.index("table")
    assert texts[ti - 1] == "before", texts
    assert texts[ti + 1] == "after", texts
    assert "| A | B |" in texts[ti], texts[ti]


def test_a_document_with_no_headings_yields_no_heading_items():
    p = _doc([("p", "just prose"), ("p", "more prose")])
    assert not [k for k, _ in extract_docx_stream(p) if k == "heading"]


def test_chunk_sections_uses_docx_heading_items():
    from wiki import chunk_sections
    p = _doc([("h", 1, "Introduction"), ("p", "intro body"),
              ("h", 2, "Purpose"), ("p", "purpose body")])
    secs = chunk_sections(extract_docx_stream(p))
    by_slug = {s["slug"]: s for s in secs}
    assert "introduction" in by_slug, by_slug.keys()
    assert "introduction--purpose" in by_slug, by_slug.keys()
    assert by_slug["introduction"]["num"] == "1"
    assert by_slug["introduction--purpose"]["title"] == "Purpose"
    assert by_slug["introduction--purpose"]["heading_path"] == [
        "1 Introduction", "1.1 Purpose"]
    assert by_slug["introduction--purpose"]["body"] == "purpose body"


def test_parse_prd_reads_a_docx():
    from wiki import parse_prd
    p = _doc([("h", 1, "Scope"), ("p", "the scope body")])
    secs = parse_prd(p)
    assert [s["slug"] for s in secs] == ["scope"], secs


def test_parse_prd_refuses_a_docx_with_no_headings():
    from wiki import parse_prd
    p = _doc([("p", "prose only"), ("p", "more prose")])
    try:
        parse_prd(p)
    except SystemExit as e:
        assert "no headings" in str(e).lower(), str(e)
    else:
        assert False, "expected SystemExit for a headingless docx"


def _prd_root(blocks, name="spec.docx"):
    """A throwaway repo root holding `blocks` as inputs/prd/v1/<name> and an
    empty manifest, so ingest-prd can be driven through the real CLI."""
    import json
    src = _doc(blocks)                     # rebuilds TMP
    dest = TMP / "inputs/prd/v1" / name
    dest.parent.mkdir(parents=True, exist_ok=True)
    src.replace(dest)
    (TMP / "manifest.json").write_text(
        json.dumps({"schema_version": 1, "adopted_prd_version": None,
                    "staged_prd_version": None, "id_config_frozen": False,
                    "counters": {}, "sources": {}, "concepts": {},
                    "bindings": {}, "tc_hashes": {}, "edges": []}),
        encoding="utf-8", newline="\n")
    return TMP


def _cli(root, *argv):
    import os
    import subprocess
    env = dict(os.environ, TC_ROOT_OVERRIDE=str(root))
    return subprocess.run(
        [sys.executable, str(ROOT / "tools" / "wiki.py"), *argv, "--no-commit"],
        cwd=ROOT, capture_output=True, text=True, env=env)


def test_ingest_refuses_two_sections_whose_heading_paths_collide():
    """`File Validator > Configurations` appearing twice yields one slug
    twice. Without a check the second write_concept overwrites the first, the
    manifest entry overwrites, and ingest exits 0 having lost a section --
    the silent merge the design says must refuse."""
    from wiki import chunk_sections, extract_docx_stream
    root = _prd_root([("h", 1, "File Validator"),
                      ("h", 2, "Configurations"), ("p", "first config body"),
                      ("h", 2, "Windows Task Scheduler"), ("p", "sched body"),
                      ("h", 2, "Configurations"), ("p", "second config body")])
    docx = root / "inputs/prd/v1/spec.docx"
    slugs = [s["slug"] for s in chunk_sections(extract_docx_stream(docx))]
    assert slugs.count("file-validator--configurations") == 2, slugs
    r = _cli(root, "ingest-prd")
    out = r.stdout + r.stderr
    assert r.returncode != 0, out
    assert "file-validator--configurations" in out, out
    assert "1.1 Configurations" in out, out
    assert "1.3 Configurations" in out, out
    prd_dir = root / "sources/prd"
    assert not prd_dir.exists() or not list(prd_dir.glob("*.md")), \
        sorted(prd_dir.rglob("*"))


def test_a_slug_collision_writes_nothing_at_all():
    """The refusal must fire BEFORE any concept lands, not partway through."""
    import json
    root = _prd_root([("h", 1, "Intro"), ("p", "intro body"),
                      ("h", 1, "File Validator"),
                      ("h", 2, "Configurations"), ("p", "a"),
                      ("h", 2, "Configurations"), ("p", "b")])
    assert _cli(root, "ingest-prd").returncode != 0
    prd_dir = root / "sources/prd"
    assert not prd_dir.exists() or not list(prd_dir.glob("*.md")), \
        sorted(prd_dir.rglob("*"))
    m = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    assert m["sources"] == {}, m["sources"]
    assert m["adopted_prd_version"] is None, m


def test_write_prd_sections_refuses_a_numeric_collision_too():
    """The check lives in write_prd_sections, so it guards pdf/md (numeric
    slugs) as well as docx (heading-path slugs).

    wiki.ROOT is redirected at the module for the duration: this calls the
    writer directly, and a regression that let it write would otherwise drop
    a concept into the REPO's own sources/prd/."""
    import wiki
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)
    secs = [{"num": "3.1", "slug": "3-1", "title": "Ordering",
             "heading_path": ["3 Orders", "3.1 Ordering"], "body": "a",
             "content_hash": "sha256:a"},
            {"num": "3.1", "slug": "3-1", "title": "Ordering (again)",
             "heading_path": ["3 Orders", "3.1 Ordering (again)"], "body": "b",
             "content_hash": "sha256:b"}]
    manifest = {"sources": {}}
    real_root = wiki.ROOT
    wiki.ROOT = TMP
    try:
        wiki.write_prd_sections(secs, 1, "/inputs/prd/v1/x.pdf", manifest)
    except SystemExit as e:
        assert "3-1" in str(e), str(e)
        assert "3.1 Ordering" in str(e), str(e)
    else:
        assert False, "expected SystemExit for a duplicate PRD slug"
    finally:
        wiki.ROOT = real_root
    assert manifest["sources"] == {}, manifest
    assert not (TMP / "sources").exists(), sorted(TMP.rglob("*"))


def test_distinct_heading_paths_still_ingest():
    root = _prd_root([("h", 1, "File Validator"),
                      ("h", 2, "Configurations"), ("p", "a"),
                      ("h", 1, "File Summary"),
                      ("h", 2, "Configurations"), ("p", "b")])
    r = _cli(root, "ingest-prd")
    out = r.stdout + r.stderr
    assert r.returncode == 0, out
    assert (root / "sources/prd/file-validator--configurations.md").exists(), \
        sorted((root / "sources/prd").glob("*"))
    assert (root / "sources/prd/file-summary--configurations.md").exists(), \
        sorted((root / "sources/prd").glob("*"))


def test_md_and_numeric_chunking_is_unchanged():
    from wiki import chunk_sections
    secs = chunk_sections([("text", "3.1 Ordering"), ("text", "order body")])
    assert [s["slug"] for s in secs] == ["3-1"], secs


if __name__ == "__main__":
    try:
        for name, fn in sorted(globals().items()):
            if name.startswith("test_") and callable(fn):
                fn()
                print(f"[PASS] {name}")
        print("test_docx_stream OK")
    finally:
        if TMP.exists():
            shutil.rmtree(TMP)
