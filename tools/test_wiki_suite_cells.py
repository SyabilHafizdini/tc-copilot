#!/usr/bin/env python3
"""The workbook's row cells as importable functions
(run: py tools/test_wiki_suite_cells.py). The operator app's grid calls the
same functions, so these tests pin that the sheet writes exactly what they
return."""
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import render_sit
import render_uat
import wiki
import wiki_suite
from wiki_coverage import ELEMENT_BLOCK_LEAD

COMMON = "1. Shared setup.\n2. Shared data."


def _concepts():
    return {
        "modules/m": ({"type": "Module", "title": "Alpha Module",
                       "description": "PRD section §1"}, "", None),
        "stories/US-X": ({"type": "User Story", "title": "Story X",
                          "description": "§1.1", "acceptance_criteria": [
                              {"id": "AC1", "text": "Log in. Given a user, when "
                               "they log in, then the page opens."}]}, "", None),
    }


def _tc(i, run="Main flow", chain="Start of run: no test case to continue from.",
        extra=None, data="**User** = tenant", expected="1. Done."):
    fm = {"id": f"1.1-AC01-{i:02d}", "module": "/modules/m.md", "title": "t",
          "covers": ["/stories/US-X.md#AC1"], "section": "S", "run": run,
          "scenario_id": f"SC-{i}", "confidence": "Medium",
          "confidence_parts": {
              "scenario": {"level": "High", "remark": "Source: AC1."},
              "steps": {"level": "Medium", "remark": "Inferred: x. Verify: y."},
              "data": {"level": "High", "remark": ""},
              "expected": {"level": "High", "remark": "Source: AC1."}}}
    pre = f"{COMMON}\n3. {chain}" + (f"\n4. {extra}" if extra else "")
    body = (f"# Objective\n\nO.\n\n# Preconditions\n\n{pre}\n\n# Test Data\n\n{data}\n\n"
            f"# Steps\n\n1. Do **it**.\n\n# Expected Results\n\n{expected}\n\n"
            f"# Postconditions\n\nNone.\n")
    return (f"testcases/sit/m/{fm['id']}", fm, body)


def test_sheet_common_is_the_shared_lines_without_the_chain_line():
    tcs = [_tc(1), _tc(2, chain="Continue from TC-1.1-AC01-01: Home.", extra="Own line.")]
    assert wiki_suite.sheet_common(tcs) == ["Shared setup.", "Shared data."]
    # an all-start sheet repeats the chain line in every case: still not shared
    assert wiki_suite.sheet_common([_tc(1), _tc(2)]) == ["Shared setup.", "Shared data."]
    assert wiki_suite.sheet_common([]) == []


def test_tc_cells_lifts_the_chain_into_steps_and_the_extra_into_data():
    tc = _tc(2, chain="Continue from TC-1.1-AC01-01: Home.", extra="Own line.")
    c = wiki_suite.tc_cells(tc[1], tc[2], _concepts(), ["Shared setup.", "Shared data."])
    assert c["id"] == "TC-1.1-AC01-02"
    assert c["chain"] == ["Continue from TC-1.1-AC01-01: Home."]
    assert c["extra"] == ["Own line."]
    assert c["steps"] == "Continue from TC-1.1-AC01-01: Home.\n\n1. Do **it**."
    assert c["data"] == "**User** = tenant\n\nOwn line."
    assert c["scenario"].startswith("**Scenario:** Log in.\n\n**Given:** A user.")
    assert c["confidence"] == "Medium"
    assert c["remarks"].split("\n")[1] == "**Test Steps**: Medium - Inferred: x. Verify: y."


def test_tc_cells_blanks_a_dash_and_tags_the_question():
    tc = _tc(1, data="-")
    c = wiki_suite.tc_cells(tc[1], tc[2], _concepts(), ["Shared setup.", "Shared data."],
                            {"SC-1#steps": "Q-07"})
    assert c["data"] == "" and c["extra"] == []
    assert "**Test Steps**: Medium [Q-07] - " in c["remarks"]


def test_split_expected_separates_the_element_block():
    block = f"3. {ELEMENT_BLOCK_LEAD}\n    a. Button : **Save**\n    b. Label : **Name**"
    own, got = wiki_suite.split_expected(f"1. Saved.\n2. Shown.\n{block}")
    assert own == "1. Saved.\n2. Shown." and got == block
    assert wiki_suite.split_expected("1. Saved.") == ("1. Saved.", "")
    assert wiki_suite.split_expected("") == ("", "")
    # a human item that merely quotes the sentence is not the block
    quoted = f"1. The text '{ELEMENT_BLOCK_LEAD}' is shown."
    assert wiki_suite.split_expected(quoted) == (quoted, "")


def test_both_renderers_write_the_lead_split_expected_looks_for():
    saved = (render_sit.element_verification_block_multi,
             render_uat.element_verification_block)
    render_sit.element_verification_block_multi = lambda fm, ids: "    a. Button : **Go**"
    render_uat.element_verification_block = lambda fm, ac: "    a. Button : **Go**"
    try:
        for text in (render_sit.with_element_block({}, "1. A.\n2. B.", ["AC1"]),
                     render_uat.with_element_block({}, "1. A.\n2. B.", "AC1")):
            assert text == ("1. A.\n2. B.\n3. The following elements are displayed "
                            "and labelled correctly:\n    a. Button : **Go**"), text
            assert wiki_suite.split_expected(text)[0] == "1. A.\n2. B."
    finally:
        (render_sit.element_verification_block_multi,
         render_uat.element_verification_block) = saved


def _sheet_rows(wb):
    """{display id: [A..G as plain text]} over every C-TC sheet."""
    out = {}
    for ws in wb:
        if not ws.title.startswith("C-TC"):
            continue
        for row in ws.iter_rows(min_row=wiki_suite.TC_HEADER_ROW + 1):
            if str(row[0].value or "").startswith("TC-"):
                out[str(row[0].value)] = [str(c.value or "") for c in row[:7]]
    return out


def _assert_sheet_equals_cells(tcs, concepts, manifest, summaries=None):
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "w.xlsx"
        wiki_suite.render_xlsx(tcs, "t", out, concepts=concepts, manifest=manifest,
                               doubt_summaries=summaries)
        got = _sheet_rows(load_workbook(out))
    qmap = wiki_suite._question_map(summaries)
    n = 0
    for _name, _mod, subs, _dash in wiki_suite.plan_sheets(tcs, concepts):
        sheet = [tc for _s, _f, tcl in subs for tc in tcl]
        common = wiki_suite.sheet_common(sheet)
        for _rel, fm, body in sheet:
            c = wiki_suite.tc_cells(fm, body, concepts, common, qmap)
            want = [c["id"]] + [wiki_suite._plain(c[k]) for k in
                                ("scenario", "steps", "data", "expected")] + \
                   [c["confidence"], wiki_suite._plain(c["remarks"])]
            assert got[c["id"]] == want, (c["id"], got[c["id"]], want)
            n += 1
    assert n == len(got) == len(tcs), (n, len(got), len(tcs))


def test_the_sheet_writes_exactly_what_tc_cells_returns():
    tcs = [_tc(1), _tc(2, chain="Continue from TC-1.1-AC01-01: Home.", extra="Own line."),
           _tc(3, run="Variant flow", data="-")]
    _assert_sheet_equals_cells(tcs, _concepts(), {})


def test_run_order_holds_for_a_story_number_that_is_not_numeric():
    """Rows and sheets follow the spec's run order (`order`) whatever the id
    looks like: an id with a lettered story number must not fall back to id
    order, which would put a later run's first case among the first run's."""
    def fm(tid, order, run):
        return {"id": tid, "module": "/modules/m.md", "order": order, "run": run,
                "covers": ["/stories/US-X.md#HS-01"]}
    tcs = [fm("ABC-001-AC01-01", 1, "Main"), fm("ABC-001-AC05-02", 2, "Main"),
           fm("ABC-001-AC05-01", 3, "Main"), fm("ABC-001-AC01-02", 4, "Variant")]
    got = sorted(reversed(tcs), key=wiki_suite.tc_sort_key)
    assert [t["id"] for t in got] == [t["id"] for t in tcs], [t["id"] for t in got]
    runs = [name for name, _t in wiki_suite.run_groups([(None, t, "") for t in got])]
    assert runs == ["Main", "Variant"], runs


def test_plan_sheets_is_one_sheet_per_run_else_one_per_module():
    tcs = [_tc(1), _tc(2), _tc(3, run="Variant flow")]
    names = [s[0] for s in wiki_suite.plan_sheets(tcs, _concepts())]
    assert names == [wiki_suite.run_sheet_name("Main flow", 0),
                     wiki_suite.run_sheet_name("Variant flow", 1)], names
    runless = [(r, {k: v for k, v in fm.items() if k != "run"}, b) for r, fm, b in tcs]
    sheets = wiki_suite.plan_sheets(runless, _concepts())
    assert [s[0] for s in sheets] == \
        [wiki_suite._sheet_name(_concepts()["modules/m"][0], 0)]
    assert sheets[0][3] is None, "a module sheet has no run dashboard title"


def test_the_real_export_writes_exactly_what_tc_cells_returns():
    import wiki_doubts
    concepts, manifest = wiki.load_all()
    summaries = [wiki_doubts.story_summary(s)
                 for s in wiki_doubts.stories_with_doubts()]
    done = 0
    for kind in ("sit", "uat"):
        selected, _retired, _stale = wiki_suite.select({"kind": kind}, concepts)
        if selected:
            _assert_sheet_equals_cells(selected, concepts, manifest, summaries)
            done += 1
    if not done:
        print("[SKIP] test_the_real_export_writes_exactly_what_tc_cells_returns: "
              "bundle has no active test cases")


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"[PASS] {name}")
    print("test_wiki_suite_cells OK")
