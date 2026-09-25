#!/usr/bin/env python3
"""Plain-assert tests for the two-delivery export: workbook marks, the Change
Log sheet, the draft snapshot and the final-export gate
(run: py tools/test_wiki_export.py). No pytest - matches tools/smoke.py.

Everything writes under a throwaway directory; the wiki is never read."""
import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import rubric_judge as rj  # noqa: E402
from wiki_suite import change_rows, render_xlsx  # noqa: E402

BODY = ("# Objective\n\nSee the page.\n\n# Preconditions\n\n1. Signed in.\n\n"
        "# Test Data\n\n**Amount** = 100\n\n# Steps\n\n1. Open the **Page**.\n\n"
        "# Expected Results\n\n1. The **Page** is shown.\n\n# Postconditions\n\n"
        "-\n\n# Traceability\n\n- AC01\n")
STORY = {"type": "User Story", "id": "US-9", "title": "Nine",
         "description": "PRD §5.1", "module": "/modules/m.md",
         "acceptance_criteria": [{"id": "AC01", "text": "Shown."}]}
MODULE = {"type": "Module", "id": "m", "title": "Mod", "description": "§5"}


def _tc(tid, title="Open the page"):
    fm = {"type": "Test Case", "id": tid, "title": title, "kind": "sit",
          "status": "active", "technique": "UC", "priority": "P1",
          "module": "/modules/m.md", "covers": ["/stories/US-9.md#AC01"],
          "coverage_items": ["SC-01"]}
    return (f"testcases/sit/US-9/{tid}", fm, BODY)


CONCEPTS = {"stories/US-9": (STORY, ""), "modules/m": (MODULE, "")}


def _tmpdir():
    return Path(tempfile.mkdtemp(prefix="wx-"))


def _sheet(path, name):
    from openpyxl import load_workbook
    return load_workbook(path)[name]


def test_draft_workbook_is_marked_in_the_header_and_change_log():
    d = _tmpdir()
    try:
        out = d / "d.xlsx"
        render_xlsx([_tc("9-AC01-01")], "Nine", out, concepts=CONCEPTS,
                    manifest={}, draft=True)
        toc = _sheet(out, "A - Table of Contents")
        assert toc["B1"].value.startswith("DRAFT (ungraded, round 0) - "), toc["B1"].value
        assert toc["C8"].value == "DRAFT r0", toc["C8"].value
        assert "not yet graded" in toc["C18"].value, toc["C18"].value
        log = _sheet(out, "Change Log")
        assert log["C2"].value == "Draft round 0 - no changes yet.", log["C2"].value
    finally:
        shutil.rmtree(d)


def test_final_workbook_carries_the_grade_and_the_change_rows():
    d = _tmpdir()
    try:
        out = d / "f.xlsx"
        changes = {"to": {"round": 2, "score": 93.44},
                   "rounds": [{"round": 1, "score": 70.77}, {"round": 2, "score": 93.44}],
                   "delta": 22.67, "regression_restored": False, "rubric_version": "v1",
                   "changes": [
                       {"tc": "9-AC01-01", "kind": "changed", "field": "Test Data",
                        "old": "a", "new": "b", "closes": "G3.1"},
                       {"tc": "9-AC01-01", "kind": "changed", "field": "Expected Results",
                        "old": "a", "new": "b", "closes": "G3.1"},
                       {"tc": "9-AC01-02", "kind": "added", "title": "Rejects zero",
                        "closes": "G1.19, G2.2"},
                       {"tc": "9-AC01-03", "kind": "removed", "reason": "superseded"}]}
        render_xlsx([_tc("9-AC01-01"), _tc("9-AC01-02", "Rejects zero")], "Nine", out,
                    concepts=CONCEPTS, manifest={},
                    grade={"round": 2, "score": 93.44, "threshold": 70, "version": "v1"},
                    changes=changes)
        toc = _sheet(out, "A - Table of Contents")
        assert toc["C8"].value == "r2 · rubric 93.44", toc["C8"].value
        assert "Graded round 2: score 93.44 (threshold 70), rubric v1" == toc["C18"].value, toc["C18"].value
        log = _sheet(out, "Change Log")
        rows = [(log.cell(r, 1).value, log.cell(r, 2).value, log.cell(r, 3).value)
                for r in range(2, 6)]
        # openpyxl writes an empty-string cell value as a blank cell (no <is>
        # element at all - see openpyxl.cell._writer), so an empty Dashboard/
        # Test Case ID round-trips as None, not "". These two rows assert
        # that real, observed behaviour rather than the pre-round-trip value.
        assert rows[0] == (None, None, "Graded round 2: 93.44 (round 1: 70.77, +22.67), rubric v1"), rows[0]
        assert rows[1] == ("Nine", "TC-9-AC01-01", "Changed Test Data, Expected Results (closes G3.1)"), rows[1]
        assert rows[2] == ("Nine", "TC-9-AC01-02", "Added - Rejects zero (closes G1.19, G2.2)"), rows[2]
        assert rows[3] == (None, "TC-9-AC01-03", "Removed - superseded"), rows[3]
    finally:
        shutil.rmtree(d)


def test_final_workbook_without_changes_writes_the_courtesy_row():
    d = _tmpdir()
    try:
        out = d / "f.xlsx"
        render_xlsx([_tc("9-AC01-01")], "Nine", out, concepts=CONCEPTS, manifest={},
                    grade={"round": 2, "score": 80.0, "threshold": 70, "version": "v1"})
        log = _sheet(out, "Change Log")
        assert log["C2"].value == ("No current change log - run eval_rubric --diff on "
                                   "this sealed set."), log["C2"].value
    finally:
        shutil.rmtree(d)


def test_change_rows_reports_a_restored_round_one():
    rows = change_rows({"to": {"round": 1, "score": 80.0},
                        "rounds": [{"round": 1, "score": 80.0}, {"round": 2, "score": 75.0}],
                        "delta": -5.0, "regression_restored": True,
                        "rubric_version": "v1", "changes": []}, {})
    assert rows[0][2] == ("Graded round 1: 80.0 (round 2 regressed to 75.0, -5.0; "
                          "round 1 kept), rubric v1"), rows[0]


def test_change_rows_maps_a_br_id_through_display_id_and_dash_of():
    # A BR- id must delegate to the canonical _display_id convention (which
    # special-cases BR- alongside TC-), not the inline TC-prefixing that used
    # to turn it into "TC-BR-..." and miss the dash_of lookup.
    rows = change_rows({"to": {"round": 1, "score": 90.0}, "rounds": [],
                        "delta": None, "regression_restored": False,
                        "rubric_version": "v1",
                        "changes": [{"tc": "BR-1", "kind": "changed",
                                    "field": "Test Data"}]},
                       {"BR-1": "Nine"})
    assert rows[1] == ("Nine", "BR-1", "Changed Test Data"), rows[1]


def test_change_rows_refuses_an_unknown_change_kind():
    try:
        change_rows({"to": {"round": 1, "score": 90.0}, "rounds": [],
                     "delta": None, "regression_restored": False,
                     "rubric_version": "v1",
                     "changes": [{"tc": "9-AC01-01", "kind": "renamed"}]}, {})
    except ValueError as e:
        assert "unknown change kind" in str(e), e
        assert "renamed" in str(e), e
        assert "9-AC01-01" in str(e), e
    else:
        raise AssertionError("an unrecognized change kind must raise ValueError")


# ------------------------------------------------ snapshot + gate
import wiki_export as wx  # noqa: E402

CFG = {"rubric": {"threshold": 70, "version": "v1"}}


def test_select_scope_tcs_keeps_active_cases_of_the_scope_in_sort_order():
    retired = dict(_tc("9-AC01-02")[1], status="retired")
    other = dict(_tc("8-AC01-01")[1], covers=["/stories/US-8.md#AC01"])
    concepts = dict(CONCEPTS)
    concepts["testcases/sit/US-9/9-AC01-02"] = (_tc("9-AC01-02")[1], BODY)
    concepts["testcases/sit/US-9/9-AC01-01"] = (_tc("9-AC01-01")[1], BODY)
    concepts["testcases/sit/US-9/9-AC01-03"] = (retired, BODY)
    concepts["testcases/sit/US-8/8-AC01-01"] = (other, BODY)
    # a flow's UAT case covering the story is not part of the story's set (I5)
    concepts["testcases/uat/f/UAT-9-01"] = (dict(_tc("UAT-9-01")[1], kind="uat"), BODY)
    rel, sfm, tcs = wx.select_scope_tcs(concepts, "story", "US-9")
    assert rel == "stories/US-9" and sfm["title"] == "Nine"
    assert [fm["id"] for _r, fm, _b in tcs] == ["9-AC01-01", "9-AC01-02"], tcs


def test_draft_snapshot_records_digest_and_every_case():
    d = _tmpdir()
    try:
        tcs = [_tc("9-AC01-01"), _tc("9-AC01-02")]
        manifest = {"tc_hashes": {"testcases/sit/US-9/9-AC01-01": "sha256:1",
                                  "testcases/sit/US-9/9-AC01-02": "sha256:2"}}
        snap = wx.write_draft_snapshot("US-9", "story", tcs, manifest,
                                       "build/inventory/sit/x.xlsx", "abc123", "v1",
                                       d / "US-9-draft.json")
        on_disk = json.loads((d / "US-9-draft.json").read_text(encoding="utf-8"))
        assert on_disk == snap
        assert set(snap["test_cases"]) == {"9-AC01-01", "9-AC01-02"}
        assert snap["test_cases"]["9-AC01-01"]["hash"] == "sha256:1"
        assert snap["seal_commit"] == "abc123" and snap["score"] is None
        from wiki_rubric import sealed_digest
        assert snap["sealed_digest"] == sealed_digest(manifest["tc_hashes"],
                                                      [r for r, _f, _b in tcs])
        # M9: the score field records the current score when one exists
        snap2 = wx.write_draft_snapshot("US-9", "story", tcs, manifest, "x.xlsx", "abc123",
                                        "v1", d / "US-9-draft.json", score=81.5)
        assert snap2["score"] == 81.5, snap2["score"]
    finally:
        shutil.rmtree(d)


def _scores(d, *rounds):
    for n, score, digest, partial in rounds:
        (d / f"US-9-r{n}-score.json").write_text(json.dumps(
            {"score": score, "partial": partial, "sealed_digest": digest,
             "rubric_version": "v1"}), encoding="utf-8")


def test_export_gate_refuses_ungraded_round_one_only_and_stale_scores():
    old = rj.RUBRIC_BUILD
    d = _tmpdir()
    try:
        rj.RUBRIC_BUILD = d
        cur, err = wx.export_gate("US-9", "story", "sha256:cur", CFG)
        assert cur is None and "has not been graded" in err and "--draft" in err, err
        _scores(d, (1, 80.0, "sha256:cur", False))
        cur, err = wx.export_gate("US-9", "story", "sha256:cur", CFG)
        assert cur is None and "improve round has not run" in err, err
        _scores(d, (1, 80.0, "sha256:cur", False), (2, 90.0, "sha256:old", False))
        cur, err = wx.export_gate("US-9", "story", "sha256:cur", CFG)
        assert cur == {"round": 1, "score": 80.0} and err is None, (cur, err)
        _scores(d, (1, 80.0, "sha256:old1", False), (2, 90.0, "sha256:old", False))
        cur, err = wx.export_gate("US-9", "story", "sha256:cur", CFG)
        assert cur is None and "different sealed set" in err and "--round 2" in err, err
        _scores(d, (2, 90.0, "sha256:cur", True))
        cur, err = wx.export_gate("US-9", "story", "sha256:cur", CFG)
        assert cur is None and "PARTIAL" in err, err
        _scores(d, (2, 60.0, "sha256:cur", False))
        cur, err = wx.export_gate("US-9", "story", "sha256:cur", CFG)
        assert cur is None and "below the threshold" in err, err
    finally:
        rj.RUBRIC_BUILD = old
        shutil.rmtree(d)


def test_final_export_ignores_a_stale_changes_file():
    """A changes file computed on an older sealed set must not decorate the
    final workbook with the wrong change log."""
    stale = {"to_digest": "sha256:older", "changes": []}
    assert wx.changes_for("sha256:cur", stale) is None
    fresh = {"to_digest": "sha256:cur", "changes": []}
    assert wx.changes_for("sha256:cur", fresh) == fresh


def test_suite_workbook_has_a_header_only_change_log():
    """M1: a suite workbook is not a scope delivery - no courtesy row."""
    d = _tmpdir()
    try:
        out = d / "s.xlsx"
        render_xlsx([_tc("9-AC01-01")], "Suite: s", out, concepts=CONCEPTS, manifest={},
                    change_log=False)
        log = _sheet(out, "Change Log")
        assert log.max_row == 1, log.max_row      # before any cell access creates row 2
        assert log["C1"].value == "Description of Change", log["C1"].value
    finally:
        shutil.rmtree(d)


def test_draft_refuses_a_graded_set():
    """A2: the current sealed set already has a round score -> not a first cut."""
    state = {"rounds": {1: {"digest": "sha256:d0"}, 2: {"digest": "sha256:d1"}}}
    assert wx.graded_round_for("sha256:d1", state) == 2
    assert wx.graded_round_for("sha256:d0", state) == 1
    assert wx.graded_round_for("sha256:d2", state) is None, "a changed, ungraded set is a new first cut"
    assert wx.graded_round_for("sha256:d0", {"rounds": {}}) is None
    both = {"rounds": {1: {"digest": "sha256:d0"}, 2: {"digest": "sha256:d0"}}}
    assert wx.graded_round_for("sha256:d0", both) == 2
    msg = wx.draft_refusal("US-9", "story", 2, "US-9-sit")
    assert msg.startswith("export refused: the sealed set of US-9 was graded in round 2; "
                          "a graded set is not a first cut."), msg
    assert ("Next: py tools/eval_rubric.py --story US-9 --diff, then "
            "py tools/wiki.py export --story US-9 --name US-9-sit") in msg, msg


def test_a_malformed_changes_file_is_a_refusal_naming_it():
    """M3: never a traceback."""
    d = _tmpdir()
    try:
        assert wx.load_changes(d / "missing.json", "--story", "US-9") is None
        (d / "ok.json").write_text('{"to_digest": "x"}', encoding="utf-8")
        assert wx.load_changes(d / "ok.json", "--story", "US-9") == {"to_digest": "x"}
        for name, content in (("bad.json", '{"to_'), ("list.json", "[]"), ("null.json", "null")):
            (d / name).write_text(content, encoding="utf-8")
            try:
                wx.load_changes(d / name, "--story", "US-9")
            except SystemExit as e:
                assert name in str(e) and "--story US-9 --diff" in str(e), e
            else:
                raise AssertionError(f"malformed changes file accepted: {name}")
    finally:
        shutil.rmtree(d)


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"[PASS] {name}")
    print("test_wiki_export OK")
