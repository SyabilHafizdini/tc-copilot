#!/usr/bin/env python3
"""Tester route: `wiki doubts observe` and observations on the doubts card
(run: py tools/test_wiki_doubts_observe.py). Whole-project cases run as
subprocesses of a scratch copy of the tools; the real repository is never
touched."""
import hashlib
import json
import shutil
import sys
import tempfile
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import wiki
import wiki_doubts
import wiki_suite
import doubts_scratch
import test_wiki_doubts as T

STORY, Q1, Q2, SID, BY = T.STORY, T.Q1, T.Q2, T.SID, T.BY
OBS = f"doubts/observations/{STORY}.yaml"
TEXT1 = "  Profile A  (ID-1234567).\nMailbox is  inbox B \n"
TEXT2 = "The page shows done and a receipt."


def _workbook(d, name="w.xlsx"):
    """The real AI Doubts sheet (wiki_suite code) of the scratch project."""
    s = wiki_doubts.story_summary(STORY, root=d)
    tcs = []
    for p in sorted((d / "testcases" / "fixture").glob("TC-*.md")):
        fm, body = wiki.read_concept(p)[:2]
        tcs.append((p.relative_to(d).with_suffix("").as_posix(), fm, body))
    out = d / "build" / name
    out.parent.mkdir(exist_ok=True)
    wiki_suite.render_xlsx(tcs, "t", out, concepts={}, manifest={},
                           doubt_summaries=[s])
    return out


def _fill(path, rows):
    """rows: {question id: (observation, observed by, date)}; extra rows
    (id, observation) appended for ids not on the sheet."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["AI Doubts"]
    head = {c.value: c.column for c in ws[1]}
    seen = set()
    for r in range(2, ws.max_row + 1):
        qid = ws.cell(row=r, column=head["Question ID"]).value
        if qid in rows:
            seen.add(qid)
            for col, v in zip(("Observation", "Observed by", "Date"), rows[qid]):
                ws.cell(row=r, column=head[col], value=v)
    for qid, vals in rows.items():
        if qid not in seen:
            r = ws.max_row + 1
            ws.cell(row=r, column=head["Question ID"], value=qid)
            for col, v in zip(("Observation", "Observed by", "Date"), vals):
                ws.cell(row=r, column=head[col], value=v)
    wb.save(path)


def _obs(d):
    return yaml.safe_load((d / OBS).read_text(encoding="utf-8"))


def _observe(d, wb, by="carol", *extra):
    return T._sw(d, "doubts", "observe", "--workbook", str(wb), "--by", by,
                 *extra)


def test_observe_round_trip_idempotent_and_skips():
    def go(d):
        wb = _workbook(d)
        _fill(wb, {Q1: (TEXT1, None, None),
                   Q2: (TEXT2, "alice", "2026-10-06"),
                   "-": ("ungrouped note", None, None),
                   "Q-NOPE-01": ("who is this", None, None)})
        digest = hashlib.sha256(wb.read_bytes()).hexdigest()
        before = T._commits(d)
        code, out, err = _observe(d, wb)
        assert code == 0, out + err
        doc = _obs(d)
        assert doc["story"] == f"/stories/{STORY}.md"
        by_q = {o["question"]: o for o in doc["observations"]}
        assert sorted(by_q) == [Q1, Q2], by_q
        assert by_q[Q1]["text"] == TEXT1, by_q[Q1]["text"]
        assert by_q[Q1]["by"] == "carol" and by_q[Q2]["by"] == "alice"
        assert by_q[Q2]["observed_on"] == "2026-10-06"
        assert by_q[Q1]["observed_on"] is None
        assert by_q[Q1]["workbook"] == "w.xlsx"
        assert by_q[Q1]["workbook_hash"] == digest
        assert by_q[Q1]["at"]
        assert "skipped" in out and "-" in out and "Q-NOPE-01" in out, out
        assert T._commits(d) == before + 1
        msg = T._git_out(d, "log", "-1", "--format=%B")
        assert msg.startswith(f"doubts observe: {STORY} 2 observation(s) "
                              f"by carol"), msg
        assert "Assertion-Event" not in msg and "Co-Authored" not in msg
        first = (d / OBS).read_bytes()
        # idempotent: nothing appended, nothing committed
        code, out, err = _observe(d, wb)
        assert code == 0 and "no new observation" in out, out + err
        assert (d / OBS).read_bytes() == first
        assert T._commits(d) == before + 1
        # a changed workbook (new hash) appends again, old entries untouched
        wb2 = d / "build" / "w2.xlsx"
        shutil.copy(wb, wb2)
        _fill(wb2, {Q2: ("another view", "bob", None)})
        code, out, err = _observe(d, wb2)
        assert code == 0, out + err
        got = _obs(d)["observations"]
        assert got[:2] == doc["observations"], got
        assert len(got) == 4 and got[-1]["text"] == "another view", got
        code, lint, err = T._sw(d, "lint")
        assert code == 0 and "0 error(s)" in lint, lint + err
    T._with_scratch(go)


def test_observe_refusals():
    d = T._mkroot()
    try:
        junk = d / "junk.xlsx"
        junk.write_text("not a zip", encoding="utf-8")
        from openpyxl import Workbook
        nosheet = d / "nosheet.xlsx"
        w = Workbook()
        w.active.title = "Other"
        w.save(nosheet)
        nohead = d / "nohead.xlsx"
        w = Workbook()
        w.active.title = "AI Doubts"
        w.active.append(["Question ID", "Question"])
        w.save(nohead)
        noid = d / "noid.xlsx"
        w = Workbook()
        w.active.title = "AI Doubts"
        w.active.append(["Question", "Observation"])
        w.save(noid)
        cases = [
            ([], "--workbook"),
            (["--workbook", str(nosheet)], "--by"),
            (["--by", "t"], "--workbook"),
            (["--workbook", str(d / "missing.xlsx"), "--by", "t"], "not found"),
            (["--workbook", str(junk), "--by", "t"], "not an xlsx"),
            (["--workbook", str(nosheet), "--by", "t"], "AI Doubts"),
            (["--workbook", str(nohead), "--by", "t"], "Observation"),
            (["--workbook", str(noid), "--by", "t"], "Question ID"),
        ]
        for args, why in cases:
            code, out = T._run(["observe", *args], d)
            assert code == 1, (args, code, out)
            assert out.startswith("doubts observe refused:"), out
            assert why in out, (why, out)
            assert not (d / "doubts" / "observations").exists()
    finally:
        T._cleanup(d)


def _write_obs(d, entries, story=STORY):
    p = d / "doubts" / "observations" / f"{story}.yaml"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(yaml.safe_dump({"story": f"/stories/{story}.md",
                                 "observations": entries},
                                sort_keys=False, allow_unicode=True),
                 encoding="utf-8")
    return p


def _entry(q, text, at="2026-10-06T10:00:00+08:00", by="carol"):
    return {"question": q, "text": text, "by": by, "at": at,
            "observed_on": None, "workbook": "w.xlsx", "workbook_hash": "ab" * 32}


def test_card_carries_the_observation_as_proposed():
    d = T._mkroot()
    try:
        s = wiki_doubts.story_summary(STORY, root=d)
        plain = wiki_doubts.build_card(s, "h")
        assert wiki_doubts.build_card(s, "h", observations=[]) == plain
        new = {"proposed_origin", "observed_by", "observation_at",
               "register_proposed"}
        assert all(not (new & set(e)) for e in plain["open_questions"])
        obs = [_entry(Q1, "old", at="2026-10-05T09:00:00+08:00"),
               _entry(Q1, TEXT1, by="zed"), _entry(Q2, TEXT2)]
        card = wiki_doubts.build_card(s, "h", observations=obs)
        e = {x["id"]: x for x in card["open_questions"]}
        assert e[Q1]["proposed"] == TEXT1
        assert e[Q1]["proposed_origin"] == "tester-observation"
        assert e[Q1]["observed_by"] == "zed"
        assert e[Q1]["observation_at"] == "2026-10-06T10:00:00+08:00"
        assert e[Q1]["register_proposed"] is None
        assert e[Q2]["proposed"] == TEXT2
        assert e[Q2]["register_proposed"] == "The page shows done."
        # an observation of an unknown question is ignored
        other = wiki_doubts.build_card(s, "h",
                                       observations=[_entry("Q-X-01", "x")])
        assert other == plain
        # a confirm question is never affected by an observation
        s2 = json.loads(json.dumps(s))
        for q in s2["questions"]:
            if q["id"] == Q2:
                for m in q["members"]:
                    m["state"] = "rewritten"
                    m["resolution"] = "R-1"
        c2 = wiki_doubts.build_card(s2, "h", observations=[_entry(Q2, TEXT2)])
        q2 = [x for x in c2["open_questions"] if x["id"] == Q2][0]
        assert q2["kind"] == "confirm" and "proposed_origin" not in q2
    finally:
        T._cleanup(d)


def test_card_cli_and_answer_accept_is_a_correction_and_never_lifts():
    def go(d):
        _write_obs(d, [_entry(Q2, TEXT2)])
        T._commit_all(d, "scratch: an observation")
        card = T._emit(d, {Q2: "accept"})
        c = json.loads(card.read_text(encoding="utf-8"))
        e = [x for x in c["open_questions"] if x["id"] == Q2][0]
        assert e["proposed_origin"] == "tester-observation"
        code, out, err = T._answer(d, card)
        assert code == 0, out + err
        new = [n for n in T._res_files(d) if n != f"R-{STORY}-01.md"]
        fm, body = T._fm_body(d / "resolutions" / new[0])
        assert fm["effect"] == "corrects" and fm["origin"] == "tester-observed"
        assert "confirmed_parts" not in fm and fm["member_basis"], fm
        assert fm["answers"] == Q2 and fm["card"] == card.name
        assert "> " + TEXT2 in body.split("# Adopted answer")[1]
        assert (f"accepted tester carol's observation" in body
                and card.name in body), body
        assert wiki_doubts.load_lifts(d) == wiki_doubts.lifts_from([]), \
            "an observation produced a lift"
        code, out, err = T._sw(d, "doubts", "list", "--json", "--all",
                               "--story", STORY)
        got = {x["doubt"]: x["state"]
               for x in json.loads(out)["stories"][0]["doubts"]}
        assert got[f"{SID}-AC2-02#scenario"] == "answered", got
        assert got[f"{SID}-AC2-02#expected"] == "answered", got
        assert "closed" not in got.values()
        code, out, err = T._sw(d, "lint")
        assert code == 0 and "0 error(s)" in out, out + err
    T._with_scratch(go)


def test_text_answer_on_an_observed_entry_is_human_stated():
    def go(d):
        _write_obs(d, [_entry(Q2, TEXT2)])
        T._commit_all(d, "scratch: an observation")
        card = T._emit(d, {Q2: "The page shows a different thing."})
        code, out, err = T._answer(d, card)
        assert code == 0, out + err
        new = [n for n in T._res_files(d) if n != f"R-{STORY}-01.md"]
        fm, body = T._fm_body(d / "resolutions" / new[0])
        assert fm["effect"] == "corrects" and fm["origin"] == "human-stated"
    T._with_scratch(go)


def test_observation_added_after_the_card_invalidates_it():
    def go(d):
        card = T._emit(d, {Q2: "accept"})
        _write_obs(d, [_entry(Q2, TEXT2)])
        T._expect_refusal(d, card, f"{Q2} differs from the current state")
        # the same card is valid again once the observation is gone
        (d / OBS).unlink()
        code, out, err = T._answer(d, card)
        assert code == 0, out + err
    T._with_scratch(go)


def test_render_after_an_observed_answer_prints_rendered_0():
    def go(d):
        rs, rf = doubts_scratch.RSTORY, doubts_scratch.RFLOW
        q2 = f"Q-{rs}-02"
        _write_obs(d, [_entry(q2, TEXT2)], story=rs)
        T._commit_all(d, "scratch: an observation")
        code, out, err = T._sw(d, "doubts", "card", "--story", rs,
                               "--question", q2)
        assert code == 0, out + err
        rel = out.splitlines()[0].split(":")[0]
        code, out, err = T._sw(d, "card", "revise", rel, "--by", BY,
                               "--answer", f"{q2}=accept")
        assert code == 0, out + err
        code, out, err = T._sw(d, "doubts", "answer", "--card", rel, "--by", BY)
        assert code == 0, out + err
        for a in (["render_sit.py", "--story", rs],
                  ["render_uat.py", "--flow", rf]):
            code, out, err = T._sw_path(d, a)
            assert code == 0 and "rendered 0" in out, out + err
        code, out, err = T._sw(d, "lint")
        assert code == 0 and "0 error(s)" in out, out + err
    T._rwith(go)


def test_l15_checks_the_observations_file():
    d = T._mkroot()
    try:
        before = wiki_doubts.stories_with_doubts(d)
        _write_obs(d, [_entry(Q1, "fine")])
        assert wiki_doubts.l15_errors(d) == [], wiki_doubts.l15_errors(d)
        assert wiki_doubts.stories_with_doubts(d) == before
        bad = _entry(Q1, "x")
        cases = [
            ([{**bad, "question": "Q-NOPE-01"}], "Q-NOPE-01"),
            ([{**bad, "text": "  "}], "text"),
            ([{k: v for k, v in bad.items() if k != "by"}], "by"),
            ([{k: v for k, v in bad.items() if k != "at"}], "at"),
            ([{k: v for k, v in bad.items() if k != "workbook_hash"}],
             "workbook_hash"),
        ]
        for entries, needle in cases:
            _write_obs(d, entries)
            errs = wiki_doubts.l15_errors(d)
            assert any("L15 doubts/observations/" in e and needle in e
                       for e in errs), (needle, errs)
        p = _write_obs(d, [bad])
        p.write_text(p.read_text(encoding="utf-8").replace(
            f"/stories/{STORY}.md", "/stories/US-OTHER.md"), encoding="utf-8")
        assert any("story" in e for e in wiki_doubts.l15_errors(d))
        p.write_text("story: [unclosed\n", encoding="utf-8")
        assert any("invalid YAML" in e for e in wiki_doubts.l15_errors(d))
    finally:
        T._cleanup(d)


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn(); print(f"[PASS] {name}")
            except T.Skip as e:
                print(f"[SKIP] {name} ({e})")
            except Exception as e:
                fails += 1; print(f"[FAIL] {name}: {e}")
    sys.exit(1 if fails else 0)
