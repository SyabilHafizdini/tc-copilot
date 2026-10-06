#!/usr/bin/env python3
"""Doubt collectors: every Medium/Low confidence part of a test case spec is a
doubt with an identity (`<scenario_id>#<part>`) and a basis hash of the authored
text. Doubts are derived from the specs, never authored. Pure functions over
data; importing this module reads nothing."""
import hashlib
import json
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
import wiki
from render_sit import CONF_PARTS, ac_id_of

# (kind, part) -> the spec fields whose authored text the basis hashes.
PART_TEXT = {
    ("sit", "scenario"): ("title", "objective"),
    ("sit", "steps"): ("steps",),
    ("sit", "data"): ("data", "pre_extra"),
    ("sit", "expected"): ("expected",),
    ("uat", "scenario"): ("title", "objective"),
    ("uat", "steps"): ("steps",),
    ("uat", "data"): ("steps",),
    ("uat", "expected"): ("expected",),
}


def part_text(kind, entry, part):
    """The authored text a doubt's basis hashes. kind is "sit" or "uat"."""
    fields = PART_TEXT[(kind, part)]
    return "\n".join(str(entry.get(f) or "") for f in fields)


def part_basis(kind, entry, part):
    """sha256(normalize(part_text(...))) using wiki.sha256 and wiki.normalize."""
    return wiki.sha256(wiki.normalize(part_text(kind, entry, part)))


def _story_of(ref):
    """'/stories/US-X.md#HS-01' -> 'stories/US-X'; None for any other ref."""
    if not isinstance(ref, str):
        return None
    rel = ref.split("#", 1)[0].lstrip("/")
    if not rel.startswith("stories/"):
        return None
    return rel[:-3] if rel.endswith(".md") else rel


def _rows(kind, entry, scenario_id, story):
    conf = entry.get("confidence") or {}
    rem = entry.get("remarks") or {}
    out = []
    for part in CONF_PARTS:
        level = conf.get(part)
        if level is None or level == "High":
            continue
        text = part_text(kind, entry, part)
        out.append({
            "doubt": f"{scenario_id}#{part}",
            "scenario_id": scenario_id,
            "part": part,
            "kind": kind,
            "story": story,
            "level": level,
            "remark": str(rem.get(part) or "").strip(),
            "text": text,
            "basis": wiki.sha256(wiki.normalize(text)),
        })
    return out


def collect_sit(spec):
    """Rows for every non-High part of every test case of one SIT spec dict."""
    story = _story_of(spec.get("story"))
    rows = []
    for tc in spec.get("test_cases") or []:
        sid = spec["scenario_id"].format(
            ac=tc["ac"], ac_id=ac_id_of(spec, tc["ac"]), seq=tc["seq"])
        rows.extend(_rows("sit", tc, sid, story))
    return rows


def collect_uat(spec, flow_fm):
    """Rows for every non-High part of every entry of one UAT spec dict.
    flow_fm is the flow concept's frontmatter; the journey entry whose id equals
    the entry key gives the story through its `ref`."""
    refs = {e.get("id"): e.get("ref")
            for e in ((flow_fm or {}).get("journey") or []) if isinstance(e, dict)}
    rows = []
    for jid, entry in (spec.get("entries") or {}).items():
        sid = spec["scenario_id"].format(jid=jid)
        rows.extend(_rows("uat", entry, sid, _story_of(refs.get(jid))))
    return rows


def _load_spec(p):
    """(spec mapping, None) or (None, why) for one spec file; never raises."""
    try:
        spec = yaml.safe_load(p.read_text(encoding="utf-8"))
    except (yaml.YAMLError, UnicodeDecodeError, OSError) as e:
        return None, f"not readable YAML ({type(e).__name__})"
    if not isinstance(spec, dict):
        return None, "must be a mapping"
    return spec, None


def _flow_fm(root, spec):
    """The flow concept's frontmatter of a UAT spec ({} when absent)."""
    fp = root / str(spec.get("flow")).lstrip("/")
    fm = wiki.read_concept(fp)[0] if fp.is_file() else {}
    return fm if isinstance(fm, dict) else {}


def _spec_rows(kind, p, root):
    """(rows, None) or ([], why) for one spec file. A bad spec yields no rows;
    a UAT spec whose flow concept is missing yields rows with no story."""
    spec, why = _load_spec(p)
    if spec is None:
        return [], why
    if kind == "uat" and not _text(spec.get("flow")):
        return [], "has no flow"
    try:
        if kind == "sit":
            return collect_sit(spec), None
        return collect_uat(spec, _flow_fm(root, spec)), None
    except Exception as e:  # never raise on spec content
        return [], f"cannot be read as a {kind.upper()} spec ({type(e).__name__})"


def collect_all(root=None, sit_dir=None, uat_dir=None, errors=None):
    """Every doubt row in the project, sorted by doubt id. Missing directories
    yield no rows; a spec that cannot be read yields none and, when `errors`
    is a list, appends (path, why) to it. Never raises on spec content."""
    root = Path(root) if root is not None else wiki.ROOT
    sit_dir = Path(sit_dir) if sit_dir is not None else root / "tools" / "sit_specs"
    uat_dir = Path(uat_dir) if uat_dir is not None else root / "tools" / "uat_specs"
    rows = []
    for kind, base in (("sit", sit_dir), ("uat", uat_dir)):
        for p in sorted(base.glob("*.yaml")) if base.is_dir() else []:
            got, why = _spec_rows(kind, p, root)
            rows.extend(got)
            if why and errors is not None:
                errors.append((p, why))
    return sorted(rows, key=lambda r: r["doubt"])


def _uat_specs(root):
    """(path, spec) for every readable UAT spec mapping with a flow."""
    d = root / "tools" / "uat_specs"
    for p in sorted(d.glob("*.yaml")) if d.is_dir() else []:
        spec, _why = _load_spec(p)
        if spec is not None and _text(spec.get("flow")):
            yield p, spec


def _uat_scenario_ids(spec):
    """The scenario ids of a UAT spec's entries ([] when it cannot say)."""
    try:
        return [spec["scenario_id"].format(jid=j) for j in spec.get("entries") or {}]
    except Exception:  # never raise on spec content
        return []


# ---------------------------------------------------------------- register

REGISTER_DIR = "doubts"
QUESTION_KEYS = {"id", "question", "about", "home", "proposed", "members"}
REGISTER_KEYS = {"story", "questions"}
HOME_FORMS = ("ac:<id>, rule:<id>, rule:new, component:<id>, table:<heading> "
              "or none")


def load_register(story_id, root=None):
    """The parsed register for one story, or None when the file is absent.
    Raises ValueError (message names the file) when the YAML does not parse.
    Does no validation beyond parsing."""
    root = Path(root) if root is not None else wiki.ROOT
    p = root / REGISTER_DIR / f"{story_id}.yaml"
    if not p.exists():
        return None
    try:
        return yaml.safe_load(p.read_text(encoding="utf-8"))
    except yaml.YAMLError as e:
        raise ValueError(f"{p.name}: invalid YAML: {e}")


def _ids(fm, key):
    return {e["id"] for e in (fm.get(key) or [])
            if isinstance(e, dict) and "id" in e}


def _headings(body):
    return {m.group(1).strip()
            for m in re.finditer(r"^#{1,6}[ \t]+(.+?)[ \t]*$", body or "", re.M)}


def _home_error(home, story_fm, story_body):
    if not isinstance(home, str):
        return f"home must be a string ({HOME_FORMS})"
    if home in ("none", "rule:new"):
        return None
    kind, _, target = home.partition(":")
    if kind == "ac" and target:
        return None if target in _ids(story_fm, "acceptance_criteria") else \
            f"home ac '{target}' is not an acceptance criterion of the story"
    if kind == "rule" and target:
        return None if target in _ids(story_fm, "business_rules") else \
            f"home rule '{target}' is not a business rule of the story"
    if kind == "component" and target:
        return None if target in _ids(story_fm, "components") else \
            f"home component '{target}' is not a component of the story"
    if kind == "table" and target:
        return None if target in _headings(story_body) else \
            f"home table heading '{target}' not found in the story body"
    return f"home '{home}' is not one of {HOME_FORMS}"


def register_errors(reg, story_fm, doubts, resolutions=(), story_body=""):
    """Error strings for one register (a list, empty when valid). Never raises on
    malformed input: a wrong type is an error string, not an exception."""
    if not isinstance(reg, dict):
        return ["register: must be a mapping with story and questions"]
    errs = []
    story_fm = story_fm if isinstance(story_fm, dict) else {}
    story_id = None
    st = reg.get("story")
    m = re.fullmatch(r"/stories/(.+)\.md", st) if isinstance(st, str) else None
    if m:
        story_id = m.group(1)
        if story_fm.get("id") not in (None, story_id):
            errs.append(f"register: story must be /stories/{story_fm['id']}.md")
    else:
        errs.append("register: story must be /stories/<STORY-ID>.md")
    for k in reg:
        if k not in REGISTER_KEYS:
            errs.append(f"register: unknown key '{k}'")
    qs = reg.get("questions")
    if not isinstance(qs, list):
        errs.append("register: questions must be a list")
        return errs

    answered = {r["answers"] for r in _answer_resolutions(resolutions)}
    frags = wiki.fragment_ids(story_fm)
    by_id = {}
    scen_parts = {}
    for d in doubts:
        by_id[d["doubt"]] = d
        scen_parts.setdefault(d["scenario_id"], set()).add(d["part"])
    id_re = re.compile(rf"Q-{re.escape(story_id)}-\d{{2,}}") if story_id else None
    seen_ids = set()
    holder = {}
    for i, q in enumerate(qs):
        if not isinstance(q, dict):
            errs.append(f"register: question #{i + 1} must be a mapping")
            continue
        qid = q.get("id")
        w = qid if isinstance(qid, str) and qid else f"question #{i + 1}"
        for k in q:
            if k not in QUESTION_KEYS:
                errs.append(f"{w}: unknown key '{k}'")
        if id_re is not None and not (isinstance(qid, str) and id_re.fullmatch(qid)):
            errs.append(f"{w}: id does not match Q-{story_id}-NN")
        if isinstance(qid, str):
            if qid in seen_ids:
                errs.append(f"{w}: duplicate id")
            seen_ids.add(qid)
        text = q.get("question")
        if not isinstance(text, str) or not text.strip():
            errs.append(f"{w}: question must be a non-empty string")
        elif not text.strip().endswith("?"):
            errs.append(f"{w}: question must end with '?'")
        about = q.get("about")
        if not isinstance(about, list) or not about:
            errs.append(f"{w}: about must be a non-empty list")
        else:
            for a in about:
                if not isinstance(a, str) or a not in frags:
                    errs.append(f"{w}: about '{a}' is not a fragment of the story")
        he = _home_error(q.get("home"), story_fm, story_body)
        if he:
            errs.append(f"{w}: {he}")
        pr = q.get("proposed")
        if pr is not None and (not isinstance(pr, str) or not pr.strip()):
            errs.append(f"{w}: proposed must be null or a non-empty string")
        members = q.get("members")
        if not isinstance(members, list):
            errs.append(f"{w}: members must be a list")
            members = []
        if not members and not (isinstance(qid, str) and qid in answered):
            errs.append(f"{w}: has no members and no Resolution answers it")
        in_q = set()
        for mem in members:
            if not isinstance(mem, str) or mem.count("#") != 1 \
                    or not all(mem.split("#")):
                errs.append(f"{w}: member '{mem}' is not <scenario_id>#<part>")
                continue
            sid, part = mem.split("#")
            if mem in in_q:
                errs.append(f"{w}: member '{mem}' is listed twice")
                continue
            in_q.add(mem)
            if sid not in scen_parts:
                errs.append(f"{w}: member '{mem}': unknown scenario id '{sid}'")
            elif part not in CONF_PARTS:
                errs.append(f"{w}: member '{mem}': unknown part '{part}'")
            elif mem not in by_id:
                errs.append(f"{w}: member '{mem}' is not a doubt "
                            "(that part is High or absent)")
            elif story_id and by_id[mem]["story"] != f"stories/{story_id}":
                errs.append(f"{w}: member '{mem}' belongs to another story "
                            f"({by_id[mem]['story']})")
            if mem in by_id:
                other = holder.get(mem)
                if other is not None and other != w:
                    errs.append(f"{w}: member '{mem}' is in two questions "
                                f"({other} and {w})")
                holder.setdefault(mem, w)
    return errs


# ------------------------------------------------------------------ states

def _map(v):
    return v if isinstance(v, dict) else {}


def _members(q):
    """The str members of a question mapping; anything malformed is skipped."""
    if not isinstance(q, dict) or not isinstance(q.get("members"), list):
        return []
    return [m for m in q["members"] if isinstance(m, str)]


def _asserted(resolutions):
    return [r for r in resolutions
            if isinstance(r, dict) and r.get("status") == "asserted"]


ANSWER_EFFECTS = ("confirms", "corrects")
LIFT_KEYS = ("effect", "confirmed_parts", "member_basis")


def _text(v):
    return isinstance(v, str) and bool(v.strip())


def is_answer_resolution(fm):
    """True when a Resolution's frontmatter is shaped as an answer: `answers`
    and `card` are non-empty strings and `effect` is confirms or corrects.
    The one predicate every lift, state and fence uses (whether it is
    asserted is checked separately where it matters)."""
    return (isinstance(fm, dict) and _text(fm.get("answers"))
            and _text(fm.get("card")) and fm.get("effect") in ANSWER_EFFECTS)


def _answer_resolutions(resolutions):
    """The asserted answer Resolutions: the only ones that count."""
    return [r for r in _asserted(resolutions) if is_answer_resolution(r)]


def doubt_states(doubts, register, resolutions, manifest=None):
    """{doubt_id: {"state", "question", "resolution"}} for every row in `doubts`
    whose binding is not retired. register may be None (everything ungrouped).
    closed comes from lifts_from, so a closed doubt is exactly a lifted one."""
    bindings = _map(_map(manifest).get("bindings"))
    held = {}
    qs = _map(register).get("questions")
    for q in (qs if isinstance(qs, list) else []):
        for mem in _members(q):
            held.setdefault(mem, q.get("id"))
    lifts = lifts_from(resolutions)
    res = _answer_resolutions(resolutions)
    out = {}
    for row in doubts:
        if _map(bindings.get(row["scenario_id"])).get("status") == "retired":
            continue
        d, basis = row["doubt"], row["basis"]
        state, rid = "open", None
        hits = [x for x in lifts.get(d, []) if x["basis"] == basis]
        if hits:
            state = "closed"
            rid = max(hits, key=lambda h: _natural(h["resolution"]))["resolution"]
        else:
            corr = [r for r in res if r.get("effect") == "corrects"
                    and d in _map(r.get("member_basis"))]
            same = next((r for r in corr if r["member_basis"][d] == basis), None)
            if same:
                state, rid = "answered", same.get("id")
            elif corr:
                state, rid = "rewritten", corr[0].get("id")
        out[d] = {"state": state, "question": held.get(d), "resolution": rid}
    return out


def sorted_questions(register, doubts, states):
    """Question summaries in impact order: distinct open test cases descending,
    then Low count descending, then id."""
    by_id = {d["doubt"]: d for d in doubts}
    out = []
    qs = _map(register).get("questions")
    for q in (qs if isinstance(qs, list) else []):
        if not isinstance(q, dict):
            continue
        members = []
        for mem in _members(q):
            row, st = by_id.get(mem), states.get(mem)
            if row is None or st is None:
                continue
            members.append({**row, "state": st["state"],
                            "resolution": st["resolution"]})
        live = [m for m in members if m["state"] != "closed"]
        have = {m["state"] for m in members}
        status = ("open" if "open" in have else
                  "rewrite-pending" if "answered" in have else
                  "confirm-pending" if "rewritten" in have else "closed")
        low = sum(1 for m in live if m["level"] == "Low")
        out.append({
            "id": q.get("id"), "question": q.get("question"),
            "about": q.get("about"), "home": q.get("home"),
            "proposed": q.get("proposed"), "status": status,
            "open_tcs": len({m["scenario_id"] for m in live}), "low": low,
            "lowest": ("Low" if low else "Medium" if live else None),
            "members": members,
        })
    return sorted(out, key=lambda x: (-x["open_tcs"], -x["low"], str(x["id"])))


# ----------------------------------------------------- summary, list and lint

def _root(root):
    return Path(root) if root is not None else wiki.ROOT


def _resolutions(root):
    out = []
    d = root / "resolutions"
    if d.is_dir():
        for p in sorted(d.glob("*.md")):
            if p.name == "index.md":
                continue
            fm = wiki.read_concept(p)[0]
            if isinstance(fm, dict):
                out.append(fm)
    return out


def _manifest(root):
    p = root / "manifest.json"
    if not p.exists():
        return {}
    try:
        m = json.loads(p.read_text(encoding="utf-8"))
    except ValueError:
        return {}
    return m if isinstance(m, dict) else {}


def _story_concept(root, story_id):
    p = root / "stories" / f"{story_id}.md"
    if not p.exists():
        return {}, "", False
    fm, body = wiki.read_concept(p)
    return (fm if isinstance(fm, dict) else {}), body, True


def _register_files(root):
    d = root / REGISTER_DIR
    return sorted(d.glob("*.yaml")) if d.is_dir() else []


def _tc_of(manifest, scenario_id):
    b = _map(_map(manifest.get("bindings")).get(scenario_id))
    tc = b.get("tc")
    return Path(tc).name if isinstance(tc, str) and tc else None


# ------------------------------------------------------------ lift and ratchet

_LEVEL_RANK = {"Low": 0, "Medium": 1, "High": 2}


def _natural(s):
    """Sort key putting R-X-2 before R-X-10."""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", str(s))]


def lifts_from(resolutions):
    """{doubt_id: [{"basis", "resolution", "question", "by", "date"}, ...]} from
    every asserted answer Resolution (is_answer_resolution) with effect
    confirms (its confirmed_parts). Corrects Resolutions, and any Resolution
    that is not an answer, contribute nothing. Deterministic order (by
    resolution id)."""
    out = {}
    live = [r for r in _answer_resolutions(resolutions)
            if r.get("effect") == "confirms" and _text(r.get("id"))]
    for r in sorted(live, key=lambda r: _natural(r["id"])):
        for doubt, basis in _map(r.get("confirmed_parts")).items():
            out.setdefault(doubt, []).append({
                "basis": basis, "resolution": r["id"],
                "question": r.get("answers"), "by": r.get("asserted_by"),
                "date": str(r.get("asserted_at") or "")[:10]})
    return out


def load_lifts(root=None):
    """lifts_from over the asserted Resolutions under <root>/resolutions."""
    return lifts_from(_resolutions(_root(root)))


def effective_confidence(conf, rem, sc_id, bases, lifts):
    """(eff_conf, eff_rem, lifted). conf / rem are the spec's authored mappings
    (not mutated). bases = {part: current basis}. A part is lifted when its
    authored level is not High and some lifts[f"{sc_id}#{part}"] entry has
    basis == bases[part] (the highest resolution id wins when several match).
    A lifted part gets level "High" and the remark
    `Source: <R-id> - human answer to <Q-id> (<by>, <date>).`
    lifted = {part: {"authored": <authored level>,
                     "resolution": "/resolutions/<R-id>.md"}} for lifted parts only.
    With no lift, eff_conf == conf and eff_rem == rem, so rendering is unchanged."""
    eff_conf, eff_rem, lifted = dict(conf), dict(rem or {}), {}
    for part in CONF_PARTS:
        authored = conf.get(part)
        if authored in (None, "High"):
            continue
        hits = [x for x in (lifts or {}).get(f"{sc_id}#{part}", [])
                if x["basis"] == bases.get(part)]
        if not hits:
            continue
        x = max(hits, key=lambda h: _natural(h["resolution"]))
        eff_conf[part] = "High"
        eff_rem[part] = (f"Source: {x['resolution']} - human answer to "
                         f"{x['question']} ({x['by']}, {x['date']}).")
        lifted[part] = {"authored": authored,
                        "resolution": f"/resolutions/{x['resolution']}.md"}
    return eff_conf, eff_rem, lifted


def merge_lifted(parts, lifted):
    """confidence_parts mapping with `authored` and `resolution` added (after
    level and remark) to each lifted part and to no other."""
    return {p: ({**v, **lifted[p]} if p in lifted else v)
            for p, v in parts.items()}


def lifted_parts(conf_parts_fm):
    """{part: resolution_ref} read from a test case's frontmatter
    confidence_parts: the parts that carry a `resolution` key. The single
    projection used by the skip test here and by lint/next later."""
    return {p: v["resolution"] for p, v in _map(conf_parts_fm).items()
            if isinstance(v, dict) and v.get("resolution")}


def raise_errors(prior_fm, conf, where, story=None):
    """Error strings for the ratchet. prior_fm = the existing test case file's
    frontmatter (None or no confidence_parts gives []). For each part, the prior
    authored level is confidence_parts[part]["authored"] if present, else
    ["level"]. If conf[part] outranks it: one error. Lowering is allowed."""
    prior = _map(_map(prior_fm).get("confidence_parts"))
    hint = f" --story {story}" if story else ""
    errs = []
    for part in CONF_PARTS:
        old = _map(prior.get(part))
        before = old.get("authored") or old.get("level")
        now = _map(conf).get(part)
        if before in _LEVEL_RANK and now in _LEVEL_RANK \
                and _LEVEL_RANK[now] > _LEVEL_RANK[before]:
            errs.append(f"{where}: confidence.{part} raised from {before} to "
                        f"{now} - only a human answer raises a level "
                        f"(py tools/wiki.py doubts card{hint})")
    return errs


def existing_fm(binding, root=None):
    """Frontmatter of the test case file a manifest binding points at, or None
    when there is no binding or the file is absent."""
    if not isinstance(binding, dict) or not binding.get("tc"):
        return None
    p = _root(root) / (binding["tc"] + ".md")
    if not p.exists():
        return None
    fm = wiki.read_concept(p)[0]
    return fm if isinstance(fm, dict) else None


def lift_changed(binding, lifted, root=None):
    """True when the lifted-part projection an entry would render differs from
    the one in its existing file (a new lift, or one that no longer holds).
    Compares only that projection: a spec-only remark edit still needs --force.
    The one comparison the renderers' skip and unrendered_lifts share."""
    would = {p: v["resolution"] for p, v in lifted.items()}
    fm = existing_fm(binding, root=root) or {}
    return would != lifted_parts(fm.get("confidence_parts"))


def _flow_of_scenarios(root):
    """{scenario_id: flow id} over every UAT spec (the id of the flow concept,
    else the spec's flow file stem)."""
    out = {}
    for _p, spec in _uat_specs(root):
        fid = _flow_fm(root, spec).get("id") or Path(str(spec["flow"])).stem
        for sc in _uat_scenario_ids(spec):
            out[sc] = fid
    return out


def _lift_possible(root):
    """False only when no lift can exist or be pending: no Resolution carries
    `confirmed_parts:` (a lift needs an answer Resolution with them) and no
    test case names a `resolution`. A cheap raw-text scan, a superset of the
    real condition (a false True only costs the full pass)."""
    for d, needle in (("resolutions", b"confirmed_parts:"),
                      ("testcases", b"resolution:")):
        base = root / d
        for p in base.rglob("*.md") if base.is_dir() else []:
            if needle in p.read_bytes():
                return True
    return False


def unrendered_lifts(root=None):
    """Every (test case, part) whose rendered lift differs from the lift its
    spec + asserted Resolutions now give: a list of dicts
    {"story", "kind", "flow", "scenario_id", "part", "doubt", "tc", "expected",
    "rendered"}, expected / rendered being the resolution ref or None. Covers a
    confirmed lift not yet rendered and a rendered lift that no longer holds.
    Unbound or retired scenarios and scenarios with no file are skipped. Sorted
    by doubt id."""
    root = _root(root)
    if not _lift_possible(root):
        return []
    rows = collect_all(root=root)
    lifts = load_lifts(root)
    bindings = _map(_manifest(root).get("bindings"))
    flows = _flow_of_scenarios(root)
    by_sc = {}
    for r in rows:
        by_sc.setdefault(r["scenario_id"], []).append(r)
    out = []
    for sc, rs in by_sc.items():
        b = bindings.get(sc)
        if not isinstance(b, dict) or b.get("status") == "retired":
            continue
        fm = existing_fm(b, root=root)
        if fm is None:
            continue
        _c, _r, lifted = effective_confidence(
            {r["part"]: r["level"] for r in rs}, {}, sc,
            {r["part"]: r["basis"] for r in rs}, lifts)
        rendered = lifted_parts(fm.get("confidence_parts"))
        for r in rs:
            want = lifted[r["part"]]["resolution"] if r["part"] in lifted else None
            have = rendered.get(r["part"])
            if want == have:
                continue
            out.append({"story": (r["story"] or "")[len("stories/"):],
                        "kind": r["kind"], "flow": flows.get(sc),
                        "scenario_id": sc, "part": r["part"],
                        "doubt": r["doubt"], "tc": b.get("tc"),
                        "expected": want, "rendered": have})
    return sorted(out, key=lambda e: e["doubt"])


def _register_error_text(e):
    """ValueError text from load_register without its leading file name."""
    s = str(e)
    return s.split(": ", 1)[1] if ": " in s else s


def story_summary(story_id, root=None, _ctx=None):
    """Everything known about one story's doubts. Never raises on a bad or
    missing register: problems go into "register_errors"."""
    root = _root(root)
    rows, res, man = _ctx if _ctx is not None else (
        collect_all(root=root), _resolutions(root), _manifest(root))
    mine = [r for r in rows if r["story"] == f"stories/{story_id}"]
    fm, body, _ = _story_concept(root, story_id)
    exists = (root / REGISTER_DIR / f"{story_id}.yaml").exists()
    reg, errs = None, []
    try:
        reg = load_register(story_id, root=root)
    except ValueError as e:
        errs = [_register_error_text(e)]
    if reg is not None:
        try:
            errs = register_errors(reg, fm, rows, res, body)
        except Exception as e:  # never raise on register content
            errs = [f"register: cannot be validated ({type(e).__name__})"]
    states = doubt_states(mine, reg, res, man)
    doubts = []
    for r in mine:
        st = states.get(r["doubt"])
        if st is None:
            continue
        doubts.append({**r, "state": st["state"], "question": st["question"],
                       "resolution": st["resolution"],
                       "tc": _tc_of(man, r["scenario_id"])})
    questions = sorted_questions(reg, doubts, states)
    live = [d for d in doubts if d["state"] != "closed"]
    ungrouped = [d for d in live if d["question"] is None]
    return {"story": story_id, "open": len(live),
            "questions_open": sum(1 for q in questions if q["status"] != "closed"),
            "ungrouped": len(ungrouped), "register": exists,
            "register_errors": errs, "questions": questions,
            "ungrouped_doubts": ungrouped, "doubts": doubts}


def stories_with_doubts(root=None, _rows=None):
    """Sorted story ids that have at least one doubt row or a register file."""
    root = _root(root)
    rows = _rows if _rows is not None else collect_all(root=root)
    ids = {r["story"][len("stories/"):] for r in rows if r["story"]}
    ids.update(p.stem for p in _register_files(root))
    return sorted(ids)


def summary_line(summary):
    return (f"{summary['story']}: {summary['open']} open doubts in "
            f"{summary['questions_open']} questions ({summary['ungrouped']} ungrouped)")


def _ascii(s):
    s = unicodedata.normalize("NFKD", str(s if s is not None else ""))
    return s.encode("ascii", "replace").decode("ascii")


USAGE = ("usage: wiki doubts list [--story <ID>] [--all] [--ungrouped] [--json]\n"
         "       wiki doubts card --story <ID> [--question Q1[,Q2,...]] [--top N]\n"
         "       wiki doubts answer --card <path> --by <user> [--story <ID>]\n"
         "       wiki doubts observe --workbook <xlsx> --by <tester>")


def _view(summary, show_all, ungrouped_only):
    """The summary as list shows it: closed questions and doubts left out
    unless show_all; only the ungrouped rows when ungrouped_only."""
    v = dict(summary)
    if ungrouped_only:
        v["questions"] = []
        v.pop("doubts")
        return v
    qs = []
    for q in summary["questions"]:
        if q["status"] == "closed" and not show_all:
            continue
        mem = [m for m in q["members"] if show_all or m["state"] != "closed"]
        qs.append({**q, "members": mem})
    v["questions"] = qs
    v["doubts"] = [d for d in summary["doubts"] if show_all or d["state"] != "closed"]
    return v


def _print_story(v):
    print(summary_line(v))
    if v["register_errors"]:
        print("register invalid - run: py tools/wiki.py lint")
    for q in v["questions"]:
        print(_ascii(f"{q['id']} [{q['status']}] {q['open_tcs']} open test cases, "
                     f"lowest {q['lowest'] or '-'}: {q['question']}"))
        for m in q["members"]:
            print(_ascii(f"    {m['doubt']}  {m['level']}  {m['state']}  "
                         f"{m['tc'] or '-'}"))
    if v["ungrouped_doubts"]:
        print("Ungrouped")
        for d in v["ungrouped_doubts"]:
            print(_ascii(f"    {d['doubt']}  {d['level']}  {d['tc'] or '-'}  "
                         f"{d['remark']}"))


# ------------------------------------------------------------------- card

CARD_KEYS = ("doubt", "tc", "part", "level", "remark", "text", "basis")
CARD_USAGE = "usage: wiki doubts card --story <ID> [--question Q1[,Q2,...]] [--top N]"


def register_hash(story_id, root=None):
    """sha256 of the normalized register file text ("" when absent)."""
    p = _root(root) / REGISTER_DIR / f"{story_id}.yaml"
    if not p.exists():
        return ""
    return wiki.sha256(wiki.normalize(p.read_text(encoding="utf-8")))


OBSERVATION_DIR = "doubts/observations"
OBSERVED = "tester-observation"


def observations_path(story_id, root=None):
    return _root(root) / OBSERVATION_DIR / f"{story_id}.yaml"


def load_observations(story_id, root=None):
    """The observation entries of one story in file order ([] when the file is
    absent). Raises ValueError (message names the file) when the file does not
    parse or is not {story, observations: [mapping, ...]}."""
    p = observations_path(story_id, root)
    if not p.exists():
        return []
    try:
        doc = yaml.safe_load(p.read_text(encoding="utf-8"))
    except yaml.YAMLError as e:
        raise ValueError(f"{p.name}: invalid YAML: {e}")
    obs = doc.get("observations") if isinstance(doc, dict) else None
    if not isinstance(obs, list) or any(not isinstance(o, dict) for o in obs):
        raise ValueError(f"{p.name}: must be a mapping with story and an "
                         f"observations list of mappings")
    return obs


def latest_observations(entries):
    """{question id: its latest observation} (by `at`, then file order)."""
    out = {}
    for o in entries:
        q = o.get("question")
        if isinstance(q, str) and isinstance(o.get("text"), str) and o["text"].strip():
            if q not in out or str(o.get("at")) >= str(out[q].get("at")):
                out[q] = o
    return out


def _card_entry(q, observed=None):
    """The card entry of one question, or None when it is not eligible.
    `observed` is the question's latest tester observation, if any."""
    live = [m for m in q["members"] if m["state"] == "open"]
    kind, proposed = "ask", q["proposed"]
    if not live:
        # No open member: the rewritten members, and the answered ones whose
        # text needed no rewrite, are confirmed as they now stand.
        live = [m for m in q["members"] if m["state"] in ("rewritten", "answered")]
        if not live:
            return None
        kind = "confirm"
        rewritten = [m for m in live if m["state"] == "rewritten"]
        rid = sorted(str(m["resolution"]) for m in (rewritten or live))[0]
        proposed = (f"Rewritten after {rid}; confirm the text shown for each "
                    f"member." if rewritten else
                    f"Answered by {rid}; confirm the text shown for each "
                    f"member is correct as it stands.")
    entry = {"id": q["id"], "kind": kind, "question": q["question"],
             "about": q["about"], "home": q["home"], "proposed": proposed,
             "members": [{k: m[k] for k in CARD_KEYS} for m in live]}
    if observed is not None and kind == "ask":
        entry.update({"proposed": observed["text"], "proposed_origin": OBSERVED,
                      "observed_by": observed.get("by"),
                      "observation_at": observed.get("at"),
                      "register_proposed": proposed})
    return entry


def _ineligible_reason(q):
    """Why a question has no card entry: an open, answered or rewritten
    member always gives one, so only a closed question has none."""
    return "closed"


def build_card(summary, reg_hash, question_ids=None, top=None,
               observations=None):
    """The card dict, pure: the same inputs give an identical dict.
    `observations` is the story's observation entries (a tester's observation
    becomes the proposed text of an ask question, marked as such). Raises
    ValueError for a bad --question id or --top."""
    if top is not None and (not isinstance(top, int) or isinstance(top, bool)
                            or top < 1):
        raise ValueError("--top must be a positive integer")
    by_id = {q["id"]: q for q in summary["questions"]}
    seen = latest_observations(observations or [])
    entries = {q["id"]: _card_entry(q, seen.get(q["id"]))
               for q in summary["questions"]}
    for qid in question_ids or []:
        if qid not in by_id:
            raise ValueError(f"unknown question {qid}; the register has: "
                             f"{', '.join(sorted(map(str, by_id))) or '(none)'}")
        if entries[qid] is None:
            raise ValueError(f"question {qid} is not eligible: "
                             f"{_ineligible_reason(by_id[qid])}")
    wanted = set(question_ids) if question_ids else None
    out = [e for qid, e in entries.items()
           if e is not None and (wanted is None or qid in wanted)]
    if top is not None:
        out = out[:top]
    return {"card_type": "doubts", "story": summary["story"],
            "register_hash": reg_hash, "open_questions": out}


def write_doubts_card(card, root=None):
    """Write build/cards/doubts-<STORY>-NNN.json; returns the path."""
    d = _root(root) / "build" / "cards"
    d.mkdir(parents=True, exist_ok=True)
    pre = f"doubts-{card['story']}-"
    nums = [int(m.group(1)) for p in d.glob(pre + "*.json")
            if (m := re.fullmatch(re.escape(pre) + r"(\d{3})\.json", p.name))]
    path = d / f"{pre}{max(nums, default=0) + 1:03d}.json"
    path.write_text(json.dumps(card, indent=1, ensure_ascii=False) + "\n",
                    encoding="utf-8", newline="\n")
    return path


def _refuse(msg):
    print(_ascii(msg), file=sys.stderr)
    sys.exit(1)


def _cmd_card(args, root):
    story = qs = top = None
    i = 0
    while i < len(args):
        a = args[i]
        if a in ("--story", "--question", "--top") and i + 1 < len(args):
            v = args[i + 1]
            if a == "--story":
                story = v
            elif a == "--question":
                qs = [x for x in v.split(",") if x]
            else:
                try:
                    top = int(v)
                except ValueError:
                    _refuse(f"doubts card refused: --top must be a positive "
                            f"integer, got {v!r}")
            i += 2
        else:
            print(CARD_USAGE)
            sys.exit(2)
    if not story:
        print(CARD_USAGE)
        sys.exit(2)
    s = story_summary(story, root=root)
    if not s["register"]:
        _refuse(f"doubts card refused: {story} has no register "
                f"doubts/{story}.yaml (group the doubts first: tc-resolve)")
    if s["register_errors"]:
        _refuse(f"doubts card refused: doubts/{story}.yaml is invalid:\n  "
                + "\n  ".join(s["register_errors"])
                + "\nrun: py tools/wiki.py lint")
    try:
        card = build_card(s, register_hash(story, root=root), qs, top,
                          load_observations(story, root=root))
    except ValueError as e:
        _refuse(f"doubts card refused: {e}")
    if not card["open_questions"]:
        _refuse(f"doubts card refused: no question of {story} is eligible "
                f"(every one is closed)")
    path = write_doubts_card(card, root=root)
    rel = path.relative_to(_root(root)).as_posix()
    n = len(card["open_questions"])
    first = card["open_questions"][0]["id"]
    print(f"{rel}: {n} question(s)")
    print("next:")
    print(f'  py tools/wiki.py card revise {rel} --by <user> '
          f'--answer {first}=accept   (or --answer {first}="<text>")')
    print(f"  py tools/wiki.py doubts answer --card {rel} --by <user>")
    return path


# ----------------------------------------------------------------- answer

ANSWER_REFUSED = "doubts answer refused: "
COMPARED_KEYS = ("id", "kind", "question", "about", "home", "proposed",
                 "proposed_origin", "observed_by", "observation_at",
                 "register_proposed")
COMPARED_MEMBER_KEYS = ("doubt", "part", "level", "text", "basis")
TITLE_QUESTION_CHARS = 70
ANSWER_USAGE = ("usage: wiki doubts answer --card <path> --by <user> "
                "[--story <ID>]")


def _answer_refuse(reason):
    _refuse(ANSWER_REFUSED + reason)


def _answer_args(args):
    card = by = story = None
    i = 0
    while i < len(args):
        a = args[i]
        if a in ("--card", "--by", "--story") and i + 1 < len(args):
            v = args[i + 1]
            if a == "--card":
                card = v
            elif a == "--by":
                by = v
            else:
                story = v
            i += 2
        else:
            _answer_refuse(ANSWER_USAGE)
    if not card:
        _answer_refuse("--card <path> is required")
    if not by:
        _answer_refuse("--by <user> is required (only a human answers)")
    return card, by, story


def _same_entry(a, b):
    if any(a.get(k) != b.get(k) for k in COMPARED_KEYS):
        return False
    ma, mb = a.get("members") or [], b.get("members") or []
    if len(ma) != len(mb):
        return False
    return all(all(x.get(k) == y.get(k) for k in COMPARED_MEMBER_KEYS)
               for x, y in zip(ma, mb))


def _tc_problem(root, manifest, doubt):
    """None when the member's test case is bound, present and not stale;
    otherwise the reason the member cannot be accepted."""
    sid = doubt.split("#", 1)[0]
    tc = _map(_map(manifest.get("bindings")).get(sid)).get("tc")
    if not isinstance(tc, str) or not tc:
        return "has no bound test case yet"
    p = root / (tc + ".md")
    if not p.exists():
        return f"has no test case file ({tc}.md)"
    fm = wiki.read_concept(p)[0]
    fm = fm if isinstance(fm, dict) else {}
    if fm.get("status") == "stale" or fm.get("stale_because"):
        return f"has a stale test case ({Path(tc).name})"
    return None


def _next_number(root, story):
    pre = f"R-{story}-"
    d = root / "resolutions"
    nums = [int(m.group(1)) for p in (d.glob(pre + "*.md") if d.is_dir() else [])
            if (m := re.fullmatch(re.escape(pre) + r"(\d+)\.md", p.name))]
    return max(nums, default=0) + 1


def _resolution_title(story, qid, question):
    q = " ".join(str(question).split())
    if len(q) > TITLE_QUESTION_CHARS:
        q = q[:TITLE_QUESTION_CHARS].rstrip() + "..."
    return f"{story} {qid}: {q}"


def _quote(text):
    return "\n".join("> " + ln if ln else ">" for ln in text.split("\n"))


def _effect(entry, ans):
    """confirms only for an accepted agent proposal; a correction, and an
    accepted tester observation, are corrections (nothing lifts from them)."""
    if ans["decision"] == "accept" and entry.get("proposed_origin") != OBSERVED:
        return "confirms"
    return "corrects"


def _resolution(story, rid, entry, ans, by, card_name, at):
    """(frontmatter, body) of the Resolution for one answered question. Human
    and tester text under `# Adopted answer` is a blockquote, like `# Human
    Statement`, so a line in it can never pass for a section heading."""
    accept = ans["decision"] == "accept"
    observed = accept and entry.get("proposed_origin") == OBSERVED
    qid = entry["id"]
    bases = {m["doubt"]: m["basis"] for m in entry["members"]}
    fm = {
        "type": "Resolution", "id": rid,
        "title": _resolution_title(story, qid, entry["question"]),
        "description": f"Human answer to doubt question {qid} on card "
                       f"{card_name}.",
        "origin": ("tester-observed" if observed else
                   "agent-proposed" if accept else "human-stated"),
        "status": "asserted", "asserted_by": by, "asserted_at": at,
        "resolves": [f"/stories/{story}.md#{frag}" for frag in entry["about"]],
        "answers": qid, "effect": _effect(entry, ans), "card": card_name,
    }
    fm["member_basis" if observed or not accept else "confirmed_parts"] = bases
    if observed:
        statement = (f"> Card {card_name}: the human accepted tester "
                     f"{entry.get('observed_by')}'s observation (recorded "
                     f"{entry.get('observation_at')}) as the answer; adopted "
                     f"verbatim below.")
        adopted = _quote(entry["proposed"])
        prop = (f"Pending. Write the answer to {entry['home']}, rewrite the "
                f"parts above, render, then confirm on a new doubts card. A "
                f"tester's observation never lifts a level.")
    elif accept:
        statement = (f'> Card {card_name} answer: "accept" -- adopts the '
                     f"proposed answer below verbatim.")
        adopted = entry["proposed"]
        prop = ("None. Confirms the current text of the parts above; the "
                "levels lift on the next render.")
    else:
        statement = _quote(ans["value"])
        adopted = _quote(ans["value"])
        prop = (f"Pending. Write the answer to {entry['home']}, rewrite the "
                f"parts above, render, then confirm on a new doubts card.")
    applies = "\n".join(f"- {m['doubt']} - {m.get('tc') or '(no test case yet)'}"
                        for m in entry["members"])
    body = (f"\n# Human Statement\n\n{statement}\n\n# Question\n\n"
            f"{entry['question']}\n\n# Adopted answer\n\n{adopted}\n\n"
            f"# Applies to\n\n{applies}\n\n# Propagation\n\n{prop}\n")
    return fm, body


def _uat_flows(root, doubt_ids):
    """Flow ids whose UAT spec holds one of the doubts."""
    out = set()
    for _p, spec in _uat_specs(root):
        sids = set(_uat_scenario_ids(spec))
        if any(x.split("#", 1)[0] in sids for x in doubt_ids):
            out.add(Path(str(spec["flow"])).stem)
    return sorted(out)


def _cmd_answer(args, root):
    card_arg, by, story_arg = _answer_args(args)
    card_path = Path(card_arg)
    if not card_path.is_absolute():
        card_path = root / card_path
    if not card_path.exists():
        _answer_refuse(f"card not found: {card_arg}")
    try:
        card = json.loads(card_path.read_text(encoding="utf-8"))
    except ValueError as e:
        _answer_refuse(f"{card_path.name} is not valid JSON ({e})")
    if not isinstance(card, dict) or card.get("card_type") != "doubts":
        _answer_refuse(f"{card_path.name} is not a doubts card "
                       f"(card_type must be 'doubts')")
    story = card.get("story")
    if not isinstance(story, str) or not story:
        _answer_refuse(f"{card_path.name} names no story")
    if story_arg is not None and story_arg != story:
        _answer_refuse(f"{card_path.name} is for story {story}, not {story_arg}")
    if not (root / REGISTER_DIR / f"{story}.yaml").exists():
        _answer_refuse(f"story {story} has no register doubts/{story}.yaml, so "
                       f"this card resolves to no story")
    hr = card.get("human_response")
    answers = hr.get("answers") if isinstance(hr, dict) else None
    if not isinstance(answers, dict) or not answers:
        _answer_refuse(f"{card_path.name} has no human_response.answers "
                       f"(answer it first: py tools/wiki.py card revise ...)")
    if hr.get("by") != by:
        _answer_refuse(f"--by {by} differs from the card's human_response.by "
                       f"({hr.get('by')})")
    if card.get("applied"):
        _answer_refuse(f"{card_path.name} is already applied; emit a new card "
                       f"for the questions still open")
    entries = {e.get("id"): e for e in card.get("open_questions") or []
               if isinstance(e, dict)}
    for qid in answers:
        if qid not in entries:
            _answer_refuse(f"answer for {qid}, which is not on {card_path.name} "
                           f"(it has: {', '.join(map(str, entries)) or '(none)'})")
    summary = story_summary(story, root=root)
    if summary["register_errors"]:
        _answer_refuse(f"doubts/{story}.yaml is invalid:\n  "
                       + "\n  ".join(summary["register_errors"])
                       + "\nrun: py tools/wiki.py lint")
    try:
        card_rel = card_path.relative_to(root).as_posix()
    except ValueError:
        card_rel = str(card_path)
    changed = (f"the register or the specs changed after the card was emitted; "
               f"discard it (py tools/wiki.py card discard {card_rel} --by "
               f"{by}) and emit a new card: py tools/wiki.py doubts card "
               f"--story {story}")
    ids = sorted(answers)
    try:
        fresh = build_card(summary, register_hash(story, root=root), ids,
                           observations=load_observations(story, root=root))
    except ValueError as e:
        _answer_refuse(f"{changed} ({e})")
    if fresh["register_hash"] != card.get("register_hash"):
        _answer_refuse(f"register_hash differs: {changed}")
    now_entries = {e["id"]: e for e in fresh["open_questions"]}
    manifest = _manifest(root)
    plan = {}
    for qid in ids:
        ans = answers[qid]
        if not isinstance(ans, dict) or ans.get("decision") not in ("accept",
                                                                    "correct"):
            _answer_refuse(f"{qid}: decision must be 'accept' or 'correct'")
        if qid not in now_entries or not _same_entry(entries[qid],
                                                     now_entries[qid]):
            _answer_refuse(f"{qid} differs from the current state: {changed}")
        entry = now_entries[qid]
        if ans["decision"] == "correct":
            if not isinstance(ans.get("value"), str) or not ans["value"].strip():
                _answer_refuse(f"{qid}: a correcting answer needs non-empty text")
        else:
            if entry["proposed"] is None:
                _answer_refuse(f"{qid}: cannot accept, the question has no "
                               f"proposed answer (give the answer as text)")
            for m in entry["members"]:
                why = _tc_problem(root, manifest, m["doubt"])
                if why:
                    _answer_refuse(f"{qid}: cannot accept, {m['doubt']} {why}")
        plan[qid] = (entry, ans)

    # Everything validated. Before the first write: the commit must hold only
    # what this command writes, and the repository must already lint clean,
    # so the commit gate cannot refuse for a reason that predates the answer.
    _answer_preflight(root)
    at = wiki.now_iso()
    qs = ",".join(ids)
    log = root / "log.md"
    saved = {log: log.read_bytes() if log.exists() else None,
             card_path: card_path.read_bytes()}
    head = _git(root, "rev-parse", "HEAD").stdout.strip()
    num = _next_number(root, story)
    made, written = {}, []
    try:
        for qid in ids:
            rid = f"R-{story}-{num:02d}"
            num += 1
            entry, ans = plan[qid]
            fm, body = _resolution(story, rid, entry, ans, by, card_path.name, at)
            written.append(root / "resolutions" / f"{rid}.md")
            wiki.write_concept(written[-1], fm, body)
            made[qid] = rid
            wiki.append_log(f"**Assertion ({by})**: doubt answer {rid} "
                            f"({fm['effect']}) to {qid} of {story} "
                            f"via card {card_path.name}")
        card["applied"] = {"by": by, "at": at, "resolutions": made}
        card_path.write_text(json.dumps(card, indent=1, ensure_ascii=False)
                             + "\n", encoding="utf-8", newline="\n")
        wiki.agent_commit(
            f"doubts answer: {story} {qs} by {by}\n\n"
            f"Assertion-Event: cli-doubts-answer {story} {qs} by {by} at {at} "
            f"card {card_path.name}")
        ok = ("--no-commit" in sys.argv
              or _git(root, "rev-parse", "HEAD").stdout.strip() != head)
    except (Exception, SystemExit):
        ok = False
    if not ok:
        _rollback(root, written, saved)
        _answer_refuse(f"the commit did not go through, so the answer was "
                       f"rolled back: no Resolution written, log.md and "
                       f"{card_path.name} restored (see the messages above; "
                       f"fix the cause and run the same command again)")
    state = {d["doubt"]: d["state"]
             for d in story_summary(story, root=root)["doubts"]}
    for qid in ids:
        entry, ans = plan[qid]
        eff = _effect(entry, ans)
        print(f"resolutions/{made[qid]}.md  ({qid}, {eff})")
        for m in entry["members"]:
            print(_ascii(f"    {m['doubt']}  {state.get(m['doubt'], '-')}"))
    print("next:")
    confirms = [plan[q][0] for q in ids if _effect(*plan[q]) == "confirms"]
    if confirms:
        print(f"  py tools/render_sit.py --story {story}")
        for f in _uat_flows(root, [m["doubt"] for e in confirms
                                   for m in e["members"]]):
            print(f"  py tools/render_uat.py --flow {f}")
        print("  py tools/wiki.py seal")
    for qid in ids:
        entry, ans = plan[qid]
        if _effect(entry, ans) == "corrects":
            print(f"  {qid}: write the answer to {entry['home']}, rewrite the "
                  f"parts, render, then: py tools/wiki.py doubts card "
                  f"--story {story}")


def _git(root, *args):
    return subprocess.run(["git", *args], cwd=root, capture_output=True,
                          text=True, encoding="utf-8", errors="replace")


def _answer_preflight(root):
    """Refuse (exit 1, nothing written) unless the working tree is clean and
    the repository lints clean. build/ is git-ignored, so the card itself
    never counts as a change."""
    st = _git(root, "status", "--porcelain")
    if st.returncode != 0:
        _answer_refuse(f"git status failed, so the answer could not be "
                       f"committed: {st.stderr.strip()}")
    dirty = [ln for ln in st.stdout.splitlines() if ln.strip()]
    if dirty:
        _answer_refuse("the working tree has other uncommitted changes; commit "
                       "or discard them first, so the answer's commit holds "
                       "only what it writes:\n  " + "\n  ".join(dirty))
    r = subprocess.run([sys.executable, str(Path(wiki.__file__).resolve()),
                        "lint", "--no-commit"], cwd=root, capture_output=True,
                       text=True, encoding="utf-8", errors="replace")
    if r.returncode != 0:
        errs = [ln for ln in r.stdout.splitlines() if ln.startswith("ERROR")]
        _answer_refuse("the repository does not lint clean, so the answer "
                       "could not be committed; fix these first (nothing was "
                       "written):\n  " + "\n  ".join(errs or ["(lint failed)"]))


def _rollback(root, written, saved):
    """Undo a write whose commit failed: delete the new Resolutions, restore
    the saved files to their prior bytes, unstage anything staged."""
    for p in written:
        p.unlink(missing_ok=True)
    for p, data in saved.items():
        if data is None:
            p.unlink(missing_ok=True)
        else:
            p.write_bytes(data)
    _git(root, "reset", "-q")


# ---------------------------------------------------------------- observe

OBSERVE_REFUSED = "doubts observe refused: "
OBSERVE_USAGE = "usage: wiki doubts observe --workbook <xlsx> --by <tester>"
DOUBT_SHEET = "AI Doubts"
QUESTION_ID_RE = re.compile(r"Q-(.+)-\d+")


def _observe_refuse(reason):
    _refuse(OBSERVE_REFUSED + reason)


def _observe_args(args):
    wb = by = None
    i = 0
    while i < len(args):
        a = args[i]
        if a in ("--workbook", "--by") and i + 1 < len(args):
            if a == "--workbook":
                wb = args[i + 1]
            else:
                by = args[i + 1]
            i += 2
        else:
            _observe_refuse(OBSERVE_USAGE)
    if not wb:
        _observe_refuse("--workbook <xlsx> is required")
    if not by:
        _observe_refuse("--by <tester> is required")
    return wb, by


def _cell_text(v):
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def _sheet_rows(path):
    """[{header: cell value}] of the AI Doubts sheet, columns found by header
    text. Refuses (exit 1) when the sheet or a required header is missing."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        _observe_refuse(f"{path.name} is not an xlsx workbook")
    try:
        if DOUBT_SHEET not in wb.sheetnames:
            _observe_refuse(f"{path.name} has no '{DOUBT_SHEET}' sheet")
        rows = [list(r) for r in wb[DOUBT_SHEET].iter_rows(values_only=True)]
    finally:
        wb.close()
    head = {}
    start = 0
    for n, r in enumerate(rows):
        found = {str(v).strip(): c for c, v in enumerate(r) if v is not None}
        if "Question ID" in found or "Observation" in found:
            head, start = found, n + 1
            break
    for need in ("Question ID", "Observation"):
        if need not in head:
            _observe_refuse(f"the '{DOUBT_SHEET}' sheet has no '{need}' column")
    out = []
    for r in rows[start:]:
        out.append({h: (r[c] if c < len(r) else None) for h, c in head.items()})
    return out


DEDUP_KEYS = ("question", "text", "workbook_hash")


def _fresh_observations(have, entries):
    """The entries not already recorded (same question, text and workbook
    hash) and not repeated among themselves, in order. A malformed existing
    entry (a non-string key field) is no duplicate of anything."""
    keys = {tuple(o.get(k) for k in DEDUP_KEYS) for o in have
            if isinstance(o, dict)
            and all(isinstance(o.get(k), str) for k in DEDUP_KEYS)}
    fresh = []
    for e in entries:
        k = tuple(e[f] for f in DEDUP_KEYS)
        if k not in keys:
            keys.add(k)
            fresh.append(e)
    return fresh


def _cmd_observe(args, root):
    wb_arg, by = _observe_args(args)
    path = Path(wb_arg)
    if not path.is_absolute():
        path = Path.cwd() / path
    if not path.is_file():
        _observe_refuse(f"workbook not found: {wb_arg}")
    rows = _sheet_rows(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    at = wiki.now_iso()
    skipped, found = [], {}
    for r in rows:
        text = _cell_text(r.get("Observation"))
        if text is None or not text.strip():
            continue
        qid = (_cell_text(r.get("Question ID")) or "").strip()
        m = QUESTION_ID_RE.fullmatch(qid)
        if qid in ("", "-") or not m:
            skipped.append((qid or "(blank)", "not a grouped question "
                            "(ungrouped doubt: ask the agent to group it)"))
            continue
        story = m.group(1)
        try:
            reg = load_register(story, root=root)
        except ValueError as e:
            _observe_refuse(f"doubts/{story}.yaml is invalid: "
                            f"{_register_error_text(e)}")
        qs = _map(reg).get("questions")
        known = {q.get("id") for q in (qs if isinstance(qs, list) else [])
                 if isinstance(q, dict) and isinstance(q.get("id"), str)}
        if qid not in known:
            skipped.append((qid, "not a question of any register"))
            continue
        who = (_cell_text(r.get("Observed by")) or "").strip() or by
        found.setdefault(story, []).append({
            "question": qid, "text": text, "by": who, "at": at,
            "observed_on": ((_cell_text(r.get("Date")) or "").strip() or None),
            "workbook": path.name, "workbook_hash": digest})
    plan = {}
    for story, entries in sorted(found.items()):
        try:
            have = load_observations(story, root=root)
        except ValueError as e:
            _observe_refuse(f"doubts/observations/{e}")
        fresh = _fresh_observations(have, entries)
        if fresh:
            plan[story] = (have, fresh)
    for qid, why in skipped:
        print(_ascii(f"skipped {qid}: {why}"))
    if not plan:
        print("no new observation found (nothing written)")
        return
    parts = []
    for story, (have, fresh) in plan.items():
        dest = observations_path(story, root)
        dest.parent.mkdir(parents=True, exist_ok=True)
        doc = {"story": f"/stories/{story}.md", "observations": have + fresh}
        dest.write_text(yaml.safe_dump(doc, sort_keys=False, allow_unicode=True,
                                       width=1000),
                        encoding="utf-8", newline="\n")
        parts.append(f"{story} {len(fresh)} observation(s)")
        print(f"{OBSERVATION_DIR}/{story}.yaml: {len(fresh)} observation(s) "
              f"appended")
    print("next:")
    for story in plan:
        print(f"  py tools/wiki.py doubts card --story {story}   (the "
              f"observation appears as the proposed answer; a human still "
              f"answers)")
    wiki.agent_commit(f"doubts observe: {', '.join(parts)} by {by}")


def cmd_doubts(args, root=None):
    """`wiki doubts list|card|answer|observe ...`."""
    flags = {"--all", "--ungrouped", "--json"}
    if args and args[0] == "answer":
        _cmd_answer(args[1:], _root(root))
        return
    if args and args[0] == "card":
        _cmd_card(args[1:], _root(root))
        return
    if args and args[0] == "observe":
        _cmd_observe(args[1:], _root(root))
        return
    if not args or args[0] != "list":
        print(USAGE)
        sys.exit(2)
    rest = args[1:]
    story = None
    show = set()
    i = 0
    while i < len(rest):
        a = rest[i]
        if a == "--story" and i + 1 < len(rest):
            story, i = rest[i + 1], i + 2
        elif a in flags:
            show.add(a)
            i += 1
        else:
            print(USAGE)
            sys.exit(2)
    root = _root(root)
    rows = collect_all(root=root)
    known = stories_with_doubts(root, _rows=rows)
    if story is not None:
        if story not in known:
            print(f"no doubts for story {story}; stories with doubts: "
                  f"{', '.join(known) or '(none)'}", file=sys.stderr)
            sys.exit(1)
        known = [story]
    ctx = (rows, _resolutions(root), _manifest(root))
    views = [_view(story_summary(s, root=root, _ctx=ctx),
                   "--all" in show, "--ungrouped" in show) for s in known]
    if "--json" in show:
        print(json.dumps({"stories": views}, indent=2, ensure_ascii=False))
        return
    for v in views:
        _print_story(v)


def _resolution_files(root):
    d = root / "resolutions"
    out = []
    for p in sorted(d.glob("*.md")) if d.is_dir() else []:
        if p.name == "index.md":
            continue
        fm = wiki.read_concept(p)[0]
        if isinstance(fm, dict):
            out.append((p, fm))
    return out


def _answer_errors(path, fm, root, doubt_ids):
    """L15 messages for one Resolution that carries `answers:` or any of the
    keys only an answer has (`effect`, `confirmed_parts`, `member_basis`).
    doubt_ids is the set of current doubt ids (the collected rows)."""
    pre = f"L15 resolutions/{path.name}: "
    missing = [k for k in ("answers", "card") if not _text(fm.get(k))]
    if missing:
        carried = [k for k in ("answers",) + LIFT_KEYS if k in fm]
        return [pre + f"carries {', '.join(carried)} without a valid "
                f"{' and '.join(missing)} - only `wiki doubts answer` writes "
                f"an answer Resolution"]
    qid = fm.get("answers")
    errs = []
    effect = fm.get("effect")
    if effect not in ANSWER_EFFECTS:
        errs.append(pre + f"answers {qid} but effect is {effect!r} "
                    f"(confirms or corrects)")
    field = "confirmed_parts" if effect == "confirms" else "member_basis"
    if effect in ANSWER_EFFECTS and not _map(fm.get(field)):
        errs.append(pre + f"effect {effect} needs a non-empty {field}")
    refs = fm.get("resolves")
    stories = {_story_of(r) for r in (refs if isinstance(refs, list) else [])} - {None}
    found = None
    for st in sorted(stories):
        sid = st[len("stories/"):]
        try:
            reg = load_register(sid, root=root)
        except ValueError:
            continue
        qs = _map(reg).get("questions")
        for q in (qs if isinstance(qs, list) else []):
            if isinstance(q, dict) and q.get("id") == qid:
                found = q
    if found is None:
        errs.append(pre + f"answers {qid}, which is no question of the "
                    f"register of the story it resolves")
        return errs
    if fm.get("status") != "asserted":
        return errs
    # A key whose doubt no longer exists (its test case or part left the
    # spec) is history, not an error: the register drops the member and the
    # Resolution keeps it. A key that is a current doubt must be a member.
    members = set(_members(found))
    for key in ("confirmed_parts", "member_basis"):
        for d in _map(fm.get(key)):
            if d not in members and d in doubt_ids:
                errs.append(pre + f"{key} names {d}, which is not a member "
                            f"of {qid}")
    return errs


ANSWER_EVENT = "Assertion-Event: cli-doubts-answer"


def _adding_commits(root):
    """{path: message of the latest commit that added it, or None when the
    latest event was a delete} for every file under resolutions/, from ONE
    `git log` (no rename detection, so a rename is a fresh add). {} when root
    is not in a git repository."""
    csep, nsep = "==COMMIT==", "==FILES=="
    r = subprocess.run(["git", "log", "--no-renames", "--diff-filter=AD",
                        "--name-status", "--relative",
                        f"--format={csep}%B{nsep}", "--", "resolutions"],
                       cwd=root, capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    out = {}
    if r.returncode != 0:
        return out
    for chunk in r.stdout.split(csep)[1:]:
        msg, _, names = chunk.partition(nsep)
        for ln in names.splitlines():
            status, _, path = ln.partition("\t")
            if path and path not in out:
                out[path] = msg if status == "A" else None
    return out


def _unevented_answer_errors(root, paths):
    """L15 messages for committed answer Resolutions whose adding commit is
    not a `doubts answer` commit. A file not yet committed is exempt: the
    commit gate and L10 see it when it is committed."""
    if not paths:
        return []
    adds = _adding_commits(root)
    out = []
    for p in paths:
        msg = adds.get(f"resolutions/{p.name}")
        if msg is not None and ANSWER_EVENT not in msg:
            out.append(f"L15 resolutions/{p.name}: answers a doubt question but "
                       f"was not added by `wiki doubts answer` (its adding "
                       f"commit carries no '{ANSWER_EVENT}')")
    return out


def _forged_lift_errors(root, resolutions):
    """L15 messages for rendered test cases whose lift is not backed by an
    asserted confirming answer Resolution."""
    by_id = {r.get("id"): r for r in resolutions if r.get("id")}
    out = []
    d = root / "testcases"
    for p in sorted(d.rglob("*.md")) if d.is_dir() else []:
        if p.name == "index.md":
            continue
        fm = wiki.read_concept(p)[0]
        parts = fm.get("confidence_parts") if isinstance(fm, dict) else None
        if not isinstance(parts, dict):
            continue
        pre = f"L15 {p.relative_to(root).as_posix()}: "
        sc = fm.get("scenario_id")
        for part, v in parts.items():
            v = _map(v)
            ref = v.get("resolution")
            if not ref:
                if v.get("authored") and v.get("level") == "High":
                    out.append(pre + f"confidence_parts.{part} is High with "
                               f"authored {v['authored']} but names no "
                               f"resolution")
                continue
            rid = Path(str(ref)).stem
            r = by_id.get(rid)
            if r is None:
                out.append(pre + f"confidence_parts.{part} names {ref}, which "
                           f"does not exist")
            elif r.get("status") != "asserted":
                out.append(pre + f"confidence_parts.{part} lifted by {rid}, "
                           f"which is not asserted")
            elif not is_answer_resolution(r):
                out.append(pre + f"confidence_parts.{part} lifted by {rid}, "
                           f"which is no answer Resolution (it needs answers, "
                           f"card and effect)")
            elif r.get("effect") != "confirms":
                out.append(pre + f"confidence_parts.{part} lifted by {rid}, "
                           f"which does not confirm (effect "
                           f"{r.get('effect')!r})")
            elif f"{sc}#{part}" not in _map(r.get("confirmed_parts")):
                out.append(pre + f"confidence_parts.{part} lifted by {rid}, "
                           f"which does not confirm {sc}#{part}")
            if not v.get("authored"):
                out.append(pre + f"confidence_parts.{part} has a resolution "
                           f"but no authored level")
    return out


def _observation_errors(root):
    """L15 messages for doubts/observations/*.yaml. Never raises on content:
    a non-string field is a message."""
    d = root / OBSERVATION_DIR
    out = []
    for p in sorted(d.glob("*.yaml")) if d.is_dir() else []:
        pre = f"L15 {OBSERVATION_DIR}/{p.name}: "
        try:
            obs = load_observations(p.stem, root=root)
            doc = yaml.safe_load(p.read_text(encoding="utf-8"))
        except ValueError as e:
            out.append(pre + _register_error_text(e))
            continue
        if doc.get("story") != f"/stories/{p.stem}.md":
            out.append(pre + f"story is {doc.get('story')!r}, expected "
                       f"'/stories/{p.stem}.md'")
        try:
            reg = load_register(p.stem, root=root)
        except ValueError:
            reg = None
        qs = _map(reg).get("questions")
        known = {q.get("id") for q in (qs if isinstance(qs, list) else [])
                 if isinstance(q, dict) and isinstance(q.get("id"), str)}
        for n, o in enumerate(obs, 1):
            at = f"observation {n}: "
            q = o.get("question")
            if not isinstance(q, str):
                out.append(pre + at + f"question must be a string, got {q!r}")
            elif q not in known:
                out.append(pre + at + f"question {q!r} is no question of "
                           f"doubts/{p.stem}.yaml")
            if "text" in o and not isinstance(o["text"], str):
                out.append(pre + at + "text must be a string")
            for key in ("text", "by", "at", "workbook_hash"):
                v = o.get(key)
                if not (str(v).strip() if v is not None else ""):
                    out.append(pre + at + f"{key} is missing or empty")
    return out


def l15_errors(root=None):
    """L15: every spec under tools/*_specs/ can be read; every doubts/*.yaml
    is a valid register; every Resolution carrying an answer key is a valid
    answer Resolution, names a real question and member parts, and was added
    by `wiki doubts answer`; every lifted test case part is backed by an
    asserted confirming answer Resolution."""
    root = _root(root)
    bad = []
    rows = collect_all(root=root, errors=bad)
    out = [f"L15 {_rel(root, p)}: {why}; its doubts are not collected"
           for p, why in bad]
    res = _resolutions(root)
    for p in _register_files(root):
        sid, pre = p.stem, f"L15 doubts/{p.name}: "
        try:
            reg = load_register(sid, root=root)
        except ValueError as e:
            out.append(pre + _register_error_text(e))
            continue
        fm, body, exists = _story_concept(root, sid)
        if not exists:
            out.append(pre + f"story {sid} has no concept stories/{sid}.md")
        try:
            errs = register_errors(reg, fm, rows, res, body)
        except Exception as e:  # never crash lint on register content
            errs = [f"register: cannot be validated ({type(e).__name__})"]
        out.extend(pre + e for e in errs)
    files = _resolution_files(root)
    doubt_ids = {r["doubt"] for r in rows}
    answering = []
    for path, fm in files:
        if fm.get("answers") or any(k in fm for k in LIFT_KEYS):
            out.extend(_answer_errors(path, fm, root, doubt_ids))
        if fm.get("answers"):
            answering.append(path)
    out.extend(_unevented_answer_errors(root, answering))
    out.extend(_forged_lift_errors(root, [fm for _p, fm in files]))
    out.extend(_observation_errors(root))
    return out


def _rel(root, p):
    try:
        return Path(p).relative_to(root).as_posix()
    except ValueError:
        return Path(p).name


def w8_warnings(root=None):
    """W8: one line per story with ungrouped non-closed doubts."""
    root = _root(root)
    rows = collect_all(root=root)
    ctx = (rows, _resolutions(root), _manifest(root))
    out = []
    for sid in stories_with_doubts(root, _rows=rows):
        s = story_summary(sid, root=root, _ctx=ctx)
        if s["register_errors"] or not s["ungrouped"]:
            continue
        out.append(f"W8 {sid}: {s['ungrouped']} ungrouped doubt(s) - group them "
                   f"in doubts/{sid}.yaml (tc-resolve)")
    return out


def w9_warnings(root=None):
    """W9: one line per story with a confirmed lift not yet rendered, or a
    rendered lift that no longer holds."""
    root = _root(root)
    by_story = {}
    for e in unrendered_lifts(root):
        by_story.setdefault(e["story"], []).append(e)
    out = []
    for sid, es in sorted(by_story.items()):
        cmds = []
        if any(e["kind"] == "sit" for e in es):
            cmds.append(f"py tools/render_sit.py --story {sid}")
        for fid in sorted({e["flow"] for e in es
                           if e["kind"] == "uat" and e["flow"]}):
            cmds.append(f"py tools/render_uat.py --flow {fid}")
        out.append(f"W9 {sid}: {len(es)} confirmed doubt(s) not rendered - "
                   f"run: {' ; '.join(cmds)}")
    return out
