#!/usr/bin/env python3
"""Plain-assert tests for the Workbook view's read model
(run: py tools/test_app_workbook_model.py). No pytest in this repo.

Hermetic: every workbook is rendered from a synthetic selection into a temp
inventory directory, and the wiki the model reads is a stub, so these hold on
an empty bundle and never touch build/.
"""
import json
import os
import shutil
import sys
import tempfile
import traceback
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "tools/app"))
import wiki_suite
import workbook_model as wm

MAIN = "C-TC-1 (Main flow)"
HASH = "sha256:" + "a" * 64


def _tc(i, kind="sit", status="active"):
    tid = f"1.1-AC01-{i:02d}" if kind == "sit" else f"UAT-1.1-AC01-{i:02d}"
    fm = {"type": "Test Case", "id": tid, "kind": kind, "status": status,
          "module": "/modules/m.md", "scenario_id": f"SC-{kind}-{i:02d}",
          "covers": ["/stories/US-X.md#AC1"], "title": f"Case {i}",
          "confidence": "Low",
          "confidence_parts": {"data": {"level": "Low", "remark": f"Inferred: data {i}."}}}
    if kind == "sit":
        fm["run"], fm["section"] = "Main flow", "Step 1"
    body = ("# Steps\n1. Click **Save**.\n\n"
            "# Expected Results\n1. The **Saved** banner shows.\n")
    rel = (f"testcases/sit/m/{tid}" if kind == "sit" else f"testcases/uat/{tid}")
    return (rel, fm, body)


def _wiki(tcs):
    concepts = {
        "modules/m": ({"type": "Module", "title": "Alpha Module",
                       "description": "PRD section §1"}, "", None),
        "stories/US-X": ({"type": "User Story", "title": "Story X",
                          "description": "§1.1"}, "", None),
    }
    for rel, fm, body in tcs:
        concepts[rel] = (fm, body, None)
    manifest = {"tc_hashes": {rel: HASH for rel, _f, _b in tcs}}
    return concepts, manifest


def _doubts(tcs):
    """One question whose only member is the first test case's data part."""
    sid = tcs[0][1]["scenario_id"]
    mem = {"doubt": f"{sid}#data", "scenario_id": sid, "part": "data",
           "level": "Low", "state": "open", "resolution": None, "remark": "r"}
    q = {"id": "Q-US-X-01", "question": "Which value?", "status": "open",
         "members": [mem]}
    return [{"story": "US-X", "questions": [q], "ungrouped_doubts": [],
             "doubts": [mem], "register_errors": [], "register": True}]


class Bench:
    """A temp inventory with one compiled workbook and a stub wiki."""

    def __init__(self, tcs=None, kind="sit", name="demo", sidecar=True):
        self.tcs = tcs if tcs is not None else [_tc(1), _tc(2), _tc(3)]
        self.kind, self.name = kind, name
        self.tmp = Path(tempfile.mkdtemp(prefix="wbv-"))
        self.saved = (wm.INVENTORY_DIR, wm.SUITES_DIR, wm.STORIES_DIR, wm.load_all)
        wm.INVENTORY_DIR = self.tmp / "inventory"
        wm.SUITES_DIR, wm.STORIES_DIR = self.tmp / "suites", self.tmp / "stories"
        self.concepts, self.manifest = _wiki(self.tcs)
        wm.load_all = lambda: (self.concepts, self.manifest)
        wm._CACHE.clear()
        folder = wm.INVENTORY_DIR / kind
        folder.mkdir(parents=True)
        self.path = folder / f"{name}-latest.xlsx"
        order = wiki_suite.render_xlsx(self.tcs, "t", self.path,
                                       concepts=self.concepts,
                                       manifest=self.manifest,
                                       doubt_summaries=_doubts(self.tcs))
        if sidecar:
            wiki_suite.write_sidecar(self.path, {"type": "suite", "name": name},
                                     order, self.manifest)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        (wm.INVENTORY_DIR, wm.SUITES_DIR, wm.STORIES_DIR, wm.load_all) = self.saved
        wm._CACHE.clear()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def view(self):
        return wm.workbook(self.kind, self.path.name)


def _sheet(view, name):
    return next(s for s in view["sheets"] if s["name"] == name)


def _text(cell):
    return "".join(r["text"] for r in cell["runs"])


# ---- parse: order, widths, frozen rows, fills, borders, alignment, rich text

def test_sheets_come_back_in_file_order():
    with Bench() as b:
        names = [s["name"] for s in b.view()["sheets"]]
    assert names == ["A - Table of Contents", "B - Proj Doc References", MAIN,
                     "AI Doubts", "Test Statistics", "Change Log"], names


def test_column_widths_frozen_rows_and_row_heights():
    with Bench() as b:
        ws = _sheet(b.view(), MAIN)
    assert ws["cols"] == [float(wiki_suite.COL_W[c]) for c in "ABCDEFGHIJKLMN"], ws["cols"]
    assert ws["frozen_rows"] == wiki_suite.TC_HEADER_ROW
    assert ws["rows"][wiki_suite.TC_HEADER_ROW - 1]["height"] == 38
    assert ws["rows"][0]["height"] is None        # unset in the file


def test_header_and_section_fills_borders_and_alignment():
    with Bench() as b:
        ws = _sheet(b.view(), MAIN)
    header = ws["rows"][wiki_suite.TC_HEADER_ROW - 1]["cells"]
    assert header[0]["fill"] == "D9D9D9" and header[0]["bold"], header[0]
    assert header[0]["align"] == "center" and header[0]["valign"] == "center"
    assert header[0]["wrap"] and header[0]["border"], header[0]
    assert header[7]["fill"] == "92D050", header[7]       # execution grid header
    section = next(r for r in ws["rows"]
                   if r["cells"][0]["runs"] and _text(r["cells"][0]) == "Section: Step 1")
    assert all(c["fill"] == "DDEBF7" and c["border"] for c in section["cells"]), section
    assert section["height"] == 18
    blue = ws["rows"][wiki_suite.TC_HEADER_ROW]["cells"][0]   # the module line
    assert blue["color"] == "0070C0" and blue["bold"] and blue["fill"] is None, blue


def test_test_case_row_has_bold_runs_wrap_and_an_empty_execution_grid():
    with Bench() as b:
        ws = _sheet(b.view(), MAIN)
    row = next(r for r in ws["rows"] if _text(r["cells"][0]) == "TC-1.1-AC01-01")
    steps = row["cells"][2]
    assert {"text": "Save", "bold": True} in steps["runs"], steps
    assert _text(steps) == "1. Click Save.", steps
    assert steps["wrap"] and steps["valign"] == "top" and steps["border"], steps
    conf = row["cells"][5]
    assert _text(conf) == "Low" and conf["fill"] == "FFC7CE" and conf["align"] == "center"
    for cell in row["cells"][wiki_suite.N_AUTHORED:]:
        assert cell["runs"] == [] and cell["border"] and cell["fill"] is None, cell


def test_title_cells_carry_their_font_size():
    with Bench() as b:
        ws = _sheet(b.view(), MAIN)
    title = ws["rows"][1]["cells"][1]
    assert _text(title) == "AGILE TEST SPECIFICATIONS - Test Cases" and title["size"] == 12
    assert "size" not in ws["rows"][wiki_suite.TC_HEADER_ROW - 1]["cells"][0]


def _hand_made(folder, name="hand-latest.xlsx"):
    """A workbook the exporter never writes: a merge, an unknown function, a
    circular reference, a broken sheet reference."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "merged"
    ws.merge_cells("A1:C2")
    ws["A4"] = "=SUM(B5:B6)"
    ws["B4"] = "=B4+1"
    ws["C4"] = "='No such sheet'!A1"
    ws["A5"] = "=IFERROR(1/0,\"n/a\")"
    ws["B5"] = 2
    ws["B6"] = "=B5*3-1"
    ws["C5"] = "=COUNTIF(B5:B6,\">1\")"
    ws["C6"] = "=VLOOKUP(\"x\",A1:B2,2,FALSE)"
    wb.save(folder / name)
    return name


def test_merged_cells_are_reported_as_ranges():
    with Bench() as b:
        name = _hand_made(b.path.parent)
        ws = wm.workbook("sit", name)["sheets"][0]
    assert ws["merges"] == [[1, 1, 2, 3]], ws["merges"]
    assert _text(ws["rows"][0]["cells"][0]) == "merged"
    assert ws["rows"][0]["cells"][1]["runs"] == []        # covered by the merge
    assert ws["cols"] == [wm.DEFAULT_COL_WIDTH] * 3, ws["cols"]


def test_a_formula_the_model_does_not_evaluate_keeps_its_text():
    with Bench() as b:
        name = _hand_made(b.path.parent)
        a4 = wm.workbook("sit", name)["sheets"][0]["rows"][3]["cells"][0]
    assert a4["formula"] == "=SUM(B5:B6)" and _text(a4) == "=SUM(B5:B6)", a4


def test_parse_locates_test_case_ids_and_question_ids():
    with Bench() as b:
        parsed = wm._parsed(b.path)
    names = [s["name"] for s in parsed["sheets"]]
    assert [(names[si], tid) for si, _r, tid in parsed["tc_rows"]] == [
        (MAIN, f"TC-1.1-AC01-{i:02d}") for i in (1, 2, 3)], parsed["tc_rows"]
    assert all(r > wiki_suite.TC_HEADER_ROW for _si, r, _t in parsed["tc_rows"])
    assert parsed["questions"] == {"Q-US-X-01": {"sheet": "AI Doubts", "row": 2}}


def test_parse_is_cached_by_path_and_mtime():
    with Bench() as b:
        b.view()
        calls = []
        real = wm._parse
        wm._parse = lambda p: calls.append(p) or real(p)
        try:
            b.view()
            assert calls == [], "an unchanged file must not be parsed twice"
            os.utime(b.path, ns=(b.path.stat().st_atime_ns,
                                 b.path.stat().st_mtime_ns + 5_000_000_000))
            b.view()
            assert len(calls) == 1, calls
        finally:
            wm._parse = real


# ---- path safety and unreadable files

def test_traversal_and_non_xlsx_paths_are_refused():
    with Bench() as b:
        (b.path.parent / "notes.md").write_text("x", encoding="utf-8")
        outside = b.tmp / "outside.xlsx"
        shutil.copyfile(b.path, outside)
        for kind, file in (("sit", "notes.md"), ("sit", "missing-latest.xlsx"),
                           ("..", "outside.xlsx"), ("sit", "../../outside.xlsx"),
                           ("sit/..", "outside.xlsx"), ("", "demo-latest.xlsx"),
                           ("sit", str(outside))):
            try:
                wm.workbook(kind, file)
            except wm.WorkbookNotFound:
                continue
            raise AssertionError(f"served {kind!r}/{file!r}")


def test_a_corrupt_workbook_is_unreadable_not_a_traceback():
    with Bench() as b:
        (b.path.parent / "broken-latest.xlsx").write_bytes(b"PK\x03\x04 nope")
        try:
            wm.workbook("sit", "broken-latest.xlsx")
        except wm.WorkbookUnreadable as e:
            assert "broken-latest.xlsx" in str(e), e
        else:
            raise AssertionError("a corrupt workbook was served")


def test_the_view_is_json_serialisable():
    with Bench() as b:
        json.dumps(b.view())


# ---- formulas

def test_concatenate_and_cross_sheet_references_resolve():
    with Bench() as b:
        view = b.view()
    cfg = (wiki_suite.load_config().get("export") or {})
    want = f"{cfg.get('project_code', '')}/ {cfg.get('project_name', '')}"
    refs = _sheet(view, "B - Proj Doc References")["rows"][2]["cells"][2]
    assert _text(refs) == want and "formula" not in refs, refs
    b4 = _sheet(view, MAIN)["rows"][3]["cells"][1]
    assert _text(b4) == want and "formula" not in b4, b4


def test_countif_with_a_wildcard_counts_the_test_case_rows():
    with Bench() as b:
        c3 = _sheet(b.view(), MAIN)["rows"][2]["cells"][2]
    assert _text(c3) == "Total TC = 3" and "formula" not in c3, c3


def test_statistics_block_resolves_counta_countif_vlookup_iferror_and_sums():
    with Bench() as b:
        rows = _sheet(b.view(), "Test Statistics")["rows"]
    cell = lambda coord: rows[int(coord[1:]) - 1]["cells"]["ABC".index(coord[0])]
    assert _text(cell("B12")) == "3", cell("B12")          # =COUNTA(B$22:B$24)
    assert _text(cell("B3")) == "3", cell("B3")            # =B12
    assert _text(cell("B13")) == "0", cell("B13")          # =B14+B15
    assert _text(cell("B14")) == "0", cell("B14")          # =COUNTIF(C$22:C$24,"Pass")
    assert _text(cell("C14")) == "0.0%", cell("C14")       # =IFERROR(B14/B$12,0), 0.0%
    assert _text(cell("C4")) == "0.0%", cell("C4")         # =IFERROR(B4/$B$3,0)
    assert _text(cell("B22")) == "TC-1.1-AC01-01"
    assert _text(cell("C22")) == "0", cell("C22")          # =VLOOKUP(...) of a blank Status
    assert not any("formula" in c for r in rows for c in r["cells"])


def test_a_filled_status_column_flows_through_vlookup_and_countif():
    from openpyxl import load_workbook
    with Bench() as b:
        wb = load_workbook(b.path, rich_text=True)
        ws = wb[MAIN]
        first = next(r for r in range(9, 40)
                     if str(ws.cell(row=r, column=1).value).startswith("TC-"))
        ws.cell(row=first, column=wiki_suite.STATUS_COL, value="Pass")
        wb.save(b.path)
        wm._CACHE.clear()
        rows = _sheet(b.view(), "Test Statistics")["rows"]
    assert _text(rows[21]["cells"][2]) == "Pass", rows[21]["cells"][2]
    assert _text(rows[13]["cells"][1]) == "1"              # B14 Pass count
    assert _text(rows[13]["cells"][2]) == "33.3%"          # C14
    assert _text(rows[3]["cells"][1]) == "1"               # B4 overall Pass


def test_unknown_and_failing_formulas_fall_back_to_their_text():
    with Bench() as b:
        name = _hand_made(b.path.parent)
        rows = wm.workbook("sit", name)["sheets"][0]["rows"]
    a4, b4, c4 = rows[3]["cells"]
    assert a4["formula"] == "=SUM(B5:B6)" and _text(a4) == "=SUM(B5:B6)", a4
    assert b4["formula"] == "=B4+1", b4                    # circular
    assert c4["formula"] == "='No such sheet'!A1", c4      # #REF!
    a5, b5, c5 = rows[4]["cells"]
    assert _text(a5) == "n/a" and "formula" not in a5, a5  # IFERROR catches #DIV/0!
    assert c5["formula"] == '=COUNTIF(B5:B6,">1")', c5     # comparison criteria
    b6, c6 = rows[5]["cells"][1], rows[5]["cells"][2]
    assert _text(b6) == "5" and "formula" not in b6, b6    # arithmetic precedence
    assert c6["formula"].startswith("=VLOOKUP"), c6        # #N/A


# ---- test case rows, freshness, versions, inventory

def test_every_tc_row_maps_to_its_ref():
    with Bench() as b:
        view = b.view()
    got = [(r["tc"]["ref"], r["tc"]["id"], _text(r["cells"][0]))
           for r in _sheet(view, MAIN)["rows"] if r["tc"]]
    assert got == [(f"testcases/sit/m/1.1-AC01-{i:02d}", f"1.1-AC01-{i:02d}",
                    f"TC-1.1-AC01-{i:02d}") for i in (1, 2, 3)], got
    assert all(r["tc"] is None for s in view["sheets"] if s["name"] != MAIN
               for r in s["rows"])
    assert view["freshness"] == "fresh" and view["changed_count"] == 0, view["freshness"]
    assert view["source"] == {"type": "suite", "name": "demo"}
    assert view["name"] == "demo" and view["kind"] == "sit"


def test_links_locate_test_case_rows_and_question_rows():
    with Bench() as b:
        view = b.view()
    row = next(i for i, r in enumerate(_sheet(view, MAIN)["rows"], 1) if r["tc"])
    assert view["links"]["tcs"]["TC-1.1-AC01-01"] == {"sheet": MAIN, "row": row}
    assert view["links"]["questions"] == {"Q-US-X-01": {"sheet": "AI Doubts", "row": 2}}


def test_a_changed_hash_flags_its_row_and_the_count():
    with Bench() as b:
        b.manifest["tc_hashes"]["testcases/sit/m/1.1-AC01-02"] = "sha256:" + "b" * 64
        view = b.view()
    flags = {r["tc"]["id"]: r["tc"]["changed"]
             for r in _sheet(view, MAIN)["rows"] if r["tc"]}
    assert flags == {"1.1-AC01-01": False, "1.1-AC01-02": True, "1.1-AC01-03": False}
    assert view["freshness"] == "changed" and view["changed_count"] == 1


def test_stale_retired_and_missing_test_cases_are_changed():
    with Bench() as b:
        b.concepts["testcases/sit/m/1.1-AC01-01"][0]["status"] = "stale"
        b.concepts["testcases/sit/m/1.1-AC01-02"][0]["status"] = "retired"
        del b.concepts["testcases/sit/m/1.1-AC01-03"]
        view = b.view()
    rows = [r["tc"] for r in _sheet(view, MAIN)["rows"] if r["tc"]]
    assert [t["changed"] for t in rows] == [True, True, True], rows
    assert rows[2]["id"] == "1.1-AC01-03"                 # id from the ref's stem
    assert view["changed_count"] == 3 and view["freshness"] == "changed"


def test_no_sidecar_is_unknown_and_infers_the_source_from_the_name():
    with Bench(sidecar=False) as b:
        view = b.view()
        assert view["freshness"] == "unknown" and view["changed_count"] == 0
        assert view["source"] is None
        assert [r["tc"]["changed"] for r in _sheet(view, MAIN)["rows"] if r["tc"]] \
            == [False, False, False]                      # still linked, never flagged
        wm.SUITES_DIR.mkdir()
        (wm.SUITES_DIR / "demo.yaml").write_text("kind: sit\n", encoding="utf-8")
        assert b.view()["source"] == {"type": "suite", "name": "demo"}
    with Bench(sidecar=False, name="US-X-sit") as b:
        wm.STORIES_DIR.mkdir()
        (wm.STORIES_DIR / "US-X.md").write_text("---\n---\n", encoding="utf-8")
        assert b.view()["source"] == {"type": "export", "name": "US-X-sit",
                                      "story": "US-X"}


def test_a_damaged_sidecar_is_unknown_not_a_crash():
    for text in ("{not json", "[]", '{"tcs": ["a"]}'):
        with Bench() as b:
            b.path.with_suffix(".json").write_text(text, encoding="utf-8")
            assert b.view()["freshness"] == "unknown", text


def test_mixed_workbook_maps_a_shared_display_id_to_both_refs():
    """SIT 1.1-AC01-01 and UAT-1.1-AC01-01 both display as TC-1.1-AC01-01."""
    tcs = [_tc(1, "uat"), _tc(1, "sit")]
    with Bench(tcs=tcs, kind="mixed") as b:
        view = b.view()
    refs = [r["tc"]["ref"] for s in view["sheets"] for r in s["rows"] if r["tc"]]
    assert refs == ["testcases/uat/UAT-1.1-AC01-01",
                    "testcases/sit/m/1.1-AC01-01"], refs


def test_versions_list_latest_first_then_newest_timestamp():
    with Bench() as b:
        for ts in ("20261001-090000", "20261003-101500"):
            shutil.copyfile(b.path, b.path.parent / f"demo_{ts}.xlsx")
        shutil.copyfile(b.path, b.path.parent / "other-latest.xlsx")
        view = b.view()
        old = wm.workbook("sit", "demo_20261001-090000.xlsx")
    assert [v["file"] for v in view["versions"]] == [
        "demo-latest.xlsx", "demo_20261003-101500.xlsx", "demo_20261001-090000.xlsx"]
    assert view["versions"][2]["compiled_at"].startswith("2026-10-01T09:00:00")
    assert old["file"] == "demo_20261001-090000.xlsx" and old["freshness"] == "unknown"
    assert old["name"] == "demo" and len(old["versions"]) == 3


def test_workbooks_lists_one_entry_per_name_with_tc_locations():
    with Bench() as b:
        shutil.copyfile(b.path, b.path.parent / "demo_20261003-101500.xlsx")
        (b.path.parent / "~$demo-latest.xlsx").write_bytes(b"lock")
        (b.path.parent / "broken-latest.xlsx").write_bytes(b"PK\x03\x04 nope")
        both = wm.workbooks()
        listing = both["workbooks"]
        json.dumps(both)
    assert [w["name"] for w in listing] == ["demo"], [w["name"] for w in listing]
    # the file that could not be read is named, not dropped without a word
    assert [(s["kind"], s["file"]) for s in both["skipped"]] == \
        [("sit", "broken-latest.xlsx")], both["skipped"]
    assert "cannot read broken-latest.xlsx" in both["skipped"][0]["error"]
    w = listing[0]
    assert w["file"] == "demo-latest.xlsx" and w["kind"] == "sit"
    assert w["freshness"] == "fresh" and w["source"]["type"] == "suite"
    assert [v["file"] for v in w["versions"]] == ["demo-latest.xlsx",
                                                  "demo_20261003-101500.xlsx"]
    loc = w["tcs"]["testcases/sit/m/1.1-AC01-02"]
    assert loc["sheet"] == MAIN and isinstance(loc["row"], int), loc
    assert "sheets" not in w


def test_workbooks_is_empty_without_an_inventory():
    with Bench() as b:
        shutil.rmtree(wm.INVENTORY_DIR)
        assert wm.workbooks() == {"workbooks": [], "skipped": []}


def test_a_corrupt_latest_falls_back_to_the_newest_timestamped_compile():
    """A half-written -latest must not hide the workbook's name: the listing
    serves the newest compile that can be read and says what it passed over."""
    with Bench() as b:
        for stamp in ("20261003-101500", "20261004-090000"):
            shutil.copyfile(b.path, b.path.parent / f"demo_{stamp}.xlsx")
            shutil.copyfile(b.path.with_suffix(".json"),
                            b.path.parent / f"demo_{stamp}.json")
        b.path.write_bytes(b"PK\x03\x04 half written")
        wm._CACHE.clear()
        both = wm.workbooks()
    assert [(w["name"], w["file"]) for w in both["workbooks"]] == \
        [("demo", "demo_20261004-090000.xlsx")], both["workbooks"]
    assert both["workbooks"][0]["freshness"] == "fresh"
    assert "testcases/sit/m/1.1-AC01-02" in both["workbooks"][0]["tcs"]
    assert [s["file"] for s in both["skipped"]] == ["demo-latest.xlsx"], both["skipped"]
    assert both["skipped"][0]["kind"] == "sit"
    assert str(b.tmp) not in both["skipped"][0]["error"], "an error never carries a path"


# ---- fix round 1: row mapping that cannot swap, and views that cannot crash

def _interleaved():
    """Selected in tc_sort_key order: a journey UAT case, SIT 1.1-AC01-01 (run
    "Main flow"), then the non-journey UAT-1.1-AC01-01. plan_sheets regroups
    them, so the UAT rows sit on C-TC-1 and the SIT row on C-TC-2, and the two
    ids that share the display id TC-1.1-AC01-01 swap places."""
    journey = _tc(1, "uat")
    journey[1].update(id="UAT-0.1-AC01-01", covers=["/flows/f.md#J1"])
    journey = ("testcases/uat/UAT-0.1-AC01-01", journey[1], journey[2])
    tcs = [journey, _tc(1, "sit"), _tc(1, "uat")]
    keys = [wiki_suite.tc_sort_key(fm) for _r, fm, _b in tcs]
    assert keys == sorted(keys), keys
    return tcs


def _rewrite_sidecar(b, edit):
    path = b.path.with_suffix(".json")
    doc = json.loads(path.read_text(encoding="utf-8"))
    edit(doc)
    path.write_text(json.dumps(doc), encoding="utf-8")
    wm._CACHE.clear()


def _tc_refs(view):
    return [(r["tc"] or {}).get("ref") for s in view["sheets"] for r in s["rows"]
            if s["name"].startswith("C-TC") and r["cells"] and r["cells"][0]["runs"]
            and _text(r["cells"][0]).startswith("TC-")]


def test_interleaved_selection_maps_every_row_to_its_own_ref():
    with Bench(tcs=_interleaved(), kind="mixed") as b:
        view = b.view()
    names = [s["name"] for s in view["sheets"] if s["name"].startswith("C-TC")]
    assert len(names) == 2, names
    assert _tc_refs(view) == ["testcases/uat/UAT-0.1-AC01-01",
                              "testcases/uat/UAT-1.1-AC01-01",
                              "testcases/sit/m/1.1-AC01-01"], _tc_refs(view)


def test_interleaved_selection_flags_the_changed_row_it_belongs_to():
    with Bench(tcs=_interleaved(), kind="mixed") as b:
        b.manifest["tc_hashes"]["testcases/sit/m/1.1-AC01-01"] = "sha256:" + "b" * 64
        view = b.view()
    flagged = [(s["name"], r["tc"]["ref"]) for s in view["sheets"]
               for r in s["rows"] if r["tc"] and r["tc"]["changed"]]
    assert flagged == [(names_sit(view), "testcases/sit/m/1.1-AC01-01")], flagged


def names_sit(view):
    return next(s["name"] for s in view["sheets"] if "Main flow" in s["name"])


def test_an_old_format_sidecar_links_only_unambiguous_ids():
    with Bench(tcs=_interleaved(), kind="mixed") as b:
        _rewrite_sidecar(b, lambda d: d.pop("tcs_order"))
        view = b.view()
    assert _tc_refs(view) == ["testcases/uat/UAT-0.1-AC01-01", None, None], _tc_refs(view)
    assert view["freshness"] == "fresh" and view["changed_count"] == 0


def test_an_old_format_sidecar_still_computes_freshness_for_unlinked_refs():
    with Bench(tcs=_interleaved(), kind="mixed") as b:
        b.manifest["tc_hashes"]["testcases/sit/m/1.1-AC01-01"] = "sha256:" + "b" * 64
        _rewrite_sidecar(b, lambda d: d.pop("tcs_order"))
        view = b.view()
    assert view["freshness"] == "changed" and view["changed_count"] == 1


def test_a_count_mismatch_leaves_that_id_unlinked():
    with Bench(tcs=_interleaved(), kind="mixed") as b:
        _rewrite_sidecar(b, lambda d: d["tcs"].pop("testcases/uat/UAT-1.1-AC01-01"))
        view = b.view()
    assert _tc_refs(view) == ["testcases/uat/UAT-0.1-AC01-01", None, None], _tc_refs(view)


def test_a_formula_that_raises_something_odd_keeps_its_text():
    from openpyxl import Workbook
    with Bench() as b:
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"] = "=" + "9" * 5000                        # literal too long for int()
        ws["A2"] = "=" + "9" * 400 + "*1.5"                # int too large for float
        ws["B1"] = "nan"
        ws["B2"] = "inf"
        ws["A3"] = '=VLOOKUP("x",A1:B2,B1,FALSE)'          # int(nan) -> ValueError
        ws["A4"] = '=VLOOKUP("x",A1:B2,B2,FALSE)'          # int(inf) -> OverflowError
        wb.save(b.path.parent / "odd-latest.xlsx")
        rows = wm.workbook("sit", "odd-latest.xlsx")["sheets"][0]["rows"]
    for i in range(4):
        cell = rows[i]["cells"][0]
        assert "formula" in cell and _text(cell).startswith("="), (i, cell["runs"][0]["text"][:30])


def test_a_failure_after_load_is_unreadable_and_skipped_by_the_listing():
    with Bench() as b:
        real = wm._Book
        def boom(_wb):
            raise KeyError("odd")
        wm._Book = boom
        try:
            try:
                wm.workbook("sit", "demo-latest.xlsx")
            except wm.WorkbookUnreadable as e:
                assert "demo-latest.xlsx" in str(e), e
            else:
                raise AssertionError("served a workbook whose parse failed")
            listing = wm.workbooks()
            assert listing["workbooks"] == [], listing
            assert [s["file"] for s in listing["skipped"]] == ["demo-latest.xlsx"]
        finally:
            wm._Book = real


def test_the_listing_skips_a_file_that_fails_for_any_reason():
    with Bench() as b:
        shutil.copyfile(b.path, b.path.parent / "bad-latest.xlsx")
        real = wm._describe
        def fake(path, *a):
            if path.name.startswith("bad"):
                raise RuntimeError("odd")
            return real(path, *a)
        wm._describe = fake
        try:
            listing = wm.workbooks()
            assert [w["name"] for w in listing["workbooks"]] == ["demo"]
            assert listing["skipped"] == [{
                "kind": "sit", "file": "bad-latest.xlsx",
                "error": "cannot read bad-latest.xlsx: RuntimeError: odd"}]
        finally:
            wm._describe = real


def test_an_impossible_timestamp_in_a_file_name_falls_back_to_mtime():
    with Bench() as b:
        shutil.copyfile(b.path, b.path.parent / "demo_20269999-999999.xlsx")
        view = b.view()
        listing = wm.workbooks()["workbooks"]
        odd = wm.workbook("sit", "demo_20269999-999999.xlsx")
    assert len(view["versions"]) == 2 and listing[0]["name"] == "demo"
    assert odd["compiled_at"].startswith(str(datetime.now().year)), odd["compiled_at"]


def test_a_sidecar_must_have_every_field_in_the_right_type():
    good = {"compiled_at": "2026-10-04T10:00:00+00:00", "source": {"type": "suite"},
            "tcs": {"testcases/sit/m/1.1-AC01-01": HASH}}
    bad = [{"tcs": {}}, dict(good, source=5, tcs={"a": 1}), dict(good, source=None),
           dict(good, compiled_at=7), dict(good, tcs={"a": 1}),
           {k: v for k, v in good.items() if k != "tcs"}]
    for doc in bad:
        with Bench() as b:
            b.path.with_suffix(".json").write_text(json.dumps(doc), encoding="utf-8")
            assert b.view()["freshness"] == "unknown", doc
    with Bench() as b:
        doc = dict(good, tcs={"testcases/sit/m/1.1-AC01-01": None})
        b.path.with_suffix(".json").write_text(json.dumps(doc), encoding="utf-8")
        assert b.view()["freshness"] == "changed"          # null = never sealed


def test_a_kind_outside_the_known_four_is_refused():
    with Bench() as b:
        for kind in ("SIT", "uat/../sit", "custom", "sit "):
            try:
                wm.workbook(kind, "demo-latest.xlsx")
            except wm.WorkbookNotFound:
                continue
            raise AssertionError(f"served kind {kind!r}")


def test_an_os_error_message_never_carries_the_directory():
    import openpyxl
    with Bench() as b:
        real = openpyxl.load_workbook
        def locked(path, *a, **k):
            raise PermissionError(13, "Permission denied", str(path))
        openpyxl.load_workbook = locked
        try:
            wm.workbook("sit", "demo-latest.xlsx")
        except wm.WorkbookUnreadable as e:
            msg = str(e)
            assert "demo-latest.xlsx" in msg and "Permission denied" in msg, msg
            assert str(b.path.parent) not in msg and "wbv-" not in msg, msg
        else:
            raise AssertionError("served an unreadable workbook")
        finally:
            openpyxl.load_workbook = real


if __name__ == "__main__":
    failed = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"[PASS] {name}")
            except Exception as e:
                failed += 1
                print(f"[FAIL] {name}: {e!r}")
                traceback.print_exc()
    print("test_app_workbook_model " + ("FAILED" if failed else "OK"))
    sys.exit(1 if failed else 0)
